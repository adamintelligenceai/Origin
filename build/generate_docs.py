"""Generate Evidence Room Word templates and PDFs."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from content_data import AGENTS, EXCEPTIONS, RESPONSIBILITY_LEVELS, BRAND

ROOT = Path(__file__).resolve().parent.parent / "evidence-room"
ACCENT = colors.HexColor("#0F4C5C")


def docx_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def docx_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p


def save_docx(doc, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"  Created {path.relative_to(ROOT.parent)}")


def gen_sop_template():
    doc = Document()
    doc.add_heading("Standard Operating Procedure Template", 0)
    doc.add_paragraph("Evidence Room AP Agent OS — Process SOP")
    sections = [
        ("1. Purpose", "[Describe the purpose of this AP process and its role in the invoice-to-pay lifecycle]"),
        ("2. Scope", "[Define which entities, invoice types, and value ranges this SOP covers]"),
        ("3. Definitions", "[Key terms: PO, GR, STP, exception, DOA, etc.]"),
        ("4. Roles & Responsibilities", "[RACI for each step — Human roles and assigned agents]"),
        ("5. Process Steps", "[Numbered steps with decision points, system actions, and agent involvement]"),
        ("6. Business Rules", "[Matching tolerances, approval thresholds, coding requirements]"),
        ("7. Exception Handling", "[Link to exception taxonomy codes and resolution paths]"),
        ("8. Controls", "[Preventive and detective controls with evidence requirements]"),
        ("9. Systems", "[ERP modules, automation tools, agent platforms involved]"),
        ("10. KPIs", "[Metrics used to measure this process]"),
        ("11. Escalation", "[When and how to escalate unresolved items]"),
        ("12. Revision History", "[Version, date, author, changes]"),
    ]
    for title, placeholder in sections:
        docx_heading(doc, title, 2)
        docx_para(doc, placeholder)
        docx_para(doc, "")
    docx_heading(doc, "EXAMPLE: Invoice PO Matching SOP", 1)
    docx_para(doc, "Purpose: Match supplier invoices to purchase orders and goods receipts within defined tolerances.")
    docx_para(doc, "Agent: AP-MAT-03 Matching Agent (Autonomy Level 2 — Prepare)")
    steps = ["Receive validated invoice from AP-VAL-02", "Retrieve PO and receipt data from ERP", "Perform line-level two-way/three-way match", "Apply tolerance rules by category", "Auto-match within tolerance → route to approval", "Exception → classify via AP-TRI-04 and assign owner"]
    for i, s in enumerate(steps, 1):
        docx_para(doc, f"Step {i}: {s}")
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Templates/SOP_Template.docx")


def gen_agent_charter():
    doc = Document()
    doc.add_heading("Agent Charter Template", 0)
    doc.add_paragraph("Evidence Room AP Agent OS")
    fields = [
        "Agent Name", "Agent Code", "Version", "Date", "Author",
        "Purpose", "Scope", "Inputs", "Tools & Data Access",
        "Responsibilities", "Explicit Exclusions", "Human Owner",
        "Approval Requirements", "Escalation Criteria", "Output Standard",
        "Control Requirements", "Audit Evidence", "KPIs",
        "Current Autonomy Level", "Target Autonomy Level",
        "Failure Handling", "Cost Monitoring", "Performance History",
    ]
    for f in fields:
        docx_heading(doc, f, 2)
        docx_para(doc, f"[{f}]")
    docx_heading(doc, "EXAMPLE: Invoice Validation Agent Charter", 1)
    a = AGENTS[1]
    for key in ["name", "code", "purpose", "human_owner"]:
        docx_para(doc, f"{key.replace('_',' ').title()}: {a[key]}")
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Templates/Agent_Charter_Template.docx")


def gen_raci_template():
    doc = Document()
    doc.add_heading("RACI Matrix Template", 0)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Activity", "AP Team", "Requester", "Procurement", "Approver", "Agent"]):
        hdr[i].text = h
    activities = ["Invoice receipt", "Data extraction", "PO matching", "Exception resolution", "Approval routing", "Payment proposal review", "Month-end close support"]
    for act in activities:
        row = table.add_row().cells
        row[0].text = act
        for i in range(1, 6):
            row[i].text = "[R/A/C/I]"
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Templates/RACI_Template.docx")


def gen_governance_standard():
    doc = Document()
    doc.add_heading("AP Agent Governance Standard", 0)
    doc.add_paragraph(BRAND["tagline"])
    principles = [
        "Human accountability is never transferred to an agent",
        "Agents earn responsibility through demonstrated performance",
        "Segregation of duties must be maintained in all agent workflows",
        "All agent decisions must be auditable with retained evidence",
        "Autonomy progression requires formal approval",
        "Payment authorisation remains human-controlled at all times",
        "Data access follows least-privilege principles",
        "Agent outputs must be validated before downstream action",
        "Incident response procedures must be defined before deployment",
        "Periodic certification of agent performance is mandatory",
    ]
    docx_heading(doc, "Governance Principles", 1)
    for p in principles:
        doc.add_paragraph(p, style="List Bullet")
    docx_heading(doc, "Autonomy Progression Requirements", 1)
    for level in RESPONSIBILITY_LEVELS:
        docx_para(doc, f"Level {level['level']} — {level['name']}: {level['description']}", bold=True)
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Governance/AP_Agent_Governance_Standard.docx")


def gen_uat_template():
    doc = Document()
    doc.add_heading("User Acceptance Testing Template", 0)
    doc.add_paragraph("Agent: [Agent Name] | Version: [X.X] | Date: [YYYY-MM-DD]")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for i, h in enumerate(["Test ID", "Scenario", "Input", "Expected Output", "Actual Output", "Pass/Fail", "Tester"]):
        table.rows[0].cells[i].text = h
    tests = [
        ("UAT-001", "Valid invoice — full match", "Standard PO invoice within tolerance", "Auto-matched, routed to approval", "", "", ""),
        ("UAT-002", "Price variance within tolerance", "Invoice $0.50 above PO price", "Matched with variance noted", "", "", ""),
        ("UAT-003", "Price variance above tolerance", "Invoice 15% above PO price", "Exception EXC-005 created", "", "", ""),
        ("UAT-004", "Missing PO", "Invoice without PO reference", "Exception EXC-001, routed to requester", "", "", ""),
        ("UAT-005", "Duplicate invoice", "Same invoice number resubmitted", "Exception EXC-009 flagged", "", "", ""),
        ("UAT-006", "Unreadable document", "Poor scan quality PDF", "Exception EXC-019, manual review queue", "", "", ""),
        ("UAT-007", "Multi-line partial match", "3 lines, 1 unmatched", "2 matched, 1 exception per line", "", "", ""),
        ("UAT-008", "Shadow mode validation", "Agent recommends but does not act", "Recommendation logged, no system change", "", "", ""),
    ]
    for t in tests:
        row = table.add_row().cells
        for i, v in enumerate(t):
            row[i].text = v
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Testing/UAT_Template.docx")


def gen_risk_assessment():
    doc = Document()
    doc.add_heading("AP Agent Risk Assessment", 0)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for i, h in enumerate(["Risk ID", "Risk Description", "Likelihood", "Impact", "Mitigation", "Owner"]):
        table.rows[0].cells[i].text = h
    risks = [
        ("R-001", "Duplicate payment due to agent error", "Low", "Critical", "Human review gate + duplicate detection", "AP Controls"),
        ("R-002", "Unauthorised payment execution", "Low", "Critical", "Agent cannot execute payments", "Treasury"),
        ("R-003", "Prompt injection via invoice data", "Medium", "High", "Input sanitisation + output validation", "IT Security"),
        ("R-004", "Hallucinated classification", "Medium", "Medium", "Confidence thresholds + human review", "AP Manager"),
        ("R-005", "Data privacy breach", "Low", "Critical", "Least privilege + PII masking", "DPO"),
        ("R-006", "Over-automation without governance", "Medium", "High", "5-stage autonomy model", "AP Director"),
        ("R-007", "Supplier impersonation", "Low", "Critical", "Bank change verification process", "AP/Treasury"),
        ("R-008", "Model drift degrading accuracy", "Medium", "Medium", "Periodic revalidation + monitoring", "Automation Lead"),
    ]
    for r in risks:
        row = table.add_row().cells
        for i, v in enumerate(r):
            row[i].text = v
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Governance/AP_Agent_Risk_Register.docx")


def gen_process_discovery():
    doc = Document()
    doc.add_heading("Process Discovery Guide", 0)
    doc.add_paragraph("Evidence Room 10-Step Process-to-Agent Methodology")
    steps = [
        ("STEP 1 — Observe", "Record real process walkthroughs with AP staff. Capture screen recordings, shadow sessions, and timing data."),
        ("STEP 2 — Transcribe", "Generate transcripts of process discussions. Use structured interview guides."),
        ("STEP 3 — Extract", "Identify steps, systems, decisions, business rules, inputs, outputs, exceptions, controls, dependencies."),
        ("STEP 4 — Structure", "Convert into process map, decision tree, exception taxonomy, control map, RACI, SOP."),
        ("STEP 5 — Agentise", "Determine what remains human, what AI recommends, prepares, executes, and where deterministic automation is preferable."),
        ("STEP 6 — Test", "Run historical cases against agent logic. Measure accuracy against human baseline."),
        ("STEP 7 — Shadow", "Agent operates without action permissions. Compare recommendations to actual outcomes."),
        ("STEP 8 — Controlled Pilot", "Limited users, transactions, or categories. Daily monitoring."),
        ("STEP 9 — Measure", "Compare against baseline KPIs. Audit sample of decisions."),
        ("STEP 10 — Expand Responsibility", "Only after evidence demonstrates adequate performance. Formal autonomy review."),
    ]
    for title, desc in steps:
        docx_heading(doc, title, 2)
        docx_para(doc, desc)
    save_docx(doc, ROOT / "03_AP_AGENT_OS_PRO/Process_Mapping/Process_Discovery_Guide.docx")


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ERTitle", fontSize=24, textColor=ACCENT, spaceAfter=20, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle(name="ERHeading", fontSize=14, textColor=ACCENT, spaceBefore=16, spaceAfter=8, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle(name="ERBody", fontSize=10, leading=14, spaceAfter=6))
    styles.add(ParagraphStyle(name="ERSmall", fontSize=8, textColor=colors.grey, leading=10))
    return styles


def gen_pdf(path, title, sections):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, topMargin=2*cm, bottomMargin=2*cm, leftMargin=2.5*cm, rightMargin=2.5*cm)
    styles = pdf_styles()
    story = [Paragraph(title, styles["ERTitle"]), Spacer(1, 12)]
    for heading, content in sections:
        story.append(Paragraph(heading, styles["ERHeading"]))
        if isinstance(content, list):
            for item in content:
                story.append(Paragraph(f"• {item}", styles["ERBody"]))
        else:
            story.append(Paragraph(content, styles["ERBody"]))
        story.append(Spacer(1, 8))
    doc.build(story)
    print(f"  Created {path.relative_to(ROOT.parent)}")


def gen_free_diagnostic_pdf():
    sections = [
        ("About This Diagnostic", "The Evidence Room AP AI Readiness Diagnostic helps Finance leaders assess organisational readiness for deploying governed AI agents across Accounts Payable. Complete the 37-question assessment in the accompanying spreadsheet to generate your readiness score."),
        ("Maturity Model", [
            "Level 1 — Ad Hoc (0-25%): Manual processes, no automation strategy",
            "Level 2 — Emerging (26-45%): Some automation, early AI interest",
            "Level 3 — Defined (46-65%): Documented processes, pilot-ready",
            "Level 4 — Managed (66-85%): Systematic automation, active pilots",
            "Level 5 — Optimised (86-100%): Agent workforce with measured autonomy",
        ]),
        ("What You Receive", [
            "AP AI readiness score with maturity classification",
            "Opportunity heatmap across 10 core agent categories",
            "Baseline KPI worksheet for measurement",
            "Business case starter with industry benchmarks",
            "Downloadable scorecard for executive reporting",
        ]),
        ("Industry Context", "Ardent Partners research shows the average organisation spends $9.40 per invoice (all-inclusive), while best-in-class teams achieve $2.78-$2.81. Straight-through processing averages 32.6% across organisations. These benchmarks provide context for your readiness assessment — not targets."),
        ("Next Steps", "Score below 45%: Focus on process documentation and data access. Score 46-65%: Select your first bounded agent use case. Score 66%+: Begin shadow-mode testing with the AP Agent OS Professional toolkit."),
    ]
    gen_pdf(ROOT / "01_FREE_AP_AI_READINESS/AP_AI_Readiness_Diagnostic.pdf", "Evidence Room\nAP AI Readiness Diagnostic", sections)


def gen_starter_guide_pdf():
    sections = [
        ("Overview", "The AP Agent Starter Kit provides the essential frameworks to begin designing your AP agent operating model. This kit includes the top 10 agent blueprints, exception taxonomy, governance checklist, and implementation roadmap."),
        ("The Evidence Room Principle", "AI agents should be treated like employees who earn responsibility through demonstrated performance. Never position full autonomy as the default. Start at Level 0 (Observe) and progress through evidence-based stages."),
        ("Top 10 Agents to Start", [f"{a['code']} {a['name']}: {a['purpose'][:100]}..." for a in AGENTS[:10]]),
        ("Human vs Agent Decision Framework", "Retain human control for: payment authorisation, master data changes, fraud determinations, policy exceptions, and supplier commitments. Delegate to agents: classification, matching within tolerance, follow-up drafting, reporting, and anomaly detection."),
        ("Implementation Checklist", [
            "Complete AP AI Readiness Diagnostic",
            "Secure executive sponsorship",
            "Document current-state process for first agent",
            "Establish baseline KPIs",
            "Write agent charter for pilot agent",
            "Define controls and approval gates",
            "Run historical test cases",
            "Deploy in shadow mode for 3 weeks",
            "Review performance and decide on autonomy progression",
        ]),
    ]
    gen_pdf(ROOT / "02_AP_AGENT_STARTER/AP_Agent_Starter_Guide.pdf", "Evidence Room\nAP Agent Starter Kit", sections)


def gen_professional_os_pdf():
    agent_summaries = [f"<b>{a['code']} {a['name']}</b> — Autonomy L{a['default_autonomy']}: {a['purpose']}" for a in AGENTS]
    sections = [
        ("The AP Agent Operating System", f"Evidence Room AP Agent OS is the implementation-ready toolkit for building, governing, and scaling AI agents across Accounts Payable. It contains specifications for all {len(AGENTS)} agents, complete governance frameworks, KPI definitions, testing methodologies, and business case models."),
        ("16-Agent Architecture", agent_summaries[:8]),
        ("Agents 9-16", agent_summaries[8:]),
        ("5-Stage Responsibility Model", [f"Level {l['level']} — {l['name']}: {l['description']}" for l in RESPONSIBILITY_LEVELS]),
        ("Governance Framework", "Every agent requires: job description, inputs, tools, responsibilities, exclusions, human owner, approval requirements, escalation criteria, output standard, control requirements, audit evidence, KPIs, performance history, autonomy level, failure handling, and cost monitoring."),
        ("Exception Taxonomy", f"{len(EXCEPTIONS)} exception categories defined with root cause, resolution path, responsible party, assigned agent, automation potential, and risk level. See accompanying Exception Tracker spreadsheet."),
        ("Implementation Roadmap", "10-phase deployment framework from baseline through scale. Accelerated timeline: 4-6 weeks for one well-bounded agent. Actual duration depends on systems, controls, integrations, data quality, and governance maturity."),
        ("Evidence Standard", "Industry benchmarks sourced from Ardent Partners, Forrester, PwC, and Deloitte. See Research Ledger for full citations. Evidence Room frameworks are independently authored and clearly distinguished from third-party data."),
    ]
    gen_pdf(ROOT / "03_AP_AGENT_OS_PRO/AP_Agent_OS_Professional.pdf", "Evidence Room\nAP Agent OS — Professional", sections)


def gen_team_playbook_pdf():
    sections = [
        ("Team Edition Overview", "The AP Agent OS Team Edition is designed for enterprise transformation initiatives. It includes facilitation materials, workshop decks, stakeholder interview guides, change management toolkit, and executive communication templates."),
        ("Workshop Programme", "A structured 2-day workshop covering: current-state assessment, agent architecture overview, first agent selection, governance design, KPI establishment, and 90-day implementation planning."),
        ("Stakeholder Engagement", "Interview guides for CFO, AP Manager, IT, Internal Audit, Procurement, and Shared Services leaders. Process-owner questionnaires for detailed discovery."),
        ("Change Management", "Training materials for AP team, executive communication templates, benefits realisation tracker, and steering committee reporting pack."),
        ("Enterprise Governance", "Extended governance pack including risk register, control matrix, periodic certification checklist, and incident response procedures."),
    ]
    gen_pdf(ROOT / "04_AP_AGENT_OS_TEAM/AP_Agent_OS_Team_Playbook.pdf", "Evidence Room\nAP Agent OS — Team Edition", sections)


def gen_custom_brochure_pdf():
    sections = [
        ("AP Transformation Blueprint", "A productised service delivering a bespoke AP agent transformation plan. You provide structured information; Evidence Room returns a complete implementation blueprint."),
        ("What You Receive", [
            "Current-state assessment with AP AI maturity score",
            "Opportunity map across 16 agent categories",
            "Recommended agent architecture prioritised by impact",
            "Operating model design",
            "KPI baseline framework",
            "Controls framework",
            "Business case with ROI scenarios",
            "90-day implementation plan",
            "Executive presentation (board-ready)",
        ]),
        ("How It Works", "1. Complete the structured intake questionnaire. 2. Evidence Room analyses your inputs using the AP Agent OS framework. 3. Receive your bespoke blueprint within 10 business days. 4. Optional: implementation support consultation."),
        ("Investment", "US$1,500–3,000 depending on organisation complexity. Includes one round of revisions."),
    ]
    gen_pdf(ROOT / "05_CUSTOM_BLUEPRINT/AP_Transformation_Blueprint_Brochure.pdf", "Evidence Room\nAP Transformation Blueprint", sections)


def generate_all():
    print("Generating Word templates...")
    gen_sop_template()
    gen_agent_charter()
    gen_raci_template()
    gen_governance_standard()
    gen_uat_template()
    gen_risk_assessment()
    gen_process_discovery()
    print("Generating PDFs...")
    gen_free_diagnostic_pdf()
    gen_starter_guide_pdf()
    gen_professional_os_pdf()
    gen_team_playbook_pdf()
    gen_custom_brochure_pdf()
    print("Document generation complete.")


if __name__ == "__main__":
    generate_all()
