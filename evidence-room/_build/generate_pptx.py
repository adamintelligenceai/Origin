#!/usr/bin/env python3
"""Generate Evidence Room PowerPoint decks."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "pptx"
OUT.mkdir(parents=True, exist_ok=True)
WS_OUT = ROOT / "04_AP_AGENT_OS_TEAM" / "Workshop" / "pptx"
WS_OUT.mkdir(parents=True, exist_ok=True)

INK = RGBColor(0x0B, 0x1F, 0x33)
TEAL = RGBColor(0x1F, 0x6F, 0x78)
PAPER = RGBColor(0xF7, 0xF4, 0xEF)
AMBER = RGBColor(0xC4, 0x7E, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_slide_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return tf


def add_bullets(slide, left, top, width, height, items, *, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.font.name = "Calibri"
        p.space_after = Pt(8)
    return tf


def title_slide(prs, eyebrow, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, INK)
    add_textbox(slide, 0.7, 0.5, 8, 0.4, "EVIDENCE ROOM", size=14, bold=True, color=TEAL)
    add_textbox(slide, 0.7, 1.8, 8.5, 1.5, title, size=36, bold=True, color=WHITE)
    add_textbox(slide, 0.7, 3.5, 8.5, 1.2, subtitle, size=18, color=PAPER)
    add_textbox(slide, 0.7, 6.6, 8.5, 0.4, eyebrow, size=12, color=AMBER)
    return slide


def section_slide(prs, title, bullets, footer="Evidence Room — AP Agent OS"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_textbox(slide, 0.7, 0.35, 8, 0.3, "EVIDENCE ROOM", size=12, bold=True, color=TEAL)
    add_textbox(slide, 0.7, 0.7, 8.5, 1.0, title, size=28, bold=True, color=INK)
    add_bullets(slide, 0.7, 1.9, 8.5, 4.5, bullets, size=16)
    add_textbox(slide, 0.7, 6.8, 8.5, 0.3, footer, size=10, color=TEAL)
    return slide


def build_cfo_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Use widescreen dimensions carefully with positions
    def ts(eyebrow, title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide, INK)
        add_textbox(slide, 0.8, 0.6, 11, 0.4, "EVIDENCE ROOM", size=14, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 2.0, 11, 1.5, title, size=40, bold=True, color=WHITE)
        add_textbox(slide, 0.8, 3.8, 11, 1.2, subtitle, size=20, color=PAPER)
        add_textbox(slide, 0.8, 6.7, 11, 0.4, eyebrow, size=12, color=AMBER)

    def ss(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide)
        add_textbox(slide, 0.8, 0.4, 11, 0.3, "EVIDENCE ROOM", size=12, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 0.85, 11, 0.8, title, size=28, bold=True, color=INK)
        add_bullets(slide, 0.8, 2.0, 11, 4.5, bullets, size=18)
        add_textbox(slide, 0.8, 6.9, 11, 0.3, "AP Agent OS · Confidential working draft", size=10, color=TEAL)

    ts(
        "CFO / Finance Leadership briefing",
        "AP Agent OS — from experiment to operating performance",
        "Governed agents. Measurable outcomes. Evidence over hype.",
    )
    ss(
        "The operating problem",
        [
            "AP teams face material exception work and rising AI pressure simultaneously.",
            "Industry context (Ardent Partners 2025): avg cost/invoice $9.84; BIC $2.65; exception 18.4%; STP 35.4%; AI adoption 44%.",
            "Software stacks automate capture and workflow; few organisations have an agent operating model.",
            "Question for leadership: who owns agent decisions, evidence, and autonomy expansion?",
        ],
    )
    ss(
        "What Evidence Room is — and is not",
        [
            "Is: an operating system toolkit for designing, governing, and measuring AP AI agents.",
            "Is not: an ERP, AP suite, payment network, or prompt pack.",
            "Works across Dynamics, SAP, Oracle, NetSuite, Workday, and other environments.",
            "Payment authorisation remains human-controlled.",
        ],
    )
    ss(
        "The 16-agent AP workforce (summary)",
        [
            "Intake → Validation → Matching → Exception Triage",
            "Goods Receipt → PO Quality → Approvals → Supplier & Internal follow-up",
            "Duplicate/Anomaly indicators → Statement recon → Payment proposal review",
            "Close → Reporting → Root cause → Orchestrator supervision",
        ],
    )
    ss(
        "Responsibility is earned (Levels 0–4)",
        [
            "L0 Observe · L1 Recommend · L2 Prepare · L3 Execute within guardrails · L4 Managed autonomy",
            "Never default to full autonomy.",
            "Promotion requires KPI evidence, control sign-off, and dual approval.",
            "Demotion is mandatory after material control failure.",
        ],
    )
    ss(
        "Governance & control",
        [
            "Human accountability and segregation of duties retained.",
            "Agent Control Matrix: risk, control, owner, evidence, frequency, escalation.",
            "Prompt-injection, hallucination, and output-validation controls required.",
            "Audit logs, version control, and incident response are first-class artefacts.",
        ],
    )
    ss(
        "Illustrative economics (not a guarantee)",
        [
            "Use organisation actuals in the ROI calculator; defaults include Ardent context benchmarks.",
            "Model Conservative / Base / Upside — avoid false precision.",
            "Track capacity released separately from validated cash savings.",
            "Only book benefits with Finance-validated evidence packs.",
        ],
    )
    ss(
        "Recommended 90-day move",
        [
            "Weeks 1–2: readiness diagnostic + baseline KPIs",
            "Weeks 3–6: discover process; charter first agent; shadow mode",
            "Weeks 7–10: controlled pilot on bounded scope",
            "Weeks 11–12: performance review; decide earned responsibility progression",
        ],
    )
    ss(
        "Ask of the committee",
        [
            "Approve AP Agent OS as the governance frame for Finance AI experiments.",
            "Nominate human owners for first agent and controls.",
            "Fund Team Edition facilitation ($499 toolkit) or Custom Blueprint ($1,500–$3,000) as appropriate.",
            "Require evidence packs before any autonomy expansion.",
        ],
    )
    path = OUT / "ER_CFO_AP_Transformation_Deck.pptx"
    prs.save(path)
    return path


def build_workshop_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def ts(title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide, INK)
        add_textbox(slide, 0.8, 0.6, 11, 0.4, "EVIDENCE ROOM · WORKSHOP", size=14, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 2.2, 11, 1.5, title, size=36, bold=True, color=WHITE)
        add_textbox(slide, 0.8, 4.0, 11, 1.2, subtitle, size=18, color=PAPER)

    def ss(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide)
        add_textbox(slide, 0.8, 0.4, 11, 0.3, "WORKSHOP", size=12, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 0.85, 11, 0.8, title, size=26, bold=True, color=INK)
        add_bullets(slide, 0.8, 1.9, 11, 4.6, bullets, size=17)

    ts("AP Agent OS Workshop", "Design the agent layer across the stack you already own")
    ss("Outcomes for today", [
        "Shared language for AP agents vs automation vs ERP features",
        "Draft exception taxonomy for your environment",
        "Shortlist of first two agent candidates",
        "Named human owners and control boundaries",
        "30-day action list with evidence requirements",
    ])
    ss("Exercise 1 — Process truth", [
        "Map one invoice happy path on sticky notes / board",
        "Add three recent exceptions with who touched them",
        "Mark systems, decisions, and controls on each step",
        "Facilitator captures into Process Discovery template",
    ])
    ss("Exercise 2 — Human vs agent", [
        "For each step: Human only / Agent recommend / Agent prepare / Deterministic automation",
        "No step may be 'full autonomy' without an evidence path",
        "Payment authorisation stays human",
        "Record disagreements — they are design gold",
    ])
    ss("Exercise 3 — First agent charter", [
        "Pick one bounded use case (e.g., missing GR follow-up)",
        "Fill job description, exclusions, KPIs, escalation",
        "Set starting autonomy to Level 0 or 1",
        "Identify data access and least-privilege needs",
    ])
    ss("Exercise 4 — Control & measure", [
        "Select 5 controls from the Control Matrix pattern",
        "Select 6 KPIs across operational + risk categories",
        "Define shadow-mode success criteria",
        "Schedule performance review date before any promotion",
    ])
    ss("Close — commitments", [
        "Owners confirm 30-day actions",
        "Steering date booked",
        "No production agent actions without charter + controls",
        "Evidence over hype",
    ])
    path = WS_OUT / "ER_AP_Workshop_Deck.pptx"
    prs.save(path)
    return path


def build_steering_pack():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def ss(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide)
        add_textbox(slide, 0.8, 0.4, 11, 0.3, "STEERING COMMITTEE", size=12, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 0.85, 11, 0.8, title, size=26, bold=True, color=INK)
        add_bullets(slide, 0.8, 1.9, 11, 4.6, bullets, size=17)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, INK)
    add_textbox(slide, 0.8, 2.2, 11, 1.2, "AP Agent OS — Steering Pack", size=36, bold=True, color=WHITE)
    add_textbox(slide, 0.8, 3.6, 11, 0.8, "Decision, risk, and evidence pack for executive oversight", size=18, color=PAPER)

    ss("Standard agenda", [
        "1. KPI snapshot (STP, exceptions, intervention, cost, incidents)",
        "2. Agent registry changes and autonomy requests",
        "3. Control exceptions and incidents",
        "4. Benefits validation status",
        "5. Decisions required",
    ])
    ss("Decision types", [
        "Approve / defer / reject autonomy promotion",
        "Expand or freeze scope",
        "Fund next workstream",
        "Escalate vendor/model risk",
        "Require remediation before continue",
    ])
    ss("Non-negotiables", [
        "No payment execution by agents",
        "No silent level changes",
        "No unverified savings in board materials",
        "Human accountable owner for every live agent",
    ])
    path = OUT / "ER_Steering_Committee_Pack.pptx"
    prs.save(path)
    return path


def build_business_case_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def ss(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide)
        add_textbox(slide, 0.8, 0.4, 11, 0.3, "BUSINESS CASE", size=12, bold=True, color=TEAL)
        add_textbox(slide, 0.8, 0.85, 11, 0.8, title, size=26, bold=True, color=INK)
        add_bullets(slide, 0.8, 1.9, 11, 4.6, bullets, size=17)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, INK)
    add_textbox(slide, 0.8, 2.2, 11, 1.2, "AP Agent Business Case", size=36, bold=True, color=WHITE)
    add_textbox(slide, 0.8, 3.6, 11, 0.8, "Conservative · Base · Upside — illustrative, not guaranteed", size=18, color=PAPER)

    ss("Input stack", [
        "Invoice volume, FTE, loaded cost, manual-touch %, exception rate",
        "Resolution minutes, duplicate rate, late-payment / discount opportunity",
        "AI/tool cost, implementation cost, efficiency ranges",
        "Replace defaults with your actuals before steering",
    ])
    ss("Output stack", [
        "Baseline cost and capacity",
        "Scenario benefits and net benefit after tool cost",
        "Payback and Year-1 ROI ranges (illustrative)",
        "Sensitivity to efficiency and AI cost assumptions",
    ])
    ss("How to present without hype", [
        "Show ranges, not single-point miracles",
        "Separate capacity released from cash validated",
        "Cite external benchmarks as context only (Ardent 2025 $9.84 / $2.65)",
        "Commit to measurement before claiming success",
    ])
    path = OUT / "ER_Business_Case_Deck.pptx"
    prs.save(path)
    return path


def main():
    paths = [
        build_cfo_deck(),
        build_workshop_deck(),
        build_steering_pack(),
        build_business_case_deck(),
    ]
    for p in paths:
        print(f"Wrote {p}")


if __name__ == "__main__":
    main()
