#!/usr/bin/env python3
"""Generate Evidence Room Word templates."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "docx"
OUT.mkdir(parents=True, exist_ok=True)
LEGAL = ROOT / "10_LEGAL_AND_LICENSING" / "docx"
LEGAL.mkdir(parents=True, exist_ok=True)

INK = RGBColor(0x0B, 0x1F, 0x33)
TEAL = RGBColor(0x1F, 0x6F, 0x78)


def set_run(run, *, bold=False, size=11, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, bold=True, size=16 if level == 1 else 13, color=INK)
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic, size=size)
    return p


def cover(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("EVIDENCE ROOM")
    set_run(r, bold=True, size=12, color=TEAL)
    add_heading(doc, title, 1)
    add_para(doc, subtitle, italic=True, size=11)
    add_para(doc, "AP Agent OS · Version 1.0.0 · 2026-03-20", size=10)
    add_para(
        doc,
        "AI that earns responsibility. Governed agents. Measurable outcomes.",
        italic=True,
        size=10,
    )
    doc.add_paragraph()


def field_table(doc, labels):
    table = doc.add_table(rows=len(labels), cols=2)
    table.style = "Table Grid"
    for i, label in enumerate(labels):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = ""
    doc.add_paragraph()


def write_agent_charter():
    doc = Document()
    cover(doc, "Agent Charter Template", "Blank reusable + worked example structure")
    add_heading(doc, "1. Blank charter", 2)
    field_table(
        doc,
        [
            "Agent name / ID",
            "Human owner",
            "Deputy owner",
            "Starting autonomy level (0–4)",
            "Systems & data required",
            "Primary KPIs",
            "Approval requirements",
            "Escalation path",
            "Review cadence",
            "Cost monitoring owner",
        ],
    )
    for section in [
        ("Job description", "One paragraph describing the agent’s role in the AP lifecycle."),
        ("Inputs", "List required inputs and quality expectations."),
        ("Responsibilities", "Bullet the actions the agent may perform at the approved level."),
        ("Explicit exclusions", "State what the agent must never do (e.g., authorise payment)."),
        ("Output standard", "Define acceptable output format, confidence labelling, and evidence."),
        ("Control requirements", "Link to control matrix IDs."),
        ("Audit evidence", "Specify logs, samples, and retention."),
        ("Failure handling", "Describe fallback and human takeover."),
        ("Promotion / demotion criteria", "Evidence required to change autonomy level."),
    ]:
        add_heading(doc, section[0], 3)
        add_para(doc, section[1])
        add_para(doc, "[Write here]", italic=True)

    add_heading(doc, "2. Worked example (illustrative — ACME Corp)", 2)
    add_para(
        doc,
        "Agent: AP-04 Exception Triage · Owner: Head of AP · Starting level: 0 (Observe) → target pilot level 1 (Recommend). "
        "Exclusions: no payment changes, no taxonomy edits without change control. "
        "Primary KPIs: classification accuracy ≥95%, false-positive rate ≤5%, avg triage time reduction vs baseline. "
        "This example is fictional and for training use.",
    )
    path = OUT / "ER_Agent_Charter_Template.docx"
    doc.save(path)
    return path


def write_sop():
    doc = Document()
    cover(doc, "SOP Template — AP Agent Supported Process", "Standard operating procedure with human and agent lanes")
    field_table(
        doc,
        [
            "SOP ID / Version",
            "Process name",
            "Process owner",
            "Effective date",
            "Related agent(s)",
            "Related controls",
        ],
    )
    for h, body in [
        ("Purpose", "State why the procedure exists and which AP risk or outcome it addresses."),
        ("Scope", "Inclusions, exclusions, entities, systems."),
        ("Definitions", "Exception codes, autonomy levels, STP, etc."),
        ("Roles (RACI)", "Responsible, Accountable, Consulted, Informed — human and agent."),
        ("Procedure — Human steps", "Numbered steps humans perform."),
        ("Procedure — Agent steps", "Numbered steps agent may perform at approved autonomy."),
        ("Decision points", "Where judgment remains human-only."),
        ("Exceptions & escalations", "Link taxonomy codes and SLA."),
        ("Evidence & records", "What is retained for audit."),
        ("KPIs", "How performance of this SOP is measured."),
        ("Revision history", "Version, date, author, change summary."),
    ]:
        add_heading(doc, h, 2)
        add_para(doc, body)
        add_para(doc, "[Content]", italic=True)
    path = OUT / "ER_SOP_Template.docx"
    doc.save(path)
    return path


def write_process_discovery():
    doc = Document()
    cover(doc, "Process Discovery Workbook", "Interview and observation capture for AP → agent design")
    add_heading(doc, "Session details", 2)
    field_table(
        doc,
        [
            "Date",
            "Facilitator",
            "Interviewee / role",
            "Process segment",
            "Systems observed",
            "Recording / transcript location",
        ],
    )
    add_heading(doc, "Capture grid", 2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for i, h in enumerate(["Step", "System", "Decision / rule", "Input", "Output", "Exception / control"]):
        table.rows[0].cells[i].text = h
    for _ in range(12):
        row = table.add_row().cells
        for c in row:
            c.text = ""
    doc.add_paragraph()
    add_heading(doc, "Synthesis prompts", 2)
    for q in [
        "What remains human-only judgment?",
        "What can an agent recommend safely?",
        "What can an agent prepare for approval?",
        "Where is deterministic automation better than an LLM agent?",
        "What evidence proves the agent is ready to earn more responsibility?",
    ]:
        add_para(doc, "• " + q)
    path = OUT / "ER_Process_Discovery_Template.docx"
    doc.save(path)
    return path


def write_raci():
    doc = Document()
    cover(doc, "RACI Template — AP Agent Programme", "Human accountability remains mandatory")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = ["Activity", "AP Manager", "Controller", "AI/Automation Lead", "Process Owner", "Agent (supporting)"]
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    activities = [
        "Approve agent charter",
        "Grant system access",
        "Approve autonomy promotion",
        "Review weekly KPI pack",
        "Handle payment authorisation",
        "Incident response for AI error",
        "Supplier communication send",
        "Month-end completeness sign-off",
    ]
    for a in activities:
        cells = table.add_row().cells
        cells[0].text = a
        for i in range(1, 6):
            cells[i].text = ""
    doc.add_paragraph()
    add_para(doc, "Legend: R = Responsible, A = Accountable, C = Consulted, I = Informed. Agents are never Accountable.", italic=True)
    path = OUT / "ER_RACI_Template.docx"
    doc.save(path)
    return path


def write_governance_standard():
    doc = Document()
    cover(doc, "AP Agent Governance Standard", "Editable organisational standard (customise before adoption)")
    for h, body in [
        ("1. Purpose", "Establish mandatory rules for designing, deploying, and supervising AI agents in Accounts Payable."),
        ("2. Principles", "Human accountability; least privilege; earned responsibility; evidence before autonomy; payment authorisation remains human."),
        ("3. Roles", "Define AP Owner, Control Owner, AI Product Owner, Model/Vendor Risk Owner."),
        ("4. Lifecycle", "Observe → Recommend → Prepare → Execute within guardrails → Managed autonomy. Default start at Level 0–1."),
        ("5. Change control", "Model, prompt, workflow, and taxonomy changes require versioning, testing, and approval."),
        ("6. Monitoring", "KPI scorecard, control matrix evidence, incident log, cost monitoring."),
        ("7. Exceptions", "Override and fallback procedures with audit trail."),
        ("8. Periodic certification", "Quarterly access and control certification."),
        ("9. Related documents", "Control matrix, risk register, agent registry, licence terms."),
    ]:
        add_heading(doc, h, 2)
        add_para(doc, body)
        add_para(doc, "[Organisation-specific policy text]", italic=True)
    path = OUT / "ER_Governance_Standard.docx"
    doc.save(path)
    return path


def write_uat():
    doc = Document()
    cover(doc, "UAT Template — AP Agent", "User acceptance testing for shadow and pilot phases")
    field_table(
        doc,
        ["Agent", "Build / version", "Tester", "Environment", "Date", "Pass criteria reference"],
    )
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for i, h in enumerate(["Case ID", "Scenario", "Expected result", "Actual result", "Pass/Fail", "Evidence"]):
        table.rows[0].cells[i].text = h
    for i in range(1, 11):
        cells = table.add_row().cells
        cells[0].text = f"UAT-{i:02d}"
    doc.add_paragraph()
    add_para(doc, "Sign-off: Tester ________  Owner ________  Date ________", bold=True)
    path = OUT / "ER_UAT_Template.docx"
    doc.save(path)
    return path


def write_risk_assessment():
    doc = Document()
    cover(doc, "Risk Assessment Template — AP Agent Deployment", "Assess before pilot and before each autonomy promotion")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for i, h in enumerate(["Risk", "Cause", "Impact", "Likelihood", "Score", "Mitigation", "Owner"]):
        table.rows[0].cells[i].text = h
    for risk in [
        "Incorrect exception classification",
        "Hallucinated supplier communication",
        "Unauthorised data exposure",
        "Prompt injection via invoice content",
        "Silent autonomy expansion",
        "Payment control bypass attempt",
        "Model drift after vendor update",
    ]:
        cells = table.add_row().cells
        cells[0].text = risk
    path = OUT / "ER_Risk_Assessment_Template.docx"
    doc.save(path)
    return path


def write_meeting_guide():
    doc = Document()
    cover(doc, "Meeting Guide — AP Process Walkthrough", "60–90 minute discovery session")
    add_heading(doc, "Agenda", 2)
    for item in [
        "0–5 min: Purpose and consent for notes/transcript",
        "5–20 min: Walk a recent happy-path invoice",
        "20–45 min: Walk three recent exceptions",
        "45–60 min: Controls, approvals, pain points",
        "60–75 min: Candidate agent boundaries",
        "75–90 min: Next steps and evidence needed",
    ]:
        add_para(doc, "• " + item)
    add_heading(doc, "Must-ask questions", 2)
    for q in [
        "What must never be automated without a human?",
        "Where do invoices age and why?",
        "Which exceptions repeat every week?",
        "What evidence would convince you an agent earned more responsibility?",
    ]:
        add_para(doc, "• " + q)
    path = OUT / "ER_Meeting_Guide.docx"
    doc.save(path)
    return path


def write_implementation_plan():
    doc = Document()
    cover(doc, "Implementation Plan Template", "90-day plan scaffold for one bounded AP agent")
    add_para(
        doc,
        "Note: A well-bounded agent may be achievable in approximately 4–6 weeks in favourable conditions. "
        "Actual duration depends on systems, controls, integrations, process complexity, data quality, and governance.",
        italic=True,
    )
    for h in ["Objectives", "Scope", "Out of scope", "Workstreams", "Milestones", "Risks", "RACI", "Budget", "Success metrics", "Go/no-go gates"]:
        add_heading(doc, h, 2)
        add_para(doc, "[Content]", italic=True)
    path = OUT / "ER_Implementation_Plan_Template.docx"
    doc.save(path)
    return path


def write_licence():
    doc = Document()
    cover(doc, "Licence Terms (Plain English Draft)", "Flag for professional legal review before publication")
    add_para(
        doc,
        "This document is a plain-English draft for product packaging. It is not legal advice. "
        "Have qualified counsel adapt it for your jurisdiction before use.",
        bold=True,
        size=10,
    )
    for h, body in [
        ("Individual", "One named user. Personal learning and individual professional use within the licence terms."),
        ("Professional", "One named user for commercial/professional work for a single employer or client engagement, without redistribution."),
        ("Team", "Internal use within one defined company/team subject to stated seat or entity limits in the order."),
        ("Prohibited", "No resale, redistribution, sublicensing, publishing templates as your own product, or using materials to train a competing commercial product where enforceable."),
        ("IP", "Evidence Room retains ownership of the materials. Purchasers receive a limited licence to use, not ownership of the IP."),
        ("Disclaimer", "No guarantee of savings, ROI, fraud detection, regulatory compliance, accounting accuracy, or autonomous payment safety."),
    ]:
        add_heading(doc, h, 2)
        add_para(doc, body)
    path = LEGAL / "ER_Licence_Terms_Draft.docx"
    doc.save(path)
    return path


def main():
    paths = [
        write_agent_charter(),
        write_sop(),
        write_process_discovery(),
        write_raci(),
        write_governance_standard(),
        write_uat(),
        write_risk_assessment(),
        write_meeting_guide(),
        write_implementation_plan(),
        write_licence(),
    ]
    for p in paths:
        print(f"Wrote {p}")


if __name__ == "__main__":
    main()
