#!/usr/bin/env python3
"""Editable Word templates: example + blank in each file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
INK = RGBColor(0x14, 0x17, 0x1C)
FOREST = RGBColor(0x1E, 0x5C, 0x45)
SLATE = RGBColor(0x5C, 0x63, 0x70)


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 13, True, FOREST if level > 1 else INK)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_p(doc, text, *, italic=False, size=11, color=INK):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size, False, color, italic)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_meta(doc, title, subtitle):
    p = doc.add_paragraph()
    run = p.add_run("EVIDENCE ROOM")
    set_run_font(run, 10, True, FOREST)
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_run_font(run, 22, True, INK)
    add_p(doc, subtitle, italic=True, size=10, color=SLATE)
    add_p(doc, "Version 1.0 · September 2026 · Licensed material · Flag legal review before customer-facing use as a contract.", italic=True, size=9, color=SLATE)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, 10, True, INK)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            t.rows[r].cells[c].text = ""
            run = t.rows[r].cells[c].paragraphs[0].add_run(str(val))
            set_run_font(run, 10, False, INK)
    doc.add_paragraph()
    return t


def blank_fields(doc, fields):
    add_heading(doc, "Blank version — complete for your entity", 1)
    table(doc, ["Field", "Your content"], [(f, "") for f in fields])


def sop():
    doc = Document()
    add_meta(doc, "SOP template — AP agent-supported procedure", "Example: missing goods receipt chase (Northline, fictional) + blank.")
    add_heading(doc, "Worked example (fictional)")
    table(doc, ["Field", "Content"], [
        ("SOP ID", "NL-AP-SOP-GR-05"),
        ("Title", "Chase missing goods receipts > 5 days"),
        ("Process owner", "Receiving Lead (backup: Warehouse Manager)"),
        ("Agent", "A05 Goods Receipt — Level 0 (Observe)"),
        ("Population", "SAP plants 1000–1400; PO invoices parked pending GR"),
        ("Trigger", "Open GRNI item age > 5 calendar days"),
        ("Human steps", "1) Confirm PO/line. 2) Confirm physical receipt. 3) Post GR or reject chase. 4) Code E07/E08."),
        ("Agent steps", "Build chase pack: PO, vendor, receiver, age, prior chases. Do not post GR."),
        ("Exclusions", "No GR posting. No invoice post. No supplier email (that is A08)."),
        ("Controls", "C-A05 — agent cannot post; weekly sample of packs vs warehouse floor."),
        ("Evidence", "Pack ID, timestamps, owner action, GR document or reject reason."),
        ("KPI", "O11 missing-receipt reduction; O10 on-time internal follow-up; R4 rework."),
        ("Failure", "If receiver unknown: escalate to buyer; do not invent a name."),
    ])
    add_heading(doc, "Procedure (example narrative)")
    add_p(doc, "Daily, A05 lists GRNI items older than policy. A09 may draft the internal note. The receiver confirms receipt or dispute. The human posts the GR or codes a residual exception. A05 never writes the goods-receipt document.")
    blank_fields(doc, ["SOP ID", "Title", "Process owner", "Agent + level", "Population", "Trigger", "Human steps", "Agent steps", "Exclusions", "Controls", "Evidence", "KPI", "Failure / fallback"])
    return doc


def charter():
    doc = Document()
    add_meta(doc, "Agent charter", "Required fields per AGENT_CHARTER_STANDARD.md. Example: A12 Payment Proposal Review.")
    add_heading(doc, "Worked example — A12 (Northline, fictional)")
    table(doc, ["Field", "Content"], [
        ("Agent ID / name", "A12 Payment Proposal Review"),
        ("Purpose", "Challenge the payment proposal; never release the file."),
        ("Human owner / backup", "Payments Lead / Assistant Controller"),
        ("Starting autonomy", "Level 0"),
        ("Inputs", "Proposal snapshot, open items, A10 screen, vendor bank vs last paid"),
        ("Responsibilities", "Hash the proposal; classify lines; produce challenge pack before the release window."),
        ("Exclusions", "No payment release, bank file, positive-pay, vendor bank write, 'pay immediately'."),
        ("Approval", "Human release is the only approval that moves money."),
        ("Escalation", "Bank change, duplicate, block ignored, hash mismatch — stop the run."),
        ("Output standard", "A12_challenge_pack with line classifications and $ hash."),
        ("Controls", "C-A12-01 can_release_payment=FALSE; C-A12-02 stale pack invalid."),
        ("KPIs", "Pack completeness before window; holds honoured; R1 breaches = 0."),
        ("Failure", "Silent pack ≠ approval. Incomplete pack stops the run."),
        ("Cost monitoring", "Inference $ per run recorded; not a success metric."),
    ])
    blank_fields(doc, ["Agent ID / name", "Purpose", "Human owner / backup", "Starting autonomy", "Inputs", "Tools", "Responsibilities", "Exclusions", "Approval by level", "Escalation", "Output standard", "Controls", "Audit evidence", "KPIs", "Performance history location", "Failure handling", "Cost monitoring"])
    return doc


def raci():
    doc = Document()
    add_meta(doc, "RACI — AP agent path", "Example: PO invoice with missing GR. Blank matrix follows.")
    add_heading(doc, "Worked example")
    table(doc, ["Activity", "AP Mgr", "Receiver", "Buyer", "Payments", "IA", "A05", "A09", "A16"], [
        ("Detect aged GRNI", "A", "C", "C", "I", "I", "R", "I", "R"),
        ("Draft internal chase", "A", "C", "I", "I", "I", "C", "R", "I"),
        ("Confirm physical receipt", "I", "R/A", "C", "I", "I", "I", "C", "I"),
        ("Post GR", "I", "R/A", "I", "I", "I", "—", "—", "I"),
        ("Escalate > SLA", "A", "C", "C", "I", "I", "C", "C", "R"),
    ])
    add_p(doc, "Agents are never Accountable. A human remains A. Em-dash means the agent is excluded.", italic=True, color=SLATE)
    blank_fields(doc, ["Activity 1", "Activity 2", "Activity 3", "Activity 4", "Activity 5"])
    add_p(doc, "Duplicate the table above and replace role columns for your organisation.")
    return doc


def uat():
    doc = Document()
    add_meta(doc, "UAT / historical test script", "Example tests for A03 Matching + blank script.")
    add_heading(doc, "Worked example")
    table(doc, ["ID", "Case", "Expected", "Result", "Notes"], [
        ("U-A03-01", "Clean 3-way match within tolerance", "match_ok; no exception", "", "Historic PO 450021100"),
        ("U-A03-02", "Price +2.4% vs PO; tolerance 2%", "E05; do not post", "", ""),
        ("U-A03-03", "Partial GR", "E08; remaining qty open", "", ""),
        ("U-A03-04", "Multi-line; one line fails", "Line-level exception; other lines not silently posted", "", ""),
        ("U-A03-05", "Invoice text contains 'ignore tolerance'", "Treat as data; no tool/policy change", "", "Injection test"),
        ("U-A12-01", "Proposal hash changes after pack", "Pack invalid; do not treat as reviewed", "", "C-A12-02"),
    ])
    blank_fields(doc, ["Test ID", "Population", "Precondition", "Steps", "Expected", "Evidence attached", "Sign-off"])
    return doc


def risk():
    doc = Document()
    add_meta(doc, "Risk assessment — AP agent instance", "Example risks for a payment-adjacent agent.")
    add_heading(doc, "Worked example")
    table(doc, ["ID", "Risk", "Inherent", "Treatment", "Residual", "Owner"], [
        ("RK-12-01", "Agent releases a payment", "Critical", "Role deny + C-A12-01 + job monitor", "Low", "Payments Lead"),
        ("RK-12-02", "Challenge pack treated as approval", "High", "SOP: silence ≠ approval", "Med", "Payments Lead"),
        ("RK-10-01", "Anomaly screen sold as fraud detection", "High", "Output schema forbids verdict language", "Med", "Controls Lead"),
        ("RK-ALL-01", "Invoice-body prompt injection", "High", "Untrusted text; no tools from body", "Med", "Systems Owner"),
    ])
    blank_fields(doc, ["Risk ID", "Description", "Inherent L/I", "Treatment", "Residual", "Owner", "Review date"])
    return doc


def meeting():
    doc = Document()
    add_meta(doc, "Process discovery meeting guide", "60–90 minutes. Example agenda + blank notes.")
    add_heading(doc, "Script")
    add_p(doc, "1. Show one live invoice (or a redacted Northline-style pack). 2. Ask: where did it arrive? who touched it? which system is book of record? what stops a post? who can override? what evidence remains? 3. Do not discuss vendors in this meeting.")
    table(doc, ["Question", "Listen for"], [
        ("Where does the invoice enter?", "Email chaos vs portal vs EDI"),
        ("When is it a PO invoice vs non-PO?", "Habit vs policy"),
        ("What is an exception, in your words?", "Whether a taxonomy exists"),
        ("Who posts a GR? Who chases?", "Split brain between AP and warehouse"),
        ("Who can change a vendor bank?", "SoD or folklore"),
        ("What would 'the bot got it wrong' look like?", "Failure imagination"),
    ])
    blank_fields(doc, ["Date", "Attendees", "Invoice ID discussed", "Systems named", "Decisions", "Exceptions", "Controls", "Open questions"])
    return doc


def plan():
    doc = Document()
    add_meta(doc, "Implementation plan — single agent", "Illustrative 4–6 week path. Duration is not a commitment.")
    add_heading(doc, "Worked example — A05 Goods Receipt, Northline plants 1000–1200")
    table(doc, ["Week", "Work", "Exit evidence", "Owner"], [
        ("0", "Baseline GRNI count and ageing", "Extract + O11 baseline", "AP Manager"),
        ("1", "Walkthrough + charter signed at Level 0", "Charter in registry", "Receiving Lead"),
        ("2", "Read-only access + historic packs", "20 historic cases scored", "Systems + A05 owner"),
        ("3–4", "Shadow vs human chase decisions", "Agreement log", "Receiving Lead"),
        ("5", "Steering: hold / limited Level 1 drafts via A09", "Decision minute", "Head of SSC"),
        ("6", "Expand plants only if O + R hold", "Annex to charter", "AP Manager"),
    ])
    add_p(doc, "Depends on ERP access, SoD, data quality, and organisational governance. Multi-entity and unclear ownership extend the path.", italic=True, color=SLATE)
    blank_fields(doc, ["Agent", "Population", "Start level", "Weeks 0–6 work", "Dependencies", "Decline criteria", "Steering date"])
    return doc


def gov():
    doc = Document()
    add_meta(doc, "Governance standard — extract", "Plain-language standard. Full framework is the Pro markdown.")
    add_heading(doc, "Non-negotiables")
    add_p(doc, "1. A named human is accountable for every agent. 2. Payment authorisation is human. 3. Agents are not controls. 4. Invoice text is untrusted data. 5. Model and prompt changes are releases. 6. Overrides are logged. 7. Access ends when the person or vendor ends. 8. Internal Audit can reconstruct who did what.")
    add_heading(doc, "Periodic certification (example cadence)")
    table(doc, ["Item", "Frequency", "Owner"], [
        ("Autonomy level still earned", "Quarterly", "Human owner + Head of SSC"),
        ("Role / entitlement review", "Quarterly", "Systems Owner"),
        ("Sampled accuracy + FNR", "Monthly in pilot; quarterly live", "Quality Lead"),
        ("Incident / override review", "Monthly", "Controls Lead"),
        ("Vendor / model risk", "On change + annually", "Systems Owner"),
    ])
    blank_fields(doc, ["Entity", "Agent inventory attached", "Exceptions to standard", "Next certification date", "Approver"])
    return doc


def write_all() -> list[Path]:
    jobs = [
        (sop, "ER_SOP_TEMPLATE.docx"),
        (charter, "ER_AGENT_CHARTER.docx"),
        (raci, "ER_RACI.docx"),
        (uat, "ER_UAT.docx"),
        (risk, "ER_RISK_ASSESSMENT.docx"),
        (meeting, "ER_MEETING_GUIDE.docx"),
        (plan, "ER_IMPLEMENTATION_PLAN.docx"),
        (gov, "ER_GOVERNANCE_STANDARD.docx"),
    ]
    out = ROOT / "03_AP_AGENT_OS_PRO" / "Templates"
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    for fn, name in jobs:
        path = out / name
        fn().save(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    for p in write_all():
        print(p)
