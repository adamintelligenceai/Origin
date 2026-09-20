#!/usr/bin/env python3
"""Evidence Room executive decks — institutional, not start-up theatre."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
INK = RGBColor(0x14, 0x17, 0x1C)
FOREST = RGBColor(0x1E, 0x5C, 0x45)
PAPER = RGBColor(0xF3, 0xEF, 0xE6)
CREAM = RGBColor(0xFA, 0xF7, 0xF0)
RULE = RGBColor(0xB8, 0xA8, 0x88)
SLATE = RGBColor(0x5C, 0x63, 0x70)
RUST = RGBColor(0x8C, 0x3A, 0x2F)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)


def set_run(run, text, size=18, bold=False, color=INK, italic=False, name="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run() if p.runs else p.runs[0] if False else p.add_run(), text, size, bold, color, italic)
    # python-pptx: empty paragraph has no runs until we add
    return box


def textbox(slide, l, t, w, h, lines):
    """lines: list of (text, size, bold, color, italic)."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        text, size, bold, color = line[0], line[1], line[2], line[3]
        italic = line[4] if len(line) > 4 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        set_run(run, text, size, bold, color, italic)
    return box


def footer(slide, page, total, label):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35))
    bar.fill.solid()
    bar.fill.fore_color.rgb = INK
    bar.line.fill.background()
    textbox(slide, 0.4, 7.18, 9, 0.28, [(f"EVIDENCE ROOM  ·  {label}", 10, False, WHITE, False)])
    textbox(slide, 11.2, 7.18, 1.8, 0.28, [(f"{page} / {total}", 10, False, WHITE, False)])


def bg(slide, color=PAPER):
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    fill.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def rule(slide, l, t, w=2.2):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(0.03))
    s.fill.solid()
    s.fill.fore_color.rgb = FOREST
    s.line.fill.background()


def blank():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def cover(prs, kicker, title, subtitle, meta, total, idx, label):
    s = new_slide(prs)
    bg(s, INK)
    mark = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.6), Inches(0.28), Inches(0.28))
    mark.fill.solid()
    mark.fill.fore_color.rgb = FOREST
    mark.line.fill.background()
    textbox(s, 1.15, 0.58, 6, 0.35, [("EVIDENCE ROOM", 14, True, WHITE)])
    textbox(s, 0.7, 2.1, 11, 0.35, [(kicker, 12, False, RULE)])
    textbox(s, 0.7, 2.5, 12, 2.2, [(title, 36, True, WHITE)])
    textbox(s, 0.7, 5.0, 10, 1.0, [(subtitle, 18, False, PAPER)])
    textbox(s, 0.7, 6.6, 10, 0.3, [(meta, 12, False, RULE)])
    textbox(s, 11.2, 6.6, 1.8, 0.3, [(f"{idx} / {total}", 12, False, RULE)])


def section(prs, num, title, total, idx, label):
    s = new_slide(prs)
    bg(s)
    textbox(s, 0.7, 2.4, 3, 0.4, [(num, 16, False, FOREST)])
    textbox(s, 0.7, 2.9, 11, 1.6, [(title, 32, True, INK)])
    rule(s, 0.7, 4.6)
    footer(s, idx, total, label)


def bullets(prs, title, items, total, idx, label, note=None):
    s = new_slide(prs)
    bg(s)
    textbox(s, 0.7, 0.35, 11, 0.6, [(title, 26, True, INK)])
    rule(s, 0.7, 1.0)
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(12), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        set_run(run, "–  " + item, 16, False, INK)
        p.space_after = Pt(10)
    if note:
        textbox(s, 0.7, 6.6, 12, 0.4, [(note, 11, False, SLATE, True)])
    footer(s, idx, total, label)


def two_col(prs, title, left_h, left, right_h, right, total, idx, label):
    s = new_slide(prs)
    bg(s)
    textbox(s, 0.7, 0.35, 12, 0.55, [(title, 26, True, INK)])
    rule(s, 0.7, 0.95)
    left_sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.25), Inches(5.7), Inches(5.4))
    left_sh.fill.solid()
    left_sh.fill.fore_color.rgb = CREAM
    left_sh.line.color.rgb = RULE
    right_sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.7), Inches(1.25), Inches(5.9), Inches(5.4))
    right_sh.fill.solid()
    right_sh.fill.fore_color.rgb = CREAM
    right_sh.line.color.rgb = RULE
    textbox(s, 0.95, 1.45, 5.3, 0.4, [(left_h, 14, True, FOREST)])
    textbox(s, 6.95, 1.45, 5.5, 0.4, [(right_h, 14, True, RUST)])
    lb = s.shapes.add_textbox(Inches(0.95), Inches(1.95), Inches(5.3), Inches(4.4))
    tf = lb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(left):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        set_run(run, "–  " + item, 14, False, INK)
        p.space_after = Pt(8)
    rb = s.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.5), Inches(4.4))
    tf = rb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(right):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        set_run(run, "–  " + item, 14, False, INK)
        p.space_after = Pt(8)
    footer(s, idx, total, label)


def cfo_deck() -> Presentation:
    prs = blank()
    label = "CFO AP TRANSFORMATION"
    slides = []

    def rec(fn):
        slides.append(fn)

    rec(lambda i, n: cover(
        prs,
        "AP AGENT OS  ·  EXECUTIVE BRIEFING",
        "Responsibility is earned.",
        "An operating system for building, governing and scaling AI agents across Accounts Payable — without replacing your ERP, and without autonomous payments.",
        "Evidence Room  ·  Version 1.0  ·  September 2026  ·  evidenceroom.ai",
        n, i, label,
    ))
    rec(lambda i, n: bullets(prs, "The problem, stated for a CFO", [
        "You already have an ERP, probably an AP suite, a bank portal, and a shared-services team.",
        "The gap is not another capture tool. It is the operating system for an agent layer: job, owner, autonomy, control, evidence.",
        "Pilots die in slides because nobody can say what the agent is allowed to do, or what would count as progress.",
        "That is a responsibility problem, not a model problem.",
    ], n, i, label))
    rec(lambda i, n: two_col(
        prs, "What this is / is not",
        "Evidence Room is",
        ["A method, control system, measurement model, and 16-agent architecture.",
         "ERP-agnostic: SAP, D365, Oracle, NetSuite, Workday, others.",
         "Designed so a team can start Monday after a Friday purchase.",
         "A path from Observe → Recommend → Prepare → guarded Execute."],
        "Evidence Room is not",
        ["An ERP or AP automation replacement.",
         "A ChatGPT prompt pack.",
         "An autonomous payment system.",
         "A guarantee of savings, fraud detection, compliance, or ROI."],
        n, i, label,
    ))
    rec(lambda i, n: bullets(prs, "Industry context — not a forecast", [
        "Ardent Partners (2024 Best-in-Class, via Payables Place 21 Jan 2025): BIC AP teams show 78% lower invoice cost, 82% faster cycle, 59% lower exceptions vs peers.",
        "Best-in-Class exception rate reported at 9% of invoices.",
        "STP capability in place at 69% of Best-in-Class; they are 2.4× more likely to process straight-through.",
        "Forrester (Meng Liu, 17 Mar 2025) names six AI AP clusters: capture, matching, reporting, fraud management, payment management, e-invoicing/tax.",
        "Use these as ambition context. Your baseline remains your baseline.",
    ], n, i, label, "Independent research. Not Evidence Room results. Not a promise."))
    rec(lambda i, n: bullets(prs, "The sixteen-agent workforce", [
        "Invoice path: Intake → Duplicate screen → Validation → Matching → Approval (if required) → human or guarded post.",
        "Exception path: Triage → GR / PO quality / Supplier / Internal follow-up.",
        "Periodic path: Statement recon, payment-proposal challenge, close, reporting, root cause.",
        "A16 Orchestrator dispatches and enforces gates. It cannot promote itself.",
        "A12 challenges the payment run. A human always releases.",
    ], n, i, label))
    rec(lambda i, n: bullets(prs, "Responsibility model", [
        "Level 0 Observe — reviews, cannot act.",
        "Level 1 Recommend — recommendations for humans.",
        "Level 2 Prepare — drafts actions; approval required.",
        "Level 3 Execute within guardrails — pre-approved low-risk actions only.",
        "Level 4 Managed autonomy — scoped, time-boxed, exception-based oversight.",
        "Every agent starts at 0 or 1. Promotion requires a measured pack. Full autonomy is never the default.",
    ], n, i, label))
    rec(lambda i, n: two_col(
        prs, "What the CFO should insist on",
        "Before any write-access",
        ["Named human owner and backup.",
         "Written exclusions (especially payment release).",
         "Least-privilege service identity.",
         "Shadow versus human log.",
         "FPR / FNR on a labelled sample."],
        "Before any 'savings' slide",
        ["Customer cost model, not a vendor blog.",
         "Hours-to-cash conversion stated.",
         "Control-breach count in the same pack.",
         "Finance signature on validated savings.",
         "A decline option if evidence is thin."],
        n, i, label,
    ))
    rec(lambda i, n: bullets(prs, "90-day shape (illustrative, not a commitment)", [
        "Days 1–15: baseline diagnostic, one process walkthrough, one agent charter.",
        "Days 16–35: historical test + shadow on a bounded population.",
        "Days 36–60: controlled pilot if O + R families hold.",
        "Days 61–90: steering review; promote, hold, or retire.",
        "A single well-bounded agent can move in 4–6 weeks. Multi-ERP, weak master data, or unclear SoD extends this. Say so.",
    ], n, i, label))
    rec(lambda i, n: bullets(prs, "Ask of this room", [
        "Name an AP process owner with authority to refuse a bad pilot.",
        "Fund measurement before model spend.",
        "Keep payment authorisation human.",
        "Treat agent changes as releases.",
        "Buy the operating system — then, if needed, the productised blueprint — before you buy another platform.",
    ], n, i, label))
    rec(lambda i, n: cover(
        prs,
        "NEXT STEP",
        "Assess AP Agent Readiness.",
        "Free diagnostic · 36 questions · customer baseline. Then Starter ($79), Professional ($199), Team ($499), or a scoped Blueprint.",
        "evidenceroom.ai  ·  Responsibility is earned.",
        n, n, label,
    ))

    n = len(slides)
    for i, fn in enumerate(slides, 1):
        fn(i, n)
    return prs


def workshop_deck() -> Presentation:
    prs = blank()
    label = "AP AGENT WORKSHOP"
    blocks = []
    blocks.append(lambda i, n: cover(prs, "TEAM EDITION  ·  FACILITATION", "Redesign the AP agent layer in two days.",
                                     "A working session for AP, controllership, procurement, receiving, systems, and internal audit.",
                                     "Evidence Room AP Agent OS · Team  ·  Version 1.0", n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Outcomes by 16:30 tomorrow", [
        "A current-state map for one invoice path (not the entire enterprise).",
        "Coded exceptions using E01–E27.",
        "One signed agent charter at Level 0 or 1.",
        "A control extract including C-A12-01 (human payment release).",
        "A 4–6 week test plan with named owners.",
    ], n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Day 1 — Observe to Structure", [
        "09:00  Purpose, rules, what we will not decide today.",
        "09:30  Walkthrough: one real invoice, start to park/post.",
        "11:00  Extract steps, systems, decisions, controls.",
        "13:00  Exception taxonomy coding of last month's top 20.",
        "15:00  RACI for the chosen path.",
        "16:30  Playback. No tooling debate.",
    ], n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Day 2 — Agentise to Measure", [
        "09:00  Human / recommend / prepare / execute / deterministic split.",
        "10:30  Draft the first charter (usually A04, A05, or A10).",
        "12:30  Controls and evidence list.",
        "14:00  Historical sample design and shadow protocol.",
        "15:30  Steering one-pager. Risks. Decline criteria.",
    ], n, i, label))
    blocks.append(lambda i, n: two_col(prs, "House rules",
        "We will",
        ["Use Northline-style examples if live data cannot leave the room.",
         "Write exclusions before features.",
         "Timebox tooling talk to 15 minutes.",
         "Leave with owners, not ideas."],
        "We will not",
        ["Pick a model vendor on day one.",
         "Promise savings.",
         "Let an agent approve or pay.",
         "Call a demo a control."],
        n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Exercise: first agent, not favourite agent", [
        "Score candidates on: data available, control clarity, reversible failure, volume, political heat.",
        "Prefer a high-volume, well-coded exception (missing GR, price variance) over a glamorous 'AI payables' story.",
        "If payment proposal is proposed first, start A12 at Level 0 only.",
    ], n, i, label))
    n = len(blocks)
    for i, fn in enumerate(blocks, 1):
        fn(i, n)
    return prs


def steering_deck() -> Presentation:
    prs = blank()
    label = "STEERING COMMITTEE"
    blocks = []
    blocks.append(lambda i, n: cover(prs, "STEERING PACK  ·  MONTHLY", "Agent performance, not model theatre.",
                                     "One pack: activity, operational outcomes, financial (only if signed), risk-control. Promote, hold, or retire.",
                                     "Template · replace Northline figures · Version 1.0", n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Decision required today", [
        "Hold A01 at Level 1 — extraction accuracy 93% on golden set; FNR still 6%.",
        "Do not promote A12. Challenge packs are useful. Release remains human.",
        "Open incident: one unlogged override (R1). Close before any scope expansion.",
        "Approve 20 more days of A05 shadow on plants 1000–1200 only.",
    ], n, i, label, "Sample decisions for the template. Replace with live registry."))
    blocks.append(lambda i, n: bullets(prs, "Operating measures (illustrative layout)", [
        "O5 STP 31% → 36% on the pilot population (same entity, same channel).",
        "O3 FPR 18% → 14%. O4 FNR 7% → 6%. Do not celebrate FPR if FNR rose.",
        "O11 missing receipts 420 → 360 open items > policy.",
        "R1 control breaches: 0 this period after prior 2.",
        "F5 validated savings: $0 (unsigned). F4 hours released: 240 (capacity).",
    ], n, i, label, "Northline-style illustration. Not a customer result."))
    blocks.append(lambda i, n: two_col(prs, "Promote / hold / retire",
        "Promote only if",
        ["Charter still true.",
         "Sampled accuracy above written gate.",
         "No open severity-1 control incident.",
         "Human owner signs.",
         "A16 registry updated with evidence link."],
        "Retire or pause if",
        ["FNR worse than baseline.",
         "Unlogged override.",
         "Model changed without release.",
         "Owner left and backup unsigned.",
         "The job is better as deterministic rules."],
        n, i, label))
    n = len(blocks)
    for i, fn in enumerate(blocks, 1):
        fn(i, n)
    return prs


def case_deck() -> Presentation:
    prs = blank()
    label = "BUSINESS CASE"
    blocks = []
    blocks.append(lambda i, n: cover(prs, "INVESTMENT CASE  ·  ILLUSTRATIVE", "Capacity first. Cash second. Controls always.",
                                     "A model for discussing AP agent economics without false precision.",
                                     "Use ER_AP_ROI_CALCULATOR.xlsx  ·  Version 1.0", n, i, label))
    blocks.append(lambda i, n: bullets(prs, "How to talk about money without lying", [
        "Start from the customer's cost model, not $15 vs $3 blog figures.",
        "Hours released are not savings until hours-to-cash conversion is applied and Finance signs.",
        "Show Conservative / Base / Upside. If Conservative does not work, do not scale.",
        "Put AI run-rate and implementation cost on the same page as benefits.",
        "Never present F without R.",
    ], n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Ardent context (independent)", [
        "Best-in-Class defined as the 20% with lowest processing cost and shortest cycle (Ardent 2024).",
        "Those teams: 78% lower cost, 82% faster, 59% lower exceptions; 9% exception rate.",
        "That is a description of a peer set, not a model output for your entity.",
    ], n, i, label))
    blocks.append(lambda i, n: bullets(prs, "Northline illustration (fictional)", [
        "18,000 invoices / month · 14 AP FTE · 20% exception rate · 22 minutes each.",
        "Base case in the calculator can show capacity and cost-to-serve movement.",
        "It may not show cash payback. If it does not, say so — as the workbook does.",
        "Replace every input before this slide leaves the building.",
    ], n, i, label, "ILLUSTRATIVE. Not a customer case study."))
    n = len(blocks)
    for i, fn in enumerate(blocks, 1):
        fn(i, n)
    return prs


def write_all() -> list[Path]:
    jobs = [
        (cfo_deck, ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "ER_CFO_AP_TRANSFORMATION.pptx"),
        (workshop_deck, ROOT / "04_AP_AGENT_OS_TEAM" / "Workshop" / "ER_AP_WORKSHOP_DECK.pptx"),
        (steering_deck, ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "ER_STEERING_COMMITTEE.pptx"),
        (case_deck, ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "ER_BUSINESS_CASE_DECK.pptx"),
    ]
    paths = []
    for fn, path in jobs:
        path.parent.mkdir(parents=True, exist_ok=True)
        fn().save(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    for p in write_all():
        print(p)
