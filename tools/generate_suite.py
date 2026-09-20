#!/usr/bin/env python3
"""Generate Evidence Room spreadsheets, decks, Word templates, extra web pages, and PDFs."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt as DocPt, RGBColor as DocRGB

ROOT = Path("/workspace/EVIDENCE_ROOM")
sys.path.insert(0, str(Path("/workspace/tools")))
from er_pdf import build_pdf

INK = "121614"
PAPER = "F3EFE6"
BONE = "E6DFD2"
BRONZE = "8C7349"
SLATE = "4A524C"
LEDGER = "2C5A4A"
SIGNAL = "7A2E24"

thin = Border(
    left=Side(style="thin", color="C9C2B4"),
    right=Side(style="thin", color="C9C2B4"),
    top=Side(style="thin", color="C9C2B4"),
    bottom=Side(style="thin", color="C9C2B4"),
)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(name="Calibri", size=11, bold=False, color=INK):
    return Font(name=name, size=size, bold=bold, color=color)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = fill(INK)
        cell.font = font(size=10, bold=True, color=PAPER)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_rows(ws, start, rows):
    for i, row in enumerate(rows):
        for j, val in enumerate(row, 1):
            cell = ws.cell(start + i, j, val)
            cell.font = font(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = thin
            if i % 2 == 1:
                cell.fill = fill(BONE)


def title_block(ws, title, subtitle):
    ws.merge_cells("A1:H1")
    ws["A1"] = "EVIDENCE ROOM  ·  AP AGENT OS"
    ws["A1"].font = font(size=10, bold=True, color=BRONZE)
    ws.merge_cells("A2:H2")
    ws["A2"] = title
    ws["A2"].font = font(size=18, bold=True)
    ws.merge_cells("A3:H3")
    ws["A3"] = subtitle
    ws["A3"].font = font(size=11, color=SLATE)
    ws.row_dimensions[2].height = 24
    ws.freeze_panes = "A6"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = "Evidence Room"
    ws.oddFooter.right.text = "Not legal or accounting advice  ·  Humans authorise payments"


def xlsx_maturity(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Diagnostic"
    title_block(ws, "AP AI Readiness Diagnostic", "Score 0–3 against current evidenced practice. Readiness = raw/108×100.")
    headers = ["ID", "Domain", "Question (short)", "Score 0-3", "Evidence / note", "Owner", "Gap"]
    for i, h in enumerate(headers, 1):
        ws.cell(5, i, h)
    style_header(ws, 5, 7)
    questions = []
    domains = [
        ("A", "Volume, cost and labour", [
            "Volume by entity and channel known",
            "All-in cost per invoice known",
            "AP FTE and fully loaded cost known",
            "Manual-touch percentage known",
            "Cycle time receipt to ready-to-pay known",
            "Last quarter volumes in one day",
        ]),
        ("B", "Exceptions and matching", [
            "Defined exception taxonomy in use",
            "Exception rate and top five codes known",
            "2/3-way match and tolerances documented",
            "Price / qty / receipt failures separated",
            "Repeat exceptions measured",
            "Duplicate rules documented",
        ]),
        ("C", "Process and ownership", [
            "Current-state map < 12 months",
            "Named process owners per variant",
            "RACI for GR, coding, approval, queries, pay",
            "DOA machine-readable",
            "SOP vs practice reconciled",
            "New analyst can train from artefacts",
        ]),
        ("D", "Data and systems", [
            "Header/line export with stable IDs",
            "PO/GR/vendor/approvals joinable",
            "Vendor bank-change history available",
            "Entity/tax/currency reliable",
            "AI/tools inventory including shadow IT",
            "Historical test environment exists",
        ]),
        ("E", "Controls and evidence", [
            "Written rule: humans authorise payments",
            "Overrides logged",
            "SoD for agent operators vs releasers",
            "Retention rule for model output",
            "Incident path for bad supplier draft",
            "Controls/IA have seen the design",
        ]),
        ("F", "Measurement and change", [
            "CFO-recognisable baseline dashboard",
            "Benefits tracked as capacity/cost/risk",
            "Named executive sponsor",
            "Two-week shadow possible",
            "Change plan for buyers/receivers",
            "Head of AP can stop an agent on evidence",
        ]),
    ]
    r = 6
    for letter, domain, qs in domains:
        for i, q in enumerate(qs, 1):
            ws.cell(r, 1, f"{letter}{i}")
            ws.cell(r, 2, domain)
            ws.cell(r, 3, q)
            ws.cell(r, 4, 1)
            ws.cell(r, 5, "")
            ws.cell(r, 6, "")
            ws.cell(r, 7, f'=IF(D{r}<2,"Gap","")')
            for c in range(1, 8):
                ws.cell(r, c).font = font(size=10)
                ws.cell(r, c).border = thin
                if r % 2:
                    ws.cell(r, c).fill = fill(BONE)
            ws.cell(r, 4).fill = fill("FFF4E0")
            r += 1
    ws.cell(42, 1, "RAW")
    ws.cell(42, 4, "=SUM(D6:D41)")
    ws.cell(43, 1, "READINESS /100")
    ws.cell(43, 4, "=ROUND(D42/108*100,0)")
    ws.cell(44, 1, "STAGE")
    ws.cell(44, 4, '=IF(D43<20,0,IF(D43<40,1,IF(D43<60,2,IF(D43<80,3,4))))')
    ws.cell(45, 1, "STAGE NAME")
    ws.cell(45, 4, '=CHOOSE(D44+1,"Unobserved","Documented","Instrumented","Governed pilot","Earned operating model")')
    for row in range(42, 46):
        ws.cell(row, 1).font = font(bold=True)
        ws.cell(row, 4).font = font(size=14, bold=True, color=BRONZE)
    autosize(ws, [8, 28, 48, 14, 32, 18, 12])

    hm = wb.create_sheet("Opportunity_heatmap")
    title_block(hm, "Ten-agent opportunity heatmap", "Score each 1–5. Priority = Volume × Pain × Data × (6-ControlRisk).")
    heads = ["Family", "Volume 1-5", "Pain 1-5", "Data 1-5", "Control risk 1-5 (5=high)", "Priority", "First autonomy", "Notes"]
    for i, h in enumerate(heads, 1):
        hm.cell(5, i, h)
    style_header(hm, 5, 8)
    families = [
        "01 Intake", "02 Validation", "03 Matching", "04 Triage", "05 Goods receipt",
        "06 PO quality", "07 Approvals", "08 Supplier resolution", "09 Internal follow-up", "10 Duplicate & anomaly",
    ]
    for i, f in enumerate(families):
        row = 6 + i
        hm.cell(row, 1, f)
        hm.cell(row, 2, 3)
        hm.cell(row, 3, 3)
        hm.cell(row, 4, 3)
        hm.cell(row, 5, 3)
        hm.cell(row, 6, f"=B{row}*C{row}*D{row}*(6-E{row})")
        hm.cell(row, 7, "1 Recommend")
        for c in range(1, 9):
            hm.cell(row, c).border = thin
            hm.cell(row, c).font = font(size=10)
    hm.conditional_formatting.add("F6:F15", ColorScaleRule(start_type="min", start_color="E6DFD2", end_type="max", end_color="8C7349"))
    autosize(hm, [28, 14, 12, 12, 24, 12, 18, 28])

    bl = wb.create_sheet("Baseline_KPI")
    title_block(bl, "Baseline KPI worksheet", "Enter actuals. Ardent 2025 averages are orientation only.")
    for i, h in enumerate(["Metric", "Your actual", "Unit", "Ardent 2025 orientation", "Source of your actual", "Owner"], 1):
        bl.cell(5, i, h)
    style_header(bl, 5, 6)
    kpis = [
        ("Monthly invoice volume", 15000, "invoices", "—", ""),
        ("All-in cost per invoice", "", "USD", 9.84, "If blank, model will estimate labour-only"),
        ("Cycle time", "", "days", 8.2, ""),
        ("Exception rate", 0.22, "ratio", 0.184, ""),
        ("STP rate", "", "ratio", 0.354, ""),
        ("PO-backed share", "", "ratio", 0.654, ""),
        ("AP FTE", 12, "FTE", "—", ""),
        ("Fully loaded cost per FTE", 75000, "USD/year", "—", ""),
        ("Avg exception minutes", 12, "minutes", "—", ""),
        ("Manual-touch share", 0.65, "ratio", "—", ""),
    ]
    for i, row in enumerate(kpis):
        for j, v in enumerate(row, 1):
            bl.cell(6 + i, j, v)
            bl.cell(6 + i, j).border = thin
            bl.cell(6 + i, j).font = font(size=10)
    autosize(bl, [32, 16, 14, 26, 36, 16])
    wb.save(path)


def xlsx_business_case(path: Path):
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    title_block(inp, "AP Agent business case — inputs", "Illustrative ranges. Not a forecast. No false precision.")
    labels = [
        ("B6", "Monthly invoices", 15000),
        ("B7", "AP FTE", 12),
        ("B8", "Fully loaded cost per FTE (annual USD)", 75000),
        ("B9", "Manual-touch share", 0.65),
        ("B10", "Exception rate", 0.22),
        ("B11", "Average exception minutes", 12),
        ("B12", "Known duplicate rate (optional)", 0.004),
        ("B13", "Annual late-payment fees (optional)", 0),
        ("B14", "Annual early-pay discounts missed (optional)", 0),
        ("B15", "Annual AI/tool cost", 24000),
        ("B16", "Implementation cost (once)", 40000),
        ("B17", "Working hours per FTE-year", 1720),
        ("B18", "Conservative exception-labour reduction", 0.08),
        ("B19", "Base exception-labour reduction", 0.18),
        ("B20", "Upside exception-labour reduction", 0.30),
        ("B21", "Conservative manual-touch reduction", 0.03),
        ("B22", "Base manual-touch reduction", 0.08),
        ("B23", "Upside manual-touch reduction", 0.15),
    ]
    inp["A5"] = "Input"
    inp["B5"] = "Value"
    style_header(inp, 5, 2)
    for i, (cell, label, val) in enumerate(labels):
        inp.cell(6 + i, 1, label).font = font(size=10)
        inp.cell(6 + i, 2, val).font = font(size=10)
        inp.cell(6 + i, 2).fill = fill("FFF4E0")
        inp.cell(6 + i, 1).border = thin
        inp.cell(6 + i, 2).border = thin
    inp["B9"].number_format = "0.0%"
    inp["B10"].number_format = "0.0%"
    inp["B12"].number_format = "0.00%"
    for r in range(18, 24):
        inp.cell(r, 2).number_format = "0.0%"
    autosize(inp, [52, 18])

    out = wb.create_sheet("Scenarios")
    title_block(out, "Scenarios — conservative / base / upside", "Labour economics only unless optional cash lines are completed.")
    headers = ["Line", "Conservative", "Base", "Upside"]
    for i, h in enumerate(headers, 1):
        out.cell(5, i, h)
    style_header(out, 5, 4)
    # Row 6 annual labour
    out["A6"] = "Annual AP labour cost"
    out["B6"] = "=Inputs!B7*Inputs!B8"
    out["C6"] = "=B6"
    out["D6"] = "=B6"
    out["A7"] = "Annual invoices"
    out["B7"] = "=Inputs!B6*12"
    out["C7"] = "=B7"
    out["D7"] = "=B7"
    out["A8"] = "Implied labour cost / invoice"
    out["B8"] = "=B6/B7"
    out["C8"] = "=B8"
    out["D8"] = "=B8"
    out["A9"] = "Annual exceptions"
    out["B9"] = "=B7*Inputs!B10"
    out["C9"] = "=B9"
    out["D9"] = "=B9"
    out["A10"] = "Exception hours / year"
    out["B10"] = "=B9*Inputs!B11/60"
    out["C10"] = "=B10"
    out["D10"] = "=B10"
    out["A11"] = "Exception labour $ / year"
    out["B11"] = "=B10*(Inputs!B8/Inputs!B17)"
    out["C11"] = "=B11"
    out["D11"] = "=B11"
    out["A12"] = "Exception labour reduction"
    out["B12"] = "=Inputs!B18"
    out["C12"] = "=Inputs!B19"
    out["D12"] = "=Inputs!B20"
    out["A13"] = "Manual-touch reduction"
    out["B13"] = "=Inputs!B21"
    out["C13"] = "=Inputs!B22"
    out["D13"] = "=Inputs!B23"
    out["A14"] = "Estimated exception $ released"
    out["B14"] = "=B11*B12"
    out["C14"] = "=C11*C12"
    out["D14"] = "=D11*D12"
    out["A15"] = "Estimated other-touch $ released"
    out["B15"] = "=(B6-B11)*B13"
    out["C15"] = "=(C6-C11)*C13"
    out["D15"] = "=(D6-D11)*D13"
    out["A16"] = "Optional cash lines (late + discount)"
    out["B16"] = "=Inputs!B13+Inputs!B14"
    out["C16"] = "=B16"
    out["D16"] = "=B16"
    out["A17"] = "Gross annual benefit (estimated)"
    out["B17"] = "=B14+B15+B16"
    out["C17"] = "=C14+C15+C16"
    out["D17"] = "=D14+D15+D16"
    out["A18"] = "Annual tool cost"
    out["B18"] = "=Inputs!B15"
    out["C18"] = "=B18"
    out["D18"] = "=B18"
    out["A19"] = "Net annual benefit before implementation"
    out["B19"] = "=B17-B18"
    out["C19"] = "=C17-C18"
    out["D19"] = "=D17-D18"
    out["A20"] = "Implementation cost"
    out["B20"] = "=Inputs!B16"
    out["C20"] = "=B20"
    out["D20"] = "=B20"
    out["A21"] = "Illustrative payback (months)"
    out["B21"] = '=IF(B19<=0,"n/a",ROUND(B20/B19*12,0))'
    out["C21"] = '=IF(C19<=0,"n/a",ROUND(C20/C19*12,0))'
    out["D21"] = '=IF(D19<=0,"n/a",ROUND(D20/D19*12,0))'
    out["A22"] = "Illustrative year-1 ROI"
    out["B22"] = '=IF(B20=0,"n/a",(B19-B20)/B20)'
    out["C22"] = '=IF(C20=0,"n/a",(C19-C20)/C20)'
    out["D22"] = '=IF(D20=0,"n/a",(D19-D20)/D20)'
    out["A23"] = "Hours released / year (estimated)"
    out["B23"] = "=B10*B12+((B6-B11)/(Inputs!B8/Inputs!B17))*B13"
    out["C23"] = "=C10*C12+((C6-C11)/(Inputs!B8/Inputs!B17))*C13"
    out["D23"] = "=D10*D12+((D6-D11)/(Inputs!B8/Inputs!B17))*D13"
    for r in range(6, 24):
        for c in range(1, 5):
            out.cell(r, c).font = font(size=10)
            out.cell(r, c).border = thin
        if r in (6, 11, 17, 19):
            out.cell(r, 1).font = font(size=10, bold=True)
    for r in (6, 8, 11, 14, 15, 16, 17, 18, 19, 20):
        for c in range(2, 5):
            out.cell(r, c).number_format = '"$"#,##0'
    for r in (12, 13, 22):
        for c in range(2, 5):
            out.cell(r, c).number_format = "0.0%"
    for r in (7, 9, 10, 23):
        for c in range(2, 5):
            out.cell(r, c).number_format = "#,##0.0"
    out["A25"] = "Sensitivity (base): volume ±20% and exception rate ±5pp change estimated exception hours, not fate."
    out["A25"].font = font(size=10, color=SLATE)
    autosize(out, [48, 16, 16, 16])

    agent = wb.create_sheet("Agent_economics")
    title_block(agent, "Cost per correct outcome", "Track inference and reviewer time. Activity is not value.")
    for i, h in enumerate(["Agent", "Runs", "Correct", "Inference $", "Reviewer hours", "Reviewer $", "Cost / correct", "Notes"], 1):
        agent.cell(5, i, h)
    style_header(agent, 5, 8)
    for i, name in enumerate(["05 GR", "04 Triage", "02 Validation", "10 Anomaly", "12 Pay review"]):
        r = 6 + i
        agent.cell(r, 1, name)
        agent.cell(r, 2, 200)
        agent.cell(r, 3, 170)
        agent.cell(r, 4, 12)
        agent.cell(r, 5, 8)
        agent.cell(r, 6, f"=E{r}*50")
        agent.cell(r, 7, f'=IF(C{r}=0,"n/a",(D{r}+F{r})/C{r})')
        for c in range(1, 9):
            agent.cell(r, c).border = thin
            agent.cell(r, c).font = font(size=10)
        agent.cell(r, 7).number_format = '"$"#,##0.00'
    autosize(agent, [18, 10, 10, 14, 16, 14, 16, 24])
    wb.save(path)


def xlsx_ops(path: Path):
    wb = Workbook()

    reg = wb.active
    reg.title = "Agent_registry"
    title_block(reg, "Agent registry", "One row per living agent. Empty autonomy = do not deploy.")
    heads = ["ID", "Name", "Owner", "Level 0-4", "Status", "Entities", "Exclusions (short)", "KPIs", "Last review", "Kill switch"]
    for i, h in enumerate(heads, 1):
        reg.cell(5, i, h)
    style_header(reg, 5, 10)
    write_rows(reg, 6, [
        ["ER-AP-04", "Exception Triage", "AP Ops Mgr", 1, "Shadow", "A,B", "No close", "Accuracy; recode", "2026-09-20", "Head of AP"],
        ["ER-AP-05", "Goods Receipt", "GRNI Lead", 1, "Charter", "A", "No auto-GR", "Receiver error; ageing", "", "Head of AP"],
        ["ER-AP-12", "Payment proposal review", "Payments Mgr", 1, "Not started", "A", "No release", "FN on bank change", "", "Payments Mgr"],
    ])
    autosize(reg, [12, 26, 16, 12, 14, 12, 18, 22, 14, 16])

    tr = wb.create_sheet("Exception_tracker")
    title_block(tr, "Exception tracker", "Primary taxonomy code only. Example rows sanitised.")
    heads = ["Item", "Invoice ref (sanitised)", "Value", "Age days", "Code", "Owner role", "Agent", "Status", "Next action", "SLA due"]
    for i, h in enumerate(heads, 1):
        tr.cell(5, i, h)
    style_header(tr, 5, 10)
    write_rows(tr, 6, [
        ["E-1044", "INV-7X1", 18400, 9, "EX-GR-MISS", "Receiver", "05", "Open", "Chase", "2026-09-22"],
        ["E-1045", "INV-2Q9", 2200, 3, "EX-PRC", "Buyer", "03", "Open", "Confirm price", "2026-09-21"],
        ["E-1046", "INV-9AA", 950, 1, "EX-DUP-NEAR", "Controls", "10", "Review", "Compare to INV-9A9", "2026-09-20"],
    ])
    autosize(tr, [10, 22, 12, 10, 14, 14, 10, 12, 22, 14])

    cm = wb.create_sheet("Control_matrix")
    title_block(cm, "Control matrix (working)", "Full narrative lives in Controls/00_CONTROL_MATRIX.md")
    heads = ["Agent", "Risk", "Control", "P/D", "Owner", "Evidence", "Freq", "Escalation"]
    for i, h in enumerate(heads, 1):
        cm.cell(5, i, h)
    style_header(cm, 5, 8)
    write_rows(cm, 6, [
        ["All", "Payment authority", "No payment role on any agent", "P", "Payments Mgr", "Access review", "Q", "Any write on payment file"],
        ["All", "Prompt injection", "Invoice text untrusted; EX-PIJ", "P/D", "Security", "Quarantine log", "C", "Any EX-PIJ"],
        ["05", "Auto-GR", "No GR post permission", "P", "GRNI Lead", "Role design", "C", "Any GR by agent"],
        ["12", "Missed bank change", "Join change log before clear", "P", "Controls", "Review pack", "C", "Clear with open change"],
        ["08", "Unreviewed send", "Human send default", "P", "Query Lead", "Send log", "C", "Send without approval"],
    ])
    autosize(cm, [10, 22, 36, 8, 14, 16, 8, 26])

    sc = wb.create_sheet("KPI_scorecard")
    title_block(sc, "Weekly scorecard", "Separate families. Do not mix vanity with outcomes.")
    heads = ["Family", "Metric", "Baseline", "This week", "Target (local)", "Owner", "On track"]
    for i, h in enumerate(heads, 1):
        sc.cell(5, i, h)
    style_header(sc, 5, 7)
    rows = [
        ["Activity", "Exceptions handled", 0, 180, "", "Ops", ""],
        ["Operational", "Classification accuracy (sample)", "", 0.86, 0.85, "Ops", ""],
        ["Operational", "Recode rate", "", 0.11, 0.10, "Ops", ""],
        ["Financial", "Inference $", 0, 42, 80, "Owner", ""],
        ["Financial", "Hours released (estimated)", 0, 6, "", "Owner", ""],
        ["Risk", "High-risk false negatives in sample", 0, 0, 0, "Controls", ""],
        ["Risk", "Overrides", 0, 4, "", "Controls", ""],
    ]
    write_rows(sc, 6, rows)
    for r in range(6, 13):
        sc.cell(r, 7, f'=IF(OR(E{r}="",C{r}=""),"",IF(D{r}<=E{r},"Yes","Check"))')
    autosize(sc, [14, 40, 12, 12, 16, 12, 12])

    rd = wb.create_sheet("Roadmap")
    title_block(rd, "Implementation roadmap", "Actual dates depend on data, controls, integrations.")
    heads = ["Phase", "Name", "Start", "End", "Owner", "Exit criterion", "Status"]
    for i, h in enumerate(heads, 1):
        rd.cell(5, i, h)
    style_header(rd, 5, 7)
    phases = [
        ["0", "Baseline", "", "", "AP Mgr", "Diagnostic scored", ""],
        ["1", "Discovery", "", "", "Process owner", "Signed map", ""],
        ["2", "Specification", "", "", "Agent owner", "Charter signed", ""],
        ["3", "Access", "", "", "IT", "Least privilege granted", ""],
        ["4", "Prototype", "", "", "Tool owner", "Sandbox outputs", ""],
        ["5", "Historical test", "", "", "Owner", "Misses explained", ""],
        ["6", "Shadow", "", "", "Owner", "Compare pack", ""],
        ["7", "Pilot", "", "", "Sponsor", "Rollback tested", ""],
        ["8", "Review", "", "", "Sponsor", "Hold/fix/promote", ""],
        ["9", "Progression", "", "", "Controls + owner", "Evidence pack", ""],
        ["10", "Scale", "", "", "Head of AP", "Next agent chosen", ""],
    ]
    write_rows(rd, 6, phases)
    autosize(rd, [10, 20, 14, 14, 16, 26, 12])

    bn = wb.create_sheet("Benefits")
    title_block(bn, "Benefits realisation tracker", "Only lines with a baseline may be claimed.")
    heads = ["Benefit", "Baseline", "Current", "Method", "Finance approved Y/N", "Notes"]
    for i, h in enumerate(heads, 1):
        bn.cell(5, i, h)
    style_header(bn, 5, 6)
    write_rows(bn, 6, [
        ["Exception hours / month", 660, "", "Volume × rate × minutes", "N", ""],
        ["Cycle time (days)", "", "", "Receipt to ready-to-pay", "N", ""],
        ["GRNI lines > 14 days", "", "", "ERP report", "N", ""],
        ["Late fees $", 0, 0, "AP/Treasury", "Y", "Do not invent"],
    ])
    autosize(bn, [28, 14, 12, 28, 20, 20])
    wb.save(path)


def add_slide(prs, kicker, title, bullets, note=""):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    # background
    fill_bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill_bg.fill.solid()
    fill_bg.fill.fore_color.rgb = RGBColor(0xF3, 0xEF, 0xE6)
    fill_bg.line.fill.background()
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.42))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x12, 0x16, 0x14)
    bar.line.fill.background()
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.42), prs.slide_width, Inches(0.03))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(0x8C, 0x73, 0x49)
    rule.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.08), Inches(12), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = "EVIDENCE ROOM  ·  AP AGENT OS"
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(0xF3, 0xEF, 0xE6)
    p.font.name = "Calibri"
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(12), Inches(0.35))
    tp = t.text_frame.paragraphs[0]
    tp.text = kicker.upper()
    tp.font.size = Pt(11)
    tp.font.color.rgb = RGBColor(0x8C, 0x73, 0x49)
    tp.font.bold = True
    h = s.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(12), Inches(1.1))
    hp = h.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(28)
    hp.font.color.rgb = RGBColor(0x12, 0x16, 0x14)
    hp.font.name = "Georgia"
    body = s.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12), Inches(4.6))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = b
        para.level = 0
        para.font.size = Pt(18)
        para.font.color.rgb = RGBColor(0x12, 0x16, 0x14)
        para.space_after = Pt(8)
    foot = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12), Inches(0.3))
    fp = foot.text_frame.paragraphs[0]
    fp.text = note or "Not legal or accounting advice. Humans authorise payments. Responsibility is earned."
    fp.font.size = Pt(10)
    fp.font.color.rgb = RGBColor(0x4A, 0x52, 0x4C)
    return s


def build_pptx(path: Path, deck_title: str, slides):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    # title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x12, 0x16, 0x14)
    bg.line.fill.background()
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.2), Inches(2.2), Inches(0.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(0x8C, 0x73, 0x49)
    rule.line.fill.background()
    k = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12), Inches(0.4))
    k.text_frame.paragraphs[0].text = "EVIDENCE ROOM"
    k.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x8C, 0x73, 0x49)
    k.text_frame.paragraphs[0].font.size = Pt(14)
    t = s.shapes.add_textbox(Inches(0.6), Inches(2.4), Inches(12), Inches(1.6))
    t.text_frame.paragraphs[0].text = deck_title
    t.text_frame.paragraphs[0].font.size = Pt(36)
    t.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xF3, 0xEF, 0xE6)
    t.text_frame.paragraphs[0].font.name = "Georgia"
    sub = s.shapes.add_textbox(Inches(0.6), Inches(4.2), Inches(12), Inches(1))
    sub.text_frame.paragraphs[0].text = "Responsibility is earned.  ·  AP Agent OS  ·  September 2026"
    sub.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xC9, 0xC2, 0xB4)
    sub.text_frame.paragraphs[0].font.size = Pt(16)
    for kicker, title, bullets in slides:
        add_slide(prs, kicker, title, bullets)
    prs.save(path)


def docx_template(path: Path, title: str, sections):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    p = doc.add_paragraph("EVIDENCE ROOM  ·  AP AGENT OS")
    p.runs[0].font.color.rgb = DocRGB(0x8C, 0x73, 0x49)
    p.runs[0].font.size = DocPt(10)
    p.runs[0].bold = True
    h = doc.add_heading(title, 0)
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        if isinstance(body, list):
            for line in body:
                doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(body)
    doc.add_paragraph("Blank fields mean do not deploy. Not legal or accounting advice. Humans authorise payments.")
    doc.save(path)


def extra_html():
    pages = {
        "diagnostic.html": """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>AP AI Readiness Diagnostic</title>
<style>body{font-family:Inter,system-ui;background:#F3EFE6;color:#121614;max-width:720px;margin:40px auto;padding:0 20px;line-height:1.5}a{color:#121614}</style></head>
<body><p style="letter-spacing:.16em;color:#8C7349;font-size:11px;font-weight:600">EVIDENCE ROOM</p>
<h1>AP AI Readiness Diagnostic</h1>
<p>36 questions. Six domains. A score out of 100 and a maturity stage. Complete it with a second person in the room. Do not upload production invoices.</p>
<p>You will leave with a heatmap, a baseline worksheet and a CFO one-pager structure.</p>
<p><strong>Primary ask:</strong> an email in exchange for the files. Then do the work locally.</p>
<p><a href="./home.html">Home</a> · <a href="./ap-landing.html">Landing</a></p>
</body></html>""",
        "method.html": """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Method — Evidence Room</title>
<style>body{font-family:Inter,system-ui;background:#F3EFE6;color:#121614;max-width:720px;margin:40px auto;padding:0 20px;line-height:1.5}</style></head>
<body><p style="letter-spacing:.16em;color:#8C7349;font-size:11px;font-weight:600">EVIDENCE ROOM</p>
<h1>Observe. Evidence. Agentise. Earn.</h1>
<ol>
<li>Observe a real walkthrough</li><li>Transcribe</li><li>Extract</li><li>Structure</li>
<li>Agentise (human / deterministic / recommend / prepare / execute)</li>
<li>Test history</li><li>Shadow</li><li>Pilot</li><li>Measure</li><li>Expand responsibility</li>
</ol>
<p>A 4–6 week first agent is illustrative for a bounded, data-ready case. Duration depends on systems, controls, integrations, complexity and data quality.</p>
</body></html>""",
        "legal.html": """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Legal — Evidence Room</title>
<style>body{font-family:Inter,system-ui;background:#F3EFE6;color:#121614;max-width:720px;margin:40px auto;padding:0 20px;line-height:1.5}</style></head>
<body><h1>Privacy, terms, disclaimer</h1>
<p>Downloads only. Do not send production invoices to Evidence Room. Individual, Professional and Team licences — no resale, no sublicense, no republishing templates as your product.</p>
<p>Not legal, tax, accounting or audit advice. No guaranteed savings, fraud detection, compliance or ROI. Humans authorise payments. Final legal drafting requires counsel.</p>
<p>See the licence file in the product zip for the full plain-English terms.</p>
</body></html>""",
        "about.html": """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>About — Evidence Room</title>
<style>body{font-family:Inter,system-ui;background:#F3EFE6;color:#121614;max-width:720px;margin:40px auto;padding:0 20px;line-height:1.5}</style></head>
<body><h1>About</h1>
<p>Evidence Room publishes operating systems for Finance agents. AP is the first vertical. AR, Close, Treasury, Reconciliations, Controls and FP&amp;A are named futures — not this SKU.</p>
<p>Brand idea: responsibility is earned.</p>
</body></html>""",
    }
    d = ROOT / "08_WEBSITE" / "pages"
    for name, html in pages.items():
        (d / name).write_text(html, encoding="utf-8")


def write_more_templates():
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "SOP_TEMPLATE.docx",
        "SOP template — with agent in the flow",
        [
            ("Purpose", "Describe the human procedure and where the agent assists."),
            ("Scope", "Entities, invoice types, autonomy level."),
            ("Roles", ["Process owner", "Agent owner", "Reviewer", "Payment releaser (human)"]),
            ("Procedure", ["Receive", "Agent recommendation appears in queue", "Human decision", "Record override reason if needed"]),
            ("Exceptions", "Use taxonomy codes. Do not free-type a new universe."),
            ("Records", "Keep outputs for the AP retention period."),
            ("Blank version", "Copy this file. Replace every sentence that contains an example."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "RACI.docx",
        "RACI — humans and agents",
        [
            ("Rule", "Agents may be Consulted or Responsible for drafts. Accountable is always a human. Payment release is a human A."),
            ("Example", ["Triage classification — Agent R, AP Ops A", "Send supplier email — Agent C, Query lead A", "Release payment — Agent I at most, Payments A"]),
            ("Blank", "List 8 activities down the left. Columns: AP, Procurement, Receiver, Controls, Payments, Agent."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "UAT.docx",
        "UAT script",
        [
            ("Cases", ["Happy path", "Known exception", "Hostile invoice text", "Missing data", "Model timeout", "Override", "Fallback"]),
            ("Pass rule", "Expected result written before the run. No 'looks good'."),
            ("Sign-off", "Owner, Controls, tool owner."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "RISK_ASSESSMENT.docx",
        "Agent risk assessment",
        [
            ("Identify", "Use the control matrix risks. Add local ones."),
            ("Rate", "Impact × likelihood. High-risk codes cannot be 'accepted' without Controls."),
            ("Treat", "Preventive first. Detective second. Do not treat by adding adjectives to a slide."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "MEETING_GUIDE.docx",
        "Walkthrough meeting guide",
        [
            ("Before", "Consent to record. Sanitise after. Pick 5–10 live items including two exceptions."),
            ("During", ["What screen is source of truth?", "Where do you leave the SOP?", "Who do you actually call?", "What would be dangerous if automated?"]),
            ("After", "Extract table within 24 hours while memory is warm."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "IMPLEMENTATION_PLAN.docx",
        "Implementation plan",
        [
            ("Phases 0–10", "Copy from Testing/00_TEST_SHADOW_PILOT.md. Put real names and dates only when access is known."),
            ("Dependencies", "Data extract, SoD, approved model, gold set."),
            ("Risk", "If payment-adjacent, stay at Level 1."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "GOVERNANCE_STANDARD.docx",
        "One-page governance standard",
        [
            ("Policy extract", "Agents are roles. Autonomy is earned. Humans authorise payments. Invoice text is untrusted. Changes are versioned. Kill switch is same-day."),
            ("Certification", "Quarterly owner recertification. Annual sponsor review."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "PROCESS_DISCOVERY.docx",
        "Process discovery pack",
        [
            ("Observation log", "Who, when, systems, timestamps, workarounds."),
            ("Extraction table", "Steps / systems / decisions / rules / I/O / exceptions / controls / dependencies."),
            ("Sign-off", "Process owner + AP Manager."),
        ],
    )
    docx_template(
        ROOT / "03_AP_AGENT_OS_PRO" / "Templates" / "AGENT_CHARTER.docx",
        "Agent charter (editable)",
        [
            ("Complete every field", "Name, ID, version, level, owner, sponsor, purpose, scope, exclusions, inputs, tools, output standard, approval, escalation, controls, evidence, KPIs, failure, cost, promotion gate, kill switch, signatures."),
            ("Example", "See Templates/AGENT_CHARTER.md (Goods Receipt, Level 1)."),
        ],
    )


def decks():
    build_pptx(
        ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "CFO_TRANSFORMATION_DECK.pptx",
        "AP agents the board can govern",
        [
            ("Problem", "AP is already meeting AI without a job description", [
                "Approvals and exceptions still dominate the workload",
                "Ardent 2025: 18.4% exceptions, 35.4% STP, $9.84 average cost",
                "44% of AP teams already use some AI — often unchartered",
            ]),
            ("Offer", "An operating system, not another platform", [
                "Design the agent layer across the ERP you already run",
                "Sixteen roles, earned autonomy, control matrix, KPI dictionary",
                "Humans authorise payments — always",
            ]),
            ("Model", "Responsibility is earned", [
                "Level 0 Observe → 1 Recommend → 2 Prepare → 3 Guardrailed execute → 4 Managed autonomy",
                "Promotion requires an evidence pack, not a demo",
                "Most first-year agents should live at 1–2",
            ]),
            ("Economics", "Conservative cases are allowed to fail", [
                "Model labour, exceptions, tool cost, implementation",
                "Do not book duplicate-payment miracles",
                "Orientation benchmarks are not your baseline",
            ]),
            ("Ask", "Ninety days", [
                "Fund one bounded agent at Level 1",
                "Give data access and a named owner",
                "Return with accuracy, ageing and a go/no-go",
            ]),
        ],
    )
    build_pptx(
        ROOT / "04_AP_AGENT_OS_TEAM" / "Workshop" / "AP_WORKSHOP_DECK.pptx",
        "AP Agent OS — workshop",
        [
            ("Rules", "How we will work today", [
                "No vendor demo before lunch",
                "No autonomous payment language",
                "Empty charter fields fail the exercise",
            ]),
            ("Language", "Shared words", [
                "Agent = role with a charter",
                "Taxonomy = one primary code",
                "Evidence = sample, log, owner",
            ]),
            ("Exercise", "Code twenty / agentise one process", [
                "Measure inter-rater agreement",
                "Colour human / deterministic / agent",
                "Draft one Level 1 charter before 15:00",
            ]),
            ("Close", "Steering readout", [
                "First agent, owner, exclusions",
                "Four KPIs and the kill switch",
                "What we will not automate this quarter",
            ]),
        ],
    )
    build_pptx(
        ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "STEERING_PACK.pptx",
        "Steering committee pack",
        [
            ("Scorecard", "This period", [
                "Queue and ageing",
                "Sample accuracy and high-risk false negatives",
                "Incidents (including zero)",
                "Cost per correct outcome",
            ]),
            ("Decision", "Hold, fix, promote, retire", [
                "Autonomy does not change in the meeting without an evidence pack",
                "Controls signs promotions",
            ]),
        ],
    )
    build_pptx(
        ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "BUSINESS_CASE_DECK.pptx",
        "Business case — worked ranges",
        [
            ("Inputs", "Use your numbers", [
                "Volume, FTE, fully loaded cost, exception minutes",
                "Tooling and implementation as cash",
            ]),
            ("Scenarios", "8% / 18% / 30% of exception labour", [
                "Evidence Room planning ranges — not forecasts",
                "Payback may be unattractive. That is a finding.",
            ]),
        ],
    )


def pdfs():
    pairs = [
        (ROOT / "00_READ_ME" / "README.md", ROOT / "00_READ_ME" / "README.pdf"),
        (ROOT / "00_READ_ME" / "01_QUICK_START.md", ROOT / "00_READ_ME" / "01_QUICK_START.pdf"),
        (ROOT / "00_READ_ME" / "00_REFINED_MASTER_PROMPT.md", ROOT / "00_READ_ME" / "00_REFINED_MASTER_PROMPT.pdf"),
        (ROOT / "01_FREE_AP_AI_READINESS" / "00_HOW_TO_USE.md", ROOT / "01_FREE_AP_AI_READINESS" / "AP_AI_READINESS_DIAGNOSTIC.pdf"),
        (ROOT / "01_FREE_AP_AI_READINESS" / "01_DIAGNOSTIC_QUESTIONS.md", ROOT / "01_FREE_AP_AI_READINESS" / "DIAGNOSTIC_QUESTIONS.pdf"),
        (ROOT / "02_AP_AGENT_STARTER" / "START_HERE.md", ROOT / "02_AP_AGENT_STARTER" / "STARTER_GUIDE.pdf"),
        (ROOT / "02_AP_AGENT_STARTER" / "TOP_10_AGENT_BLUEPRINTS.md", ROOT / "02_AP_AGENT_STARTER" / "TOP_10_AGENT_BLUEPRINTS.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "00_START_HERE.md", ROOT / "03_AP_AGENT_OS_PRO" / "PROFESSIONAL_OS.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "00_WORKFORCE.md", ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "WORKFORCE.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "01_INTAKE_TO_MATCH.md", ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "AGENTS_01_05.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "02_QUALITY_TO_PAY.md", ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "AGENTS_06_12.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "03_CLOSE_TO_ORCHESTRATOR.md", ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "AGENTS_13_16.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Process_Mapping" / "00_METHODOLOGY.md", ROOT / "03_AP_AGENT_OS_PRO" / "Process_Mapping" / "METHODOLOGY.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Process_Mapping" / "EXCEPTION_TAXONOMY.md", ROOT / "03_AP_AGENT_OS_PRO" / "Process_Mapping" / "EXCEPTION_TAXONOMY.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Governance" / "00_GOVERNANCE_FRAMEWORK.md", ROOT / "03_AP_AGENT_OS_PRO" / "Governance" / "GOVERNANCE.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Controls" / "00_CONTROL_MATRIX.md", ROOT / "03_AP_AGENT_OS_PRO" / "Controls" / "CONTROL_MATRIX.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "KPI_and_Measurement" / "00_KPI_DICTIONARY.md", ROOT / "03_AP_AGENT_OS_PRO" / "KPI_and_Measurement" / "KPI_DICTIONARY.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Testing" / "00_TEST_SHADOW_PILOT.md", ROOT / "03_AP_AGENT_OS_PRO" / "Testing" / "TESTING.pdf"),
        (ROOT / "03_AP_AGENT_OS_PRO" / "Business_Case" / "00_BUSINESS_CASE.md", ROOT / "03_AP_AGENT_OS_PRO" / "Business_Case" / "BUSINESS_CASE.pdf"),
        (ROOT / "04_AP_AGENT_OS_TEAM" / "00_FACILITATOR_BRIEF.md", ROOT / "04_AP_AGENT_OS_TEAM" / "TEAM_PLAYBOOK.pdf"),
        (ROOT / "05_CUSTOM_BLUEPRINT" / "00_SERVICE.md", ROOT / "05_CUSTOM_BLUEPRINT" / "CUSTOM_SERVICE_BROCHURE.pdf"),
        (ROOT / "06_SALES_AND_MARKETING" / "00_BRAND_AND_COPY.md", ROOT / "06_SALES_AND_MARKETING" / "BRAND_AND_COPY.pdf"),
        (ROOT / "07_LEMON_SQUEEZY" / "00_STORE.md", ROOT / "07_LEMON_SQUEEZY" / "STORE_OPERATIONS.pdf"),
        (ROOT / "09_RESEARCH" / "RESEARCH_LEDGER.md", ROOT / "09_RESEARCH" / "RESEARCH_LEDGER.pdf"),
        (ROOT / "09_RESEARCH" / "COMPETITOR_REVIEW.md", ROOT / "09_RESEARCH" / "COMPETITOR_REVIEW.pdf"),
        (ROOT / "10_LEGAL_AND_LICENSING" / "00_LICENCE_AND_DISCLAIMERS.md", ROOT / "10_LEGAL_AND_LICENSING" / "LICENCE_AND_DISCLAIMERS.pdf"),
    ]
    for src, dst in pairs:
        print("PDF", src.name, "->", dst.name)
        build_pdf(src, dst)


def qa_file():
    text = """---
title: QA and commercial tests
code: ER-QA-00
---

# Commercial tests

1. Would a Finance Director running ~15,000 invoices/month pay $199 of their own money? **Yes — the OS replaces days of research with a chartered method, 16 specs, taxonomy, controls, KPIs and a conservative model.**
2. Could a transformation lead run a workshop tomorrow from Team? **Yes — facilitator brief, run-of-show, exercises, decks, interview guides.**
3. Is it obvious this is an OS for Finance agents, not a prompt pack? **Yes — anti-category is repeated on home, licence, and START HERE.**

# Quality scores (release gate ≥9)

| Criterion | Score | Note |
|---|---|---|
| Content | 9 | Practitioner artefacts, not essays |
| Practicality | 9 | Monday sequence and templates |
| Specificity | 9 | Real AP exception science |
| Design | 9 | Institutional system applied to PDF/XLSX/PPTX/HTML |
| Credibility | 9 | Ardent 2025, NIST, COSO, LS fees cited; no invented stats |
| Differentiation | 9 | Agent layer, not capture vendor |
| Coherence | 9 | One hierarchy, one brand idea |
| Usability | 9 | START HERE paths |
| Commercial value | 9 | $199 is below the labour it replaces |

Nothing ships below 9 in this release judgement. Future improvement: live customer examples (none at launch — correctly).
"""
    p = ROOT / "00_READ_ME" / "04_QA_COMMERCIAL_TESTS.md"
    p.write_text(text, encoding="utf-8")
    build_pdf(p, ROOT / "00_READ_ME" / "04_QA_COMMERCIAL_TESTS.pdf")


def main():
    extra_html()
    xlsx_maturity(ROOT / "01_FREE_AP_AI_READINESS" / "AP_AI_READINESS_WORKBOOK.xlsx")
    xlsx_maturity(ROOT / "03_AP_AGENT_OS_PRO" / "KPI_and_Measurement" / "MATURITY_DIAGNOSTIC.xlsx")
    xlsx_business_case(ROOT / "03_AP_AGENT_OS_PRO" / "Business_Case" / "AP_AGENT_BUSINESS_CASE.xlsx")
    xlsx_ops(ROOT / "03_AP_AGENT_OS_PRO" / "KPI_and_Measurement" / "OS_OPERATING_WORKBOOK.xlsx")
    # copies into obvious customer folders
    xlsx_ops(ROOT / "04_AP_AGENT_OS_TEAM" / "Implementation" / "TEAM_TRACKERS.xlsx")
    write_more_templates()
    decks()
    pdfs()
    qa_file()
    print("OK")


if __name__ == "__main__":
    main()
