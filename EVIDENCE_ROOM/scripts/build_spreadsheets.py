#!/usr/bin/env python3
"""
Evidence Room — Commercial spreadsheet pack builder.
Generates openpyxl workbooks into exports/xlsx/.

Brand: ink #0B1F2A · paper #F7F4EE · amber #C47A2C · teal #1F5C5C
Title: Evidence Room — AP Agent OS
Tagline: Agents that earn responsibility.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# —— Brand ——
INK = "0B1F2A"
PAPER = "F7F4EE"
AMBER = "C47A2C"
TEAL = "1F5C5C"
RULE = "8A8580"
HAIRLINE = "D9D2C7"
MUTED = "3D4F5A"
WHITE = "FFFFFF"
ILLUST = "FFF3E0"

OUT_DIR = Path(__file__).resolve().parents[1] / "exports" / "xlsx"

thin = Border(
    left=Side(style="thin", color=HAIRLINE),
    right=Side(style="thin", color=HAIRLINE),
    top=Side(style="thin", color=HAIRLINE),
    bottom=Side(style="thin", color=HAIRLINE),
)

fill_ink = PatternFill("solid", fgColor=INK)
fill_paper = PatternFill("solid", fgColor=PAPER)
fill_amber = PatternFill("solid", fgColor=AMBER)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_illust = PatternFill("solid", fgColor=ILLUST)
fill_input = PatternFill("solid", fgColor="FFFCF7")

font_title = Font(name="Calibri", size=16, bold=True, color=INK)
font_h = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_label = Font(name="Calibri", size=10, bold=True, color=INK)
font_body = Font(name="Calibri", size=10, color=MUTED)
font_input = Font(name="Calibri", size=10, color=INK)
font_amber = Font(name="Calibri", size=10, bold=True, color=AMBER)
font_teal = Font(name="Calibri", size=10, bold=True, color=TEAL)


def style_header_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_ink
        cell.font = font_h
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def autosize(ws, min_w: int = 12, max_w: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = min_w
        for cell in col:
            if cell.value is not None:
                length = max(length, min(len(str(cell.value)), max_w))
        ws.column_dimensions[letter].width = length + 2


def brand_banner(ws, title: str, subtitle: str) -> None:
    ws.merge_cells("A1:H1")
    ws["A1"] = "Evidence Room — AP Agent OS"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=INK)
    ws["A1"].fill = fill_paper
    ws.merge_cells("A2:H2")
    ws["A2"] = title
    ws["A2"].font = font_title
    ws.merge_cells("A3:H3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color=TEAL)
    ws.merge_cells("A4:H4")
    ws["A4"] = "Agents that earn responsibility."
    ws["A4"].font = font_amber


def notes_ardent(ws) -> None:
    """Notes sheet — Ardent context only here, labeled external benchmark."""
    brand_banner(
        ws,
        "Notes & External Benchmarks",
        "Industry context only — not a guarantee of your results",
    )
    rows = [
        ("Source", "Ardent Partners — State of ePayables (2025 survey averages)"),
        ("Avg cost per invoice", "$9.84"),
        ("Exception rate", "18.4%"),
        ("Straight-through processing (STP)", "35.4%"),
        ("Label", "EXTERNAL BENCHMARK — NOT A GUARANTEE"),
        (
            "How to use",
            "Locate your baseline relative to survey averages. Do not substitute these for measured KPIs or treat them as Evidence Room performance claims.",
        ),
        (
            "Forbidden use",
            "Do not market 'guaranteed to reach $9.84'. Worksheets labeled ILLUSTRATIVE require your own inputs.",
        ),
        (
            "Product",
            "Evidence Room provides methodologies and templates. Humans remain accountable for approvals, payments, and audit representations.",
        ),
    ]
    ws["A6"] = "Field"
    ws["B6"] = "Value"
    style_header_row(ws, 6, 2)
    for i, (k, v) in enumerate(rows, start=7):
        ws.cell(row=i, column=1, value=k).font = font_label
        cell = ws.cell(row=i, column=2, value=v)
        cell.font = font_body
        ws.cell(row=i, column=1).border = thin
        cell.border = thin
        if "NOT A GUARANTEE" in str(v):
            cell.fill = fill_illust
            cell.font = font_amber
    autosize(ws)


def disclaimer_sheet(ws) -> None:
    brand_banner(ws, "Disclaimer", "Read before relying on model outputs")
    text = [
        "Evidence Room products provide educational and operational templates for designing, governing, and measuring AI agents that assist Accounts Payable work.",
        "They are not legal, tax, audit, accounting, investment, or compliance advice.",
        "They are not a substitute for your ERP, AP automation platform, payment provider, or professional advisors.",
        "AI agents can be wrong. Humans remain accountable for approvals, postings, payments, supplier master changes, and representations to auditors.",
        "We do not guarantee any reduction in cost per invoice, cycle time, exception rate, fraud losses, compliance findings, or ROI.",
        "Industry benchmarks attributed to Ardent Partners (2025) — $9.84 avg cost/invoice, 18.4% exception rate, 35.4% STP — are third-party survey averages for context only. Cited only on Notes sheets.",
        "Any cell labeled ILLUSTRATIVE is sample data for structure demonstration — replace with your measured figures before decisions.",
        "Confidential — licensed for internal use by the purchasing organisation.",
    ]
    for i, line in enumerate(text, start=6):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = Font(name="Calibri", size=9, italic=True, color=RULE)
        cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[i].height = 36
    autosize(ws)


# ═══════════════════════════════════════════════════════════════
# 1. ROI Calculator
# ═══════════════════════════════════════════════════════════════

def build_roi() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    brand_banner(ws, "ER ROI Calculator — Inputs", "Edit input cells (paper fill) only")

    headers = ["Parameter", "Value", "Unit", "Notes"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, 4)

    inputs = [
        ("Annual invoice volume", 120000, "invoices/year"),
        ("AP headcount (FTE)", 8, "FTE"),
        ("Loaded cost per FTE", 85000, "USD/year"),
        ("Manual-touch share", 0.65, "% of invoices"),
        ("Exception rate", 0.22, "% of invoices"),
        ("Minutes per exception resolution", 18, "minutes"),
        ("Duplicate payment loss rate", 0.0015, "fraction of spend"),
        ("Annual AP spend (proxy)", 48000000, "USD"),
        ("Late-payment cost / year", 120000, "USD"),
        ("Early-pay discount opportunity", 95000, "USD/year missed"),
        ("Current processing cost / invoice", 12.5, "USD"),
        ("Projected AI platform cost / year", 48000, "USD"),
        ("Implementation cost (one-time)", 75000, "USD"),
        ("Efficiency gain — Conservative", 0.12, "fraction"),
        ("Efficiency gain — Base", 0.25, "fraction"),
        ("Efficiency gain — Upside", 0.40, "fraction"),
        ("Labour cost recovery share", 0.70, "of labour savings realised"),
        ("Discount capture improvement", 0.35, "of opportunity"),
        ("Late-pay reduction", 0.40, "of late costs"),
        ("Duplicate loss reduction", 0.50, "of duplicate losses"),
        ("Model years", 3, "years"),
    ]
    for i, (name, val, unit) in enumerate(inputs, start=7):
        ws.cell(row=i, column=1, value=name).font = font_label
        cell = ws.cell(row=i, column=2, value=val)
        cell.font = font_input
        cell.fill = fill_input
        cell.border = thin
        if isinstance(val, float) and val < 1:
            cell.number_format = "0.00%"
        elif isinstance(val, (int, float)) and abs(val) >= 100:
            cell.number_format = "#,##0.00"
        ws.cell(row=i, column=3, value=unit).font = font_body
        ws.cell(row=i, column=4, value="INPUT — editable").font = font_teal
        for c in range(1, 5):
            ws.cell(row=i, column=c).border = thin

    ws["A29"] = "Derived labour pool"
    ws["B29"] = "=B8*B9"
    ws["B29"].number_format = '"$"#,##0'
    ws["C29"] = "Annual loaded AP labour"
    ws["A30"] = "Derived exception hours"
    ws["B30"] = "=B7*B11*B12/60"
    ws["B30"].number_format = "#,##0.0"
    ws["C30"] = "Hours/year on exceptions"
    autosize(ws)

    out = wb.create_sheet("Outputs")
    brand_banner(out, "Scenario Outputs", "Conservative / Base / Upside — formula-driven")
    out["A6"] = "Metric"
    out["B6"] = "Conservative"
    out["C6"] = "Base"
    out["D6"] = "Upside"
    style_header_row(out, 6, 4)

    metrics = [
        ("Labour savings (annual)", "=Inputs!B29*Inputs!B20*Inputs!B23", "=Inputs!B29*Inputs!B21*Inputs!B23", "=Inputs!B29*Inputs!B22*Inputs!B23"),
        ("Exception capacity value", "=Inputs!B30*(Inputs!B9/2080)*Inputs!B20", "=Inputs!B30*(Inputs!B9/2080)*Inputs!B21", "=Inputs!B30*(Inputs!B9/2080)*Inputs!B22"),
        ("Late-pay reduction value", "=Inputs!B15*Inputs!B25*Inputs!B20/Inputs!B21", "=Inputs!B15*Inputs!B25", "=Inputs!B15*Inputs!B25*Inputs!B22/Inputs!B21"),
        ("Early-pay capture", "=Inputs!B16*Inputs!B24*Inputs!B20/Inputs!B21", "=Inputs!B16*Inputs!B24", "=Inputs!B16*Inputs!B24*Inputs!B22/Inputs!B21"),
        ("Duplicate loss avoided", "=Inputs!B14*Inputs!B13*Inputs!B26*Inputs!B20/Inputs!B21", "=Inputs!B14*Inputs!B13*Inputs!B26", "=Inputs!B14*Inputs!B13*Inputs!B26*Inputs!B22/Inputs!B21"),
        ("Gross annual benefit", "=B7+B8+B9+B10+B11", "=C7+C8+C9+C10+C11", "=D7+D8+D9+D10+D11"),
        ("AI + run cost", "=Inputs!B18", "=Inputs!B18", "=Inputs!B18"),
        ("Net annual benefit (Y2+)", "=B12-B13", "=C12-C13", "=D12-D13"),
        ("Year-1 net (after impl.)", "=B14-Inputs!B19", "=C14-Inputs!B19", "=D14-Inputs!B19"),
        ("3-year cumulative net", "=B15+B14*2", "=C15+C14*2", "=D15+D14*2"),
        ("Payback months (approx)", '=IF(B14<=0,"n/a",Inputs!B19/(B14/12))', '=IF(C14<=0,"n/a",Inputs!B19/(C14/12))', '=IF(D14<=0,"n/a",Inputs!B19/(D14/12))'),
        ("Illustrative cost/invoice after", "=Inputs!B17*(1-Inputs!B20)", "=Inputs!B17*(1-Inputs!B21)", "=Inputs!B17*(1-Inputs!B22)"),
    ]
    for i, (name, c, b, u) in enumerate(metrics, start=7):
        out.cell(row=i, column=1, value=name).font = font_label
        for col, f in enumerate((c, b, u), start=2):
            cell = out.cell(row=i, column=col, value=f)
            cell.border = thin
            if "Payback" in name:
                cell.number_format = "0.0"
            elif "cost/invoice" in name:
                cell.number_format = '"$"#,##0.00'
            else:
                cell.number_format = '"$"#,##0'
        out.cell(row=i, column=1).border = thin

    out["A20"] = "ILLUSTRATIVE MODEL — replace Inputs with your data. Not a savings guarantee."
    out["A20"].font = font_amber
    out["A20"].fill = fill_illust
    out.merge_cells("A20:D20")
    autosize(out)

    sens = wb.create_sheet("Sensitivity")
    brand_banner(sens, "Sensitivity — Efficiency × Volume", "Base-scaled net annual benefit")
    volumes = [60000, 90000, 120000, 180000, 240000]
    gains = [0.10, 0.15, 0.20, 0.25, 0.30, 0.35, 0.40]
    sens["A6"] = "Eff. gain \\ Annual volume"
    for i, v in enumerate(volumes, start=2):
        cell = sens.cell(row=6, column=i, value=v)
        cell.number_format = "#,##0"
    style_header_row(sens, 6, 6)
    for r, g in enumerate(gains, start=7):
        sens.cell(row=r, column=1, value=g).number_format = "0%"
        sens.cell(row=r, column=1).font = font_label
        for c, v in enumerate(volumes, start=2):
            formula = (
                f"=Inputs!B8*Inputs!B9*{g}*Inputs!B23"
                f"+(Inputs!B15*Inputs!B25+Inputs!B16*Inputs!B24+Inputs!B14*Inputs!B13*Inputs!B26)"
                f"*({v}/Inputs!B7)*({g}/Inputs!B21)"
                f"-Inputs!B18"
            )
            cell = sens.cell(row=r, column=c, value=formula)
            cell.number_format = '"$"#,##0'
            cell.border = thin
    sens["A15"] = "Cells are formulas keyed off Inputs. Change Inputs to recompute."
    sens["A15"].font = font_body
    autosize(sens)

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)

    path = OUT_DIR / "ER_ROI_Calculator.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 2. KPI Scorecard
# ═══════════════════════════════════════════════════════════════

def build_kpi() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    brand_banner(ws, "ER KPI Scorecard", "Activity · Operational · Financial · Risk-control")

    cols = [
        "Category", "KPI", "Definition", "Formula / measure", "Owner",
        "Frequency", "Target", "Actual", "Sample (ILLUSTRATIVE)", "Status",
    ]
    for i, h in enumerate(cols, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, len(cols))

    rows = [
        ("Activity", "Invoices processed", "Count completed in period", "COUNT(completed)", "AP Ops Lead", "Weekly", "", "", 9840),
        ("Activity", "Agent-assisted volume %", "Share touched by agent workflow", "agent_assisted / total", "Agent Owner", "Weekly", "", "", 0.18),
        ("Activity", "Evidence packs produced", "Cases with complete evidence bundle", "COUNT(evidence_complete)", "Controls Lead", "Weekly", "", "", 412),
        ("Operational", "STP / touchless rate", "No human touch after intake", "touchless / total", "AP Manager", "Monthly", "", "", 0.31),
        ("Operational", "Exception rate", "Manual intervention after capture", "exceptions / total", "Exception Lead", "Weekly", "", "", 0.21),
        ("Operational", "First-pass match rate", "PO invoices matching without correction", "clean_match / PO_invoices", "Matching Lead", "Weekly", "", "", 0.74),
        ("Operational", "Cycle time median (days)", "Receipt → payment-ready", "MEDIAN(cycle_days)", "AP Manager", "Monthly", "", "", 4.2),
        ("Operational", "Time-to-first-touch (hrs)", "Exception queue SLA", "AVG(first_touch_hours)", "Triage Owner", "Weekly", "", "", 6.5),
        ("Financial", "Cost per invoice", "Loaded AP cost ÷ invoices", "loaded_cost / volume", "Controller", "Quarterly", "", "", 11.8),
        ("Financial", "On-time payment rate", "Paid by due / elected discount date", "on_time / total", "Treasury/AP", "Monthly", "", "", 0.89),
        ("Financial", "Early-pay discount capture", "Discounts taken ÷ available", "taken / available", "AP Manager", "Monthly", "", "", 0.42),
        ("Financial", "Late fees / year (run-rate)", "Fees from late pay", "SUM(late_fees)", "Controller", "Monthly", "", "", 98000),
        ("Risk-control", "SoD violations", "Segregation breaches detected", "COUNT(sod_breach)", "Internal Controls", "Monthly", "0", "", 0),
        ("Risk-control", "Force-pay overrides", "Logged overrides of holds", "COUNT(force_pay)", "AP Manager", "Weekly", "", "", 14),
        ("Risk-control", "Duplicate suspect precision", "True dups ÷ flags (sample)", "true_pos / flags_sampled", "A10 Owner", "Monthly", "", "", 0.61),
        ("Risk-control", "Bank-change dual control %", "Verified dual control on bank edits", "dual_ok / bank_changes", "Master Data", "Monthly", "100%", "", 0.98),
        ("Risk-control", "Agent stage demotions", "Agents demoted for evidence failure", "COUNT(demote)", "Governance Board", "Monthly", "", "", 1),
    ]
    for r, row in enumerate(rows, start=7):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val if val != "" else None)
            cell.border = thin
            cell.font = font_body
            if c == 9:
                cell.fill = fill_illust
                cell.font = font_amber
                if isinstance(val, float) and val <= 1:
                    cell.number_format = "0.0%"
                elif isinstance(val, (int, float)) and val >= 100:
                    cell.number_format = "#,##0.00"
            if c == 8:
                cell.fill = fill_input
        ws.cell(row=r, column=1).font = font_label

    ws["A25"] = "Sample column is ILLUSTRATIVE only. Leave Actual blank until you measure."
    ws["A25"].font = font_amber
    ws["A25"].fill = fill_illust
    ws.merge_cells("A25:J25")

    defs = wb.create_sheet("Definitions")
    brand_banner(defs, "KPI Definitions & Cadence", "Agree definitions before first baseline")
    defs["A6"] = "Agree definitions in writing before publishing targets. External benchmarks live only on the Notes sheet."
    defs["A6"].font = font_body
    defs.merge_cells("A6:F6")

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws)
    path = OUT_DIR / "ER_KPI_Scorecard.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 3. Agent Registry
# ═══════════════════════════════════════════════════════════════

AGENTS = [
    ("A01", "Invoice Intake Agent", "Observe→Draft→Propose", "AP Ops Lead", "Pilot", "Capture completeness %; rework rate", 4200),
    ("A02", "Invoice Validation Agent", "Observe→Propose", "Data Steward", "Design", "First-pass validation; false hold rate", 1800),
    ("A03", "Matching Agent", "Propose→Bounded execute", "Matching Lead", "Shadow", "First-pass match; false auto-match", 5100),
    ("A04", "Exception Triage Agent", "Draft→Propose", "Exception Lead", "Pilot", "Time-to-first-touch; taxonomy accuracy", 3600),
    ("A05", "Goods Receipt Agent", "Observe→Propose", "Receiving Liaison", "Design", "GR lag; false missing-GR", 900),
    ("A06", "PO Quality Agent", "Observe→Draft", "Procurement Ops", "Observe", "Preventable mismatch share", 600),
    ("A07", "Approval Agent", "Propose→Bounded execute", "AP Manager", "Design", "Approval cycle time; SoD violations", 2400),
    ("A08", "Supplier Resolution Agent", "Draft→Propose", "Vendor Desk", "Design", "First-response time; edit distance", 1500),
    ("A09", "Internal Follow-Up Agent", "Draft→Propose→Bounded", "AP Ops Lead", "Pilot", "SLA breach; reminder-to-action", 1200),
    ("A10", "Duplicate / Anomaly Agent", "Propose (signals only)", "Controls Lead", "Pilot", "Suspect precision; hold cycle time", 2100),
    ("A11", "Vendor Statement Reconciliation", "Observe→Propose", "AP Accountant", "Backlog", "Statement match %; open items ageing", 0),
    ("A12", "Payment Proposal Review", "Propose only (no release)", "Treasury/AP", "Backlog", "Proposal accuracy; hold correctness", 0),
    ("A13", "AP Close Agent", "Observe→Draft", "Close Owner", "Backlog", "Close task completion; accrual quality", 0),
    ("A14", "AP Reporting Agent", "Observe→Draft", "FP&A Partner", "Backlog", "Report freshness; rework on packs", 0),
    ("A15", "Root Cause Agent", "Observe→Draft", "Continuous Improvement", "Backlog", "Recurring defect reduction", 0),
    ("A16", "AP Manager Orchestrator", "Observe→Recommend", "AP Manager", "Backlog", "Queue balance; escalation hygiene", 0),
]


def build_registry() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Agent Registry"
    brand_banner(ws, "ER Agent Registry (A01–A16)", "Autonomy, ownership, KPIs, cost monitoring")

    headers = [
        "ID", "Agent", "Autonomy level (earn-up)", "Owner", "Status", "Primary KPIs",
        "Monthly run cost (USD) ILLUSTRATIVE", "Cost owner", "Kill-switch defined?",
        "Evidence standard link", "Last stage review", "Notes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, len(headers))

    for r, a in enumerate(AGENTS, start=7):
        vals = [
            a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[3],
            "Y" if a[4] in ("Pilot", "Shadow") else None,
            f"Evidence/{a[0]}", None, None,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin
            cell.font = font_body
            if c == 7:
                cell.fill = fill_illust
                cell.number_format = '"$"#,##0'
                cell.font = font_amber
        ws.cell(row=r, column=1).font = font_label

    ws["A24"] = "Monthly run costs are ILLUSTRATIVE placeholders for monitoring structure — not vendor quotes."
    ws["A24"].font = font_amber
    ws["A24"].fill = fill_illust
    ws.merge_cells("A24:L24")

    cost = wb.create_sheet("Cost Monitoring")
    brand_banner(cost, "Cost Monitoring Rollup", "Formula sum of registry illustrative costs")
    cost["A6"] = "Total illustrative monthly agent run cost"
    cost["B6"] = "=SUM('Agent Registry'!G7:G22)"
    cost["B6"].number_format = '"$"#,##0'
    cost["A7"] = "Annualised"
    cost["B7"] = "=B6*12"
    cost["B7"].number_format = '"$"#,##0'
    cost["A9"] = "Replace ILLUSTRATIVE costs with actual platform + labour allocation before board packs."
    cost["A9"].font = font_body

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws)
    path = OUT_DIR / "ER_Agent_Registry.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 4. Exception Tracker
# ═══════════════════════════════════════════════════════════════

TAXONOMY = [
    ("EX-PRICE", "Price variance", "Match", "Buyer", "P1"),
    ("EX-QTY", "Quantity variance", "Match", "Receiving", "P1"),
    ("EX-PO-MISS", "Missing / invalid PO", "Validation", "Requester", "P1"),
    ("EX-GR-MISS", "Missing goods receipt", "GR", "Receiving", "P1"),
    ("EX-TAX", "Tax / VAT mismatch", "Validation", "Tax Desk", "P2"),
    ("EX-VENDOR", "Vendor identity / remit mismatch", "Master data", "Vendor Master", "P1"),
    ("EX-DUP", "Duplicate suspect", "Anomaly", "Controls", "P0"),
    ("EX-BANK", "Bank detail change related", "Master data", "Treasury/MDM", "P0"),
    ("EX-APPROVAL", "Approval matrix conflict", "Approval", "AP Manager", "P1"),
    ("EX-NONPO", "Non-PO policy breach", "Policy", "Cost Centre Owner", "P2"),
    ("EX-CURR", "Currency / FX issue", "Validation", "AP Ops", "P2"),
    ("EX-CREDIT", "Credit / debit memo mismatch", "Match", "AP Accountant", "P2"),
    ("EX-CONTRACT", "Contract / rate card conflict", "Procurement", "Category Manager", "P2"),
    ("EX-EARLY", "Early invoice / date policy", "Policy", "AP Ops", "P3"),
    ("EX-OTHER", "Unclassified — requires taxonomy update", "Triage", "Exception Lead", "P2"),
]


def build_exceptions() -> Path:
    wb = Workbook()
    tax = wb.active
    tax.title = "Taxonomy"
    brand_banner(tax, "Exception Taxonomy Codes", "Stable codes for triage agents & humans")
    headers = ["Code", "Description", "Family", "Default owner role", "Default priority"]
    for i, h in enumerate(headers, 1):
        tax.cell(row=6, column=i, value=h)
    style_header_row(tax, 6, 5)
    for r, row in enumerate(TAXONOMY, start=7):
        for c, v in enumerate(row, 1):
            cell = tax.cell(row=r, column=c, value=v)
            cell.border = thin
            cell.font = font_body
        tax.cell(row=r, column=1).font = font_label
    autosize(tax)

    tr = wb.create_sheet("Tracker")
    brand_banner(tr, "Exception Tracker", "Working queue columns for live use")
    th = [
        "Case ID", "Invoice #", "Vendor", "Amount", "Taxonomy code", "Opened",
        "Owner", "Priority", "SLA due", "Status", "Agent assist?", "Root cause",
        "Resolution", "Closed", "Evidence link", "Notes",
    ]
    for i, h in enumerate(th, 1):
        tr.cell(row=6, column=i, value=h)
    style_header_row(tr, 6, len(th))

    samples = [
        ("ER-2026-001", "INV-88921", "Acme Industrial", 12450.00, "EX-PRICE", "2026-03-01", "Buyer-12", "P1", "2026-03-03", "Open", "Y", "", "", "", "Evidence/EX/001", "ILLUSTRATIVE"),
        ("ER-2026-002", "INV-88955", "Northwind Paper", 820.40, "EX-GR-MISS", "2026-03-01", "WH-East", "P1", "2026-03-02", "In progress", "Y", "", "", "", "Evidence/EX/002", "ILLUSTRATIVE"),
        ("ER-2026-003", "INV-89002", "Globex Labs", 56000.00, "EX-DUP", "2026-03-02", "Controls", "P0", "2026-03-02", "Hold", "Y", "", "", "", "Evidence/EX/003", "ILLUSTRATIVE"),
        ("ER-2026-004", "INV-89018", "Initech Services", 3400.00, "EX-APPROVAL", "2026-03-02", "AP Manager", "P1", "2026-03-04", "Open", "N", "", "", "", "", "ILLUSTRATIVE"),
        ("ER-2026-005", "INV-89044", "Umbrella Parts", 199.99, "EX-TAX", "2026-03-03", "Tax Desk", "P2", "2026-03-06", "Open", "Y", "", "", "", "Evidence/EX/005", "ILLUSTRATIVE"),
    ]
    for r, row in enumerate(samples, start=7):
        for c, v in enumerate(row, 1):
            cell = tr.cell(row=r, column=c, value=v if v != "" else None)
            cell.border = thin
            cell.font = font_body
            cell.fill = fill_illust
            if c == 4:
                cell.number_format = '"$"#,##0.00'
    tr["A13"] = "Rows 7–11 are ILLUSTRATIVE sample tickets. Add real cases below."
    tr["A13"].font = font_amber
    tr.merge_cells("A13:P13")
    autosize(tr)

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    path = OUT_DIR / "ER_Exception_Tracker.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 5. Controls Matrix
# ═══════════════════════════════════════════════════════════════

CONTROLS = [
    ("A01 Intake", "Misclassified / incomplete invoice enters queue", "Completeness gate + confidence threshold", "Preventive", "AP Ops Lead", "Source hash; field confidence; human edit log", "Per invoice", "AP Manager if confidence threshold bypassed"),
    ("A02 Validation", "Invalid tax/vendor data posted", "Fail-closed validation checklist", "Preventive", "Data Steward", "Rule version; failed checks", "Per invoice", "Steward + Controls on master conflict"),
    ("A03 Matching", "False auto-match beyond tolerance", "Tolerance policy + sampled QA", "Detective", "Matching Lead", "Line compare; policy version", "Daily sample", "Demote to Propose-only"),
    ("A04 Triage", "Wrong owner / priority", "Taxonomy accuracy sampling", "Detective", "Exception Lead", "Class features; override log", "Weekly sample", "Exception Lead review board"),
    ("A07 Approval", "SoD / self-approval", "Matrix enforcement + SoD test", "Preventive", "AP Manager", "Route path; matrix version", "Per approval + monthly test", "Internal Controls immediately"),
    ("A10 Duplicate", "Missed duplicate payment", "Signal + human investigate; no auto-void", "Detective", "Controls Lead", "Signal set; investigator notes", "Per flag + monthly precision", "Controller if confirmed duplicate paid"),
    ("A12 Pay Proposal", "Agent releases payment", "Hard forbid: propose only; no release authority", "Preventive", "Treasury", "Proposal packet; human release ID", "Per payment run", "Kill-switch + security incident"),
    ("Vendor bank change", "Fraudulent remit change", "Dual control / call-back independent of agent", "Preventive", "Master Data", "Dual approver IDs; call-back log", "Per change", "Treasury + Fraud response"),
    ("All agents", "Unlogged model/prompt change", "Change control + version pin in evidence", "Preventive", "Governance Board", "Change ticket; model hash", "Per release", "Board demotion of agent"),
    ("All agents", "PII leakage to external LLM", "Data minimisation + approved tool list", "Preventive", "InfoSec + AP Manager", "Tool allowlist; DLP sample", "Monthly", "InfoSec incident process"),
]


def build_controls() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Controls Matrix"
    brand_banner(ws, "ER Controls Matrix", "Agent | Risk | Control | Type | Owner | Evidence | Frequency | Escalation")
    headers = ["Agent / scope", "Risk", "Control", "Type", "Owner", "Evidence", "Frequency", "Escalation"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, 8)
    for r, row in enumerate(CONTROLS, start=7):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin
            cell.font = font_body
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 48
    autosize(ws, max_w=36)

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    path = OUT_DIR / "ER_Controls_Matrix.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 6. Implementation Roadmap
# ═══════════════════════════════════════════════════════════════

PHASES = [
    (0, "Foundation — readiness & baseline KPIs", "Controller / AP Manager", "2026-01-06", "2026-01-31", "Complete", "Free kit + KPI worksheet"),
    (1, "Operating model & exception taxonomy", "AP Manager", "2026-02-01", "2026-02-21", "Complete", "Starter pack"),
    (2, "A01 Intake shadow → draft", "AP Ops Lead", "2026-02-22", "2026-03-21", "In progress", "Pilot charter"),
    (3, "A04 Triage + A03 Match propose", "Exception / Matching Leads", "2026-03-22", "2026-04-30", "Planned", "Controls matrix live"),
    (4, "A10 Duplicate signals", "Controls Lead", "2026-05-01", "2026-05-31", "Planned", "Precision sampling"),
    (5, "A07 Approval propose (narrow)", "AP Manager", "2026-06-01", "2026-06-30", "Planned", "SoD tests"),
    (6, "A05/A06 GR & PO quality", "Receiving / Procurement", "2026-07-01", "2026-07-31", "Planned", "Preventable defects"),
    (7, "A08/A09 Resolution & follow-up", "Vendor Desk / Ops", "2026-08-01", "2026-08-31", "Planned", "Template approval"),
    (8, "Bounded execute for clean match class", "Governance Board", "2026-09-01", "2026-09-30", "Planned", "Earn-in review"),
    (9, "A11–A14 close & reporting drafts", "Close / FP&A", "2026-10-01", "2026-11-15", "Backlog", "Pro library"),
    (10, "Scale, Team change pack, Custom edges", "AP Manager + Sponsor", "2026-11-16", "2026-12-31", "Backlog", "Team / Custom"),
]


def build_roadmap() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Roadmap"
    brand_banner(ws, "ER Implementation Roadmap — Phases 0–10", "Gantt-like dates, owners, status")
    headers = ["Phase", "Workstream", "Owner", "Start", "End", "Status", "Exit criteria / pack", "Duration (days)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, 8)
    for r, p in enumerate(PHASES, start=7):
        for c, v in enumerate(p, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin
            cell.font = font_body
            if c in (4, 5):
                cell.number_format = "YYYY-MM-DD"
        cell = ws.cell(row=r, column=8, value=f"=E{r}-D{r}")
        cell.border = thin
        cell.number_format = "0"

    g = wb.create_sheet("Gantt")
    brand_banner(g, "Gantt (month markers)", "X = active month — illustrative schedule")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    g["A6"] = "Phase"
    for i, m in enumerate(months, start=2):
        cell = g.cell(row=6, column=i, value=m)
        cell.fill = fill_ink
        cell.font = font_h
    active = {
        0: [1], 1: [2], 2: [2, 3], 3: [3, 4], 4: [5], 5: [6],
        6: [7], 7: [8], 8: [9], 9: [10, 11], 10: [11, 12],
    }
    for r, p in enumerate(PHASES, start=7):
        g.cell(row=r, column=1, value=f"P{p[0]} {p[1][:40]}").font = font_label
        for m in range(1, 13):
            on = m in active[p[0]]
            cell = g.cell(row=r, column=m + 1, value="█" if on else "")
            cell.alignment = Alignment(horizontal="center")
            if on:
                cell.fill = fill_teal
                cell.font = Font(color=WHITE)
            cell.border = thin
    g["A19"] = "Schedule is ILLUSTRATIVE for planning structure — adjust to your capacity."
    g["A19"].font = font_amber
    g["A19"].fill = fill_illust

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws)
    path = OUT_DIR / "ER_Implementation_Roadmap.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 7. Benefits Tracker
# ═══════════════════════════════════════════════════════════════

def build_benefits() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Benefits Tracker"
    brand_banner(ws, "ER Benefits Tracker", "Baseline → current → variance — formulas")
    headers = [
        "Benefit lever", "Baseline", "Current", "Variance", "Unit", "Owner",
        "Linked agent(s)", "Evidence", "Confidence (H/M/L)", "Notes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, 10)

    levers = [
        ("Cost per invoice", 12.5, 11.8, "USD", "Controller", "A01–A04", "KPI extract", "M"),
        ("Exception rate", 0.24, 0.21, "%", "Exception Lead", "A04", "Queue report", "M"),
        ("STP rate", 0.28, 0.31, "%", "AP Manager", "A01,A03", "Workflow", "L"),
        ("Cycle time median", 5.1, 4.2, "days", "AP Manager", "A03,A07,A09", "Timestamp", "M"),
        ("Late-payment costs (run-rate)", 120000, 98000, "USD/yr", "Controller", "A07,A09", "GL", "L"),
        ("Early-pay discounts captured", 40000, 52000, "USD/yr", "AP Manager", "A12", "Treasury", "L"),
        ("Duplicate losses avoided", 0, 15000, "USD/yr", "Controls", "A10", "Investigation log", "L"),
        ("FTE hours redirected", 0, 0.6, "FTE equiv", "AP Ops Lead", "A04,A09", "Time study", "L"),
    ]
    for r, row in enumerate(levers, start=7):
        ws.cell(row=r, column=1, value=row[0]).font = font_label
        b = ws.cell(row=r, column=2, value=row[1])
        c = ws.cell(row=r, column=3, value=row[2])
        v = ws.cell(row=r, column=4, value=f"=C{r}-B{r}")
        for cell in (b, c, v):
            cell.border = thin
            cell.fill = fill_illust
            if row[3] == "%":
                cell.number_format = "0.0%"
            elif row[3].startswith("USD"):
                cell.number_format = '"$"#,##0.00'
            else:
                cell.number_format = "0.00"
        for col, val in enumerate(row[3:], start=5):
            ws.cell(row=r, column=col, value=val).border = thin
        ws.cell(row=r, column=10, value="ILLUSTRATIVE sample").border = thin
        ws.cell(row=r, column=10).font = font_amber

    ws["A16"] = "Baseline/Current values are ILLUSTRATIVE. Variance uses formulas. Not a savings guarantee."
    ws["A16"].fill = fill_illust
    ws["A16"].font = font_amber
    ws.merge_cells("A16:J16")

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws)
    path = OUT_DIR / "ER_Benefits_Tracker.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 8. Maturity Diagnostic
# ═══════════════════════════════════════════════════════════════

QUESTIONS = [
    ("Process", "P1", "Invoice intake channels are documented with an owner for each channel."),
    ("Process", "P2", "Written standard for invoice completeness before AP queue entry."),
    ("Process", "P3", "PO and non-PO paths separated with different handling rules."),
    ("Process", "P4", "Exception types catalogued with defined owners."),
    ("Process", "P5", "Approval matrices current and applied consistently."),
    ("Process", "P6", "Payment runs follow published calendar with cut-offs and holds."),
    ("Process", "P7", "Period-close AP tasks are checklist-driven."),
    ("Data", "D1", "Vendor master payment fields defined and enforced."),
    ("Data", "D2", "Duplicate vendor records reviewed on a schedule."),
    ("Data", "D3", "Invoice fields map cleanly to ERP without habitual free-text."),
    ("Data", "D4", "PO, GR, and invoice identifiers join reliably for matching."),
    ("Data", "D5", "≥12 months historical invoices queryable for patterns."),
    ("Data", "D6", "Data quality issues logged with remediation owners."),
    ("Controls", "C1", "SoD enforced: vendor setup vs entry vs payment release."),
    ("Controls", "C2", "Vendor bank changes require dual control / verification."),
    ("Controls", "C3", "Audit trails for approvals, matches, overrides, holds."),
    ("Controls", "C4", "Force-pay authority limited, logged, reviewed."),
    ("Controls", "C5", "AP access role-based and reviewed at least annually."),
    ("Controls", "C6", "Sample testing / monitoring on high-risk invoice classes."),
    ("Controls", "C7", "Policies for discounts, late fees, hold releases written & tested."),
    ("Technology", "T1", "System of record (ERP/AP) — not spreadsheets as primary path."),
    ("Technology", "T2", "Capture/OCR or structured intake for meaningful volume share."),
    ("Technology", "T3", "Workflow routing/reminders/escalation without email-only chase."),
    ("Technology", "T4", "Integrations AP–procurement–receiving stable for matching."),
    ("Technology", "T5", "Reporting extracts without heroic manual effort."),
    ("Technology", "T6", "Sandbox suitable for agent pilot testing."),
    ("People", "H1", "Roles for processors, exception owners, approvers clear & staffed."),
    ("People", "H2", "Staff can explain approvals/holds using system evidence."),
    ("People", "H3", "Capacity for process redesign, not only firefighting."),
    ("People", "H4", "Leadership: agents under human accountability."),
    ("People", "H5", "Training for tools/policy; adoption measured."),
    ("Measurement", "M1", "Cost per invoice (or proxy) tracked at least quarterly."),
    ("Measurement", "M2", "Exception rate and STP measured with agreed definitions."),
    ("Measurement", "M3", "Cycle time receipt→payment-ready by invoice class."),
    ("Measurement", "M4", "KPI reviews on fixed cadence with owners and actions."),
]


def build_maturity() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Diagnostic"
    brand_banner(ws, "ER Maturity Diagnostic — 35 questions", "Score 0–4 each; maturity auto-calcs")
    ws["A5"] = "Scoring: 0 Absent · 1 Ad hoc · 2 Emerging · 3 Established · 4 Managed"
    ws["A5"].font = font_body

    headers = ["Dimension", "ID", "Statement", "Score (0–4)", "Evidence / notes"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=6, column=i, value=h)
    style_header_row(ws, 6, 5)

    sample_scores = [
        3, 2, 3, 2, 3, 2, 2,
        3, 2, 2, 3, 2, 2,
        3, 3, 3, 2, 3, 2, 2,
        3, 2, 2, 2, 2, 1,
        3, 2, 2, 3, 2,
        2, 2, 2, 2,
    ]

    for r, (dim, qid, stmt) in enumerate(QUESTIONS, start=7):
        ws.cell(row=r, column=1, value=dim).font = font_label
        ws.cell(row=r, column=2, value=qid).font = font_teal
        ws.cell(row=r, column=3, value=stmt).font = font_body
        score = ws.cell(row=r, column=4, value=sample_scores[r - 7])
        score.fill = fill_illust
        score.font = font_amber
        score.border = thin
        note = ws.cell(row=r, column=5, value="ILLUSTRATIVE sample score — replace")
        note.font = Font(name="Calibri", size=8, italic=True, color=RULE)
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = thin

    sum_ws = wb.create_sheet("Scoring Summary")
    brand_banner(sum_ws, "Auto Maturity Scoring", "Formulas over Diagnostic!D7:D41")
    dims = [
        ("Process", 7, 13, 28),
        ("Data", 14, 19, 24),
        ("Controls", 20, 26, 28),
        ("Technology", 27, 32, 24),
        ("People", 33, 37, 20),
        ("Measurement", 38, 41, 16),
    ]
    sum_ws["A6"] = "Dimension"
    sum_ws["B6"] = "Your points"
    sum_ws["C6"] = "Max"
    sum_ws["D6"] = "% of max"
    sum_ws["E6"] = "Band"
    style_header_row(sum_ws, 6, 5)
    for i, (name, start, end, mx) in enumerate(dims, start=7):
        sum_ws.cell(row=i, column=1, value=name).font = font_label
        sum_ws.cell(row=i, column=2, value=f"=SUM(Diagnostic!D{start}:D{end})")
        sum_ws.cell(row=i, column=3, value=mx)
        sum_ws.cell(row=i, column=4, value=f"=IF(C{i}=0,0,B{i}/C{i})")
        sum_ws.cell(row=i, column=4).number_format = "0.0%"
        sum_ws.cell(
            row=i, column=5,
            value=(
                f'=IF(D{i}<0.25,"Fragile",IF(D{i}<0.5,"Forming",'
                f'IF(D{i}<0.75,"Capable",IF(D{i}<0.9,"Strong","Evidence-ready"))))'
            ),
        )
        for c in range(1, 6):
            sum_ws.cell(row=i, column=c).border = thin

    sum_ws["A14"] = "Total points"
    sum_ws["B14"] = "=SUM(B7:B12)"
    sum_ws["C14"] = 140
    sum_ws["D14"] = "=B14/C14"
    sum_ws["D14"].number_format = "0.0%"
    sum_ws["A15"] = "Maturity level"
    sum_ws["B15"] = '=IF(B14<=27,1,IF(B14<=55,2,IF(B14<=83,3,IF(B14<=111,4,5))))'
    sum_ws["C15"] = (
        '=IF(B15=1,"Manual & opaque",IF(B15=2,"Documented but fragile",'
        'IF(B15=3,"Controlled operations",IF(B15=4,"Measurable & integrable",'
        '"Agent-ready with governance"))))'
    )
    sum_ws["A17"] = "Sample scores on Diagnostic sheet are ILLUSTRATIVE. Overwrite with evidence-based scores."
    sum_ws["A17"].font = font_amber
    sum_ws["A17"].fill = fill_illust

    levels = wb.create_sheet("Level Guide")
    brand_banner(levels, "Maturity Levels 1–5", "Agent posture by level")
    guide = [
        (1, "0–27", "Manual & opaque", "Observation only"),
        (2, "28–55", "Documented but fragile", "Draft-only, 100% human action"),
        (3, "56–83", "Controlled operations", "Propose-and-confirm"),
        (4, "84–111", "Measurable & integrable", "Bounded autonomy narrow classes"),
        (5, "112–140", "Agent-ready with governance", "Portfolio with RACI & sampling"),
    ]
    levels["A6"] = "Level"
    levels["B6"] = "Points"
    levels["C6"] = "Label"
    levels["D6"] = "Agent posture"
    style_header_row(levels, 6, 4)
    for r, row in enumerate(guide, start=7):
        for c, v in enumerate(row, 1):
            levels.cell(row=r, column=c, value=v).border = thin

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws, max_w=55)
    autosize(sum_ws)
    path = OUT_DIR / "ER_Maturity_Diagnostic.xlsx"
    wb.save(path)
    return path


# ═══════════════════════════════════════════════════════════════
# 9. Business Case Model
# ═══════════════════════════════════════════════════════════════

def build_business_case() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    brand_banner(ws, "ER Business Case Model", "Board-ready structure — illustrative until you own the inputs")

    assumptions = [
        ("A1", "Annual invoice volume", 120000, "invoices"),
        ("A2", "Current cost per invoice", 12.5, "USD"),
        ("A3", "Target cost per invoice (internal goal)", 10.0, "USD — internal goal, not industry promise"),
        ("A4", "AP loaded labour pool", 680000, "USD/year"),
        ("A5", "One-time implementation", 75000, "USD"),
        ("A6", "Annual platform / AI run cost", 48000, "USD"),
        ("A7", "Change / training cost (Y1)", 25000, "USD"),
        ("A8", "Benefit realisation ramp Y1", 0.5, "fraction"),
        ("A9", "Benefit realisation ramp Y2", 0.85, "fraction"),
        ("A10", "Benefit realisation ramp Y3", 1.0, "fraction"),
        ("A11", "Working capital benefit (annual, optional)", 40000, "USD — optional"),
        ("A12", "Risk reduction value (qualitative $)", 20000, "USD — conservative proxy"),
        ("A13", "Discount rate", 0.08, "for NPV"),
        ("A14", "Horizon years", 3, "years"),
    ]
    ws["A6"] = "ID"
    ws["B6"] = "Assumption"
    ws["C6"] = "Value"
    ws["D6"] = "Unit / note"
    style_header_row(ws, 6, 4)
    for r, (i, name, val, note) in enumerate(assumptions, start=7):
        ws.cell(row=r, column=1, value=i).font = font_teal
        ws.cell(row=r, column=2, value=name).font = font_label
        cell = ws.cell(row=r, column=3, value=val)
        cell.fill = fill_input
        cell.border = thin
        if isinstance(val, float) and val <= 1:
            cell.number_format = "0.0%"
        else:
            cell.number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=note).font = font_body

    cf = wb.create_sheet("Cashflow")
    brand_banner(cf, "3-Year Cashflow (formulas)", "ILLUSTRATIVE until Assumptions are yours")
    cf["A6"] = "Line"
    cf["B6"] = "Year 1"
    cf["C6"] = "Year 2"
    cf["D6"] = "Year 3"
    style_header_row(cf, 6, 4)

    # Row map: C7=vol, C8=current$/inv, C9=target, C11=impl, C12=platform, C13=training,
    # C14=Y1 ramp, C15=Y2, C16=Y3, C17=WC, C18=risk, C19=discount
    lines = [
        ("Gross cost/invoice savings", "=Assumptions!C7*(Assumptions!C8-Assumptions!C9)*Assumptions!C14", "=Assumptions!C7*(Assumptions!C8-Assumptions!C9)*Assumptions!C15", "=Assumptions!C7*(Assumptions!C8-Assumptions!C9)*Assumptions!C16"),
        ("Working capital benefit", "=Assumptions!C17*Assumptions!C14", "=Assumptions!C17*Assumptions!C15", "=Assumptions!C17*Assumptions!C16"),
        ("Risk reduction proxy", "=Assumptions!C18*Assumptions!C14", "=Assumptions!C18*Assumptions!C15", "=Assumptions!C18*Assumptions!C16"),
        ("Total benefits", "=B7+B8+B9", "=C7+C8+C9", "=D7+D8+D9"),
        ("Platform / AI run cost", "=Assumptions!C12", "=Assumptions!C12", "=Assumptions!C12"),
        ("Implementation (one-time)", "=Assumptions!C11", 0, 0),
        ("Change / training", "=Assumptions!C13", 0, 0),
        ("Total costs", "=B11+B12+B13", "=C11+C12+C13", "=D11+D12+D13"),
        ("Net cashflow", "=B10-B14", "=C10-C14", "=D10-D14"),
    ]
    for r, (name, y1, y2, y3) in enumerate(lines, start=7):
        cf.cell(row=r, column=1, value=name).font = font_label
        for col, f in enumerate((y1, y2, y3), start=2):
            cell = cf.cell(row=r, column=col, value=f)
            cell.number_format = '"$"#,##0'
            cell.border = thin

    cf["A17"] = "Cumulative"
    cf["B17"] = "=B15"
    cf["C17"] = "=B17+C15"
    cf["D17"] = "=C17+D15"
    for col in range(2, 5):
        cf.cell(row=17, column=col).number_format = '"$"#,##0'
        cf.cell(row=17, column=col).font = font_teal

    cf["A19"] = "NPV (3yr)"
    cf["B19"] = "=B15/(1+Assumptions!C19)+C15/(1+Assumptions!C19)^2+D15/(1+Assumptions!C19)^3"
    cf["B19"].number_format = '"$"#,##0'
    cf["A20"] = "Simple ROI (3yr benefits / 3yr costs)"
    cf["B20"] = "=(B10+C10+D10)/(B14+C14+D14)"
    cf["B20"].number_format = "0.0%"

    cf["A22"] = "ILLUSTRATIVE MODEL — industry averages are external context on Notes only, not embedded targets."
    cf["A22"].font = font_amber
    cf["A22"].fill = fill_illust

    exec_s = wb.create_sheet("Exec Summary")
    brand_banner(exec_s, "Executive Summary Fields", "Narrative hooks for board pack")
    exec_s["A6"] = "Problem statement"
    exec_s["B6"] = "AP lacks a governed agent operating model; exceptions and evidence are uneven."
    exec_s["A7"] = "Proposed approach"
    exec_s["B7"] = "Evidence Room earn-up: Observe → Recommend → Draft → Execute (narrow) → Expand."
    exec_s["A8"] = "Investment asked"
    exec_s["B8"] = "Y1 implementation + change costs; ongoing platform run cost (see Assumptions)."
    exec_s["A9"] = "Decision requested"
    exec_s["B9"] = "Approve Phase 0–3 pilot with kill-switches; revisit bounded execute after earn-in."
    exec_s["A10"] = "Non-goals"
    exec_s["B10"] = "No guaranteed cost/invoice to industry average; no unsupervised payment release; no fraud certainty claims."

    notes = wb.create_sheet("Notes")
    notes_ardent(notes)
    disc = wb.create_sheet("Disclaimer")
    disclaimer_sheet(disc)
    autosize(ws)
    autosize(cf)
    path = OUT_DIR / "ER_Business_Case_Model.xlsx"
    wb.save(path)
    return path


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    builders = [
        build_roi,
        build_kpi,
        build_registry,
        build_exceptions,
        build_controls,
        build_roadmap,
        build_benefits,
        build_maturity,
        build_business_case,
    ]
    print(f"Writing workbooks to {OUT_DIR}")
    for fn in builders:
        path = fn()
        print(f"  OK {path.name} ({path.stat().st_size:,} bytes)")
    print("Done.")


if __name__ == "__main__":
    main()
