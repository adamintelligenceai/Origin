#!/usr/bin/env python3
"""Evidence Room executive decks — paper, ink, bronze. No purple AI."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from lxml import etree
from copy import deepcopy

INK = RGBColor(0x1A, 0x1C, 0x19)
PAPER = RGBColor(0xF6, 0xF3, 0xEC)
MARK = RGBColor(0x8B, 0x69, 0x14)
PERMIT = RGBColor(0x2F, 0x4A, 0x3C)
HOLD = RGBColor(0x8A, 0x3A, 0x2A)
SLATE = RGBColor(0x5C, 0x61, 0x58)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
RULE = RGBColor(0xD4, 0xCF, 0xC3)

W, H = Inches(13.333), Inches(7.5)
ROOT = Path("/workspace/EVIDENCE_ROOM")


def set_run(run, text, size=18, bold=False, color=INK, italic=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=PAPER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_rule(slide, l, t, w, h=Inches(0.03), color=MARK):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run() if p.runs else p.runs[0] if False else p.add_run(), text, size, bold, color, italic)
    # python-pptx: empty paragraph has no runs until we add
    return box


def title_bar(slide, kicker, title, sub=None):
    add_bg(slide)
    tb(slide, Inches(0.7), Inches(0.28), Inches(8), Inches(0.3), "EVIDENCE ROOM  ·  AP AGENT OS", 11, True, MARK)
    add_rule(slide, Inches(0.7), Inches(0.62), Inches(2.2))
    tb(slide, Inches(0.7), Inches(0.8), Inches(12), Inches(0.7), title, 28, True, INK)
    if sub:
        tb(slide, Inches(0.7), Inches(1.45), Inches(12), Inches(0.4), sub, 14, False, SLATE)
    tb(slide, Inches(9.6), Inches(0.28), Inches(3.2), Inches(0.3), kicker, 11, True, SLATE, align=PP_ALIGN.RIGHT)


def bullets(slide, items, top=2.0, left=0.7, width=12.0, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, "—  " + item, size, False, INK)
    return box


def cover(prs, kicker, title, sub, footer="Proof before permission."):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, INK)
    add_rule(slide, Inches(0.8), Inches(2.15), Inches(1.6), Inches(0.04), MARK)
    tb(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(0.35), kicker.upper(), 13, True, MARK)
    # title may be multiline
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(2.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, title, 40, True, WHITE)
    tb(slide, Inches(0.8), Inches(5.1), Inches(11), Inches(0.8), sub, 18, False, RGBColor(0xD4, 0xCF, 0xC3))
    tb(slide, Inches(0.8), Inches(6.8), Inches(8), Inches(0.3), footer, 12, False, MARK)
    tb(slide, Inches(9.2), Inches(6.8), Inches(3.3), Inches(0.3), "evidenceroom.ai", 12, False, RGBColor(0xD4, 0xCF, 0xC3), align=PP_ALIGN.RIGHT)
    return slide


def closer(prs, line="Assess your AP Agent readiness."):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, INK)
    tb(slide, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.4), line, 32, True, WHITE)
    tb(slide, Inches(0.8), Inches(4.0), Inches(11), Inches(0.5), "Primary: evidenceroom.ai/diagnostic    ·    Secondary: Explore the AP Agent OS", 16, False, RGBColor(0xD4, 0xCF, 0xC3))
    tb(slide, Inches(0.8), Inches(6.5), Inches(11), Inches(0.4), "Not a savings, fraud, compliance, or autonomous-payment guarantee. Counsel reviews licence and privacy before force.", 12, False, MARK)
    return slide


def cfo_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cover(prs, "CFO briefing", "Proof before permission.", "How finance leadership governs an AP agent layer — without replacing the stack or releasing cash to a model.")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "THE ASK", "The question you will be asked", "Not ‘do we have AI?’ — ‘what is it allowed to do?’")
    bullets(s, [
        "Deloitte Finance Trends 2026 (8 Oct 2025): 63% of finance teams have fully deployed AI; 21% report clear, measurable ROI; 14% use fully integrated agents.",
        "The gap is operating design: named jobs, evidence packets, autonomy that starts at Observe.",
        "APQC (16 Mar 2026): top vs bottom AP cost per $1,000 revenue is about $0.38 vs $0.92 — process quality still moves money. Independent. Not our result.",
        "You do not need another capture vendor to start. You need a rule for permission.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "WHAT IT IS", "AP Agent OS", "An operating system for the agent layer. ERP-agnostic.")
    bullets(s, [
        "Sixteen named agents — intake through orchestrator — each with a charter, exclusions, and a human owner.",
        "Five autonomy levels. Default is L0 Observe. Promotion is earned per agent × entity × slice.",
        "Four classes that stay human: payment release, vendor bank-change, policy exceptions, legal disputes.",
        "It does not post your ledger, pay your suppliers, or replace SAP, Oracle, D365, NetSuite, or Workday.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "WORKFORCE", "The stack, not a chatbot")
    bullets(s, [
        "Capture: Intake.",
        "Qualify: Validation, Matching, Duplicate & Anomaly.",
        "Resolve: Exception triage, GR, PO quality, supplier and internal follow-up, statements.",
        "Control: Approval routing, payment-proposal review (human authorises the run).",
        "Steer: Close, reporting, root cause. Orchestrator sequences and stops — it does not match or pay.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "PERMISSION", "Responsibility is earned")
    bullets(s, [
        "L0 Observe — reviews, cannot act.",
        "L1 Recommend — advice for a human.",
        "L2 Prepare — drafts; send/post requires approval.",
        "L3 Execute within guardrails — pre-approved low-risk writes only.",
        "L4 Managed autonomy — bounded, sampled, kill-switched. Still no cash movement.",
        "A kickoff deck is not a promotion file.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "CONTROL", "What you can show Audit")
    bullets(s, [
        "Packet standard: case_id, evidence_refs, autonomy applied, human_required.",
        "Control matrix: risk, control, owner, evidence, frequency, escalation.",
        "Kill-switch and fallback SOP if the model or vendor is dark.",
        "Unsigned ‘savings’ do not enter this pack.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "ECONOMICS", "How to talk about money without lying")
    bullets(s, [
        "Hackett (19 Nov 2025) found software-adopter averages of 60% touchless and 3.5× productivity at ≥30% touchless. That is their research on AP solutions — not Evidence Room performance.",
        "Your case uses your volume, your exception minutes, your loaded cost.",
        "Conservative / Base / Upside. Capacity is not cash until you sign it.",
        "If Conservative net is negative, that is information.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "90 DAYS", "One bounded agent — not the estate")
    bullets(s, [
        "Days 1–15: diagnostic, owners, slice, observe as-done work.",
        "Days 16–35: charter at L0, access, gold-first historical test.",
        "Days 36–50: shadow. No action permissions.",
        "Days 51–70: controlled pilot only if evidence. Allow-list. At-most-once writes.",
        "Days 71–90: measure vs baseline. Promote one increment or stay.",
        "Actual duration depends on data, integrations, DoA, and governance. State that every time.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "DECISION", "What we are asking")
    bullets(s, [
        "Appoint a human owner for Intake and for Duplicate & Anomaly.",
        "Name one legal entity and one channel as the first slice.",
        "Approve L0 shadow budget (time, access, privacy review) — not L3.",
        "Refuse any success criterion that is ‘autonomous payment release’.",
    ])
    closer(prs)
    out = ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "ER_CFO_AP_Transformation.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def workshop_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cover(prs, "Workshop  ·  Team Edition", "Design the agent layer.", "A one-day working session. Not a demo day. Not a prompt bake-off.")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "PURPOSE", "By 16:30 we will have")
    bullets(s, [
        "A scored readiness snapshot (even if incomplete).",
        "One first-slice definition: entity × channel × invoice type.",
        "Two named human owners (Intake + Duplicate).",
        "A draft charter at L0 — or a written reason we will not start.",
        "A parking lot for ‘autonomous payments’ and other out-of-scope asks.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "AGENDA", "One day")
    bullets(s, [
        "09:00  Proof before permission — the rule, not the product tour.",
        "09:30  As-done AP: channels, exceptions, who actually touches what.",
        "10:30  Taxonomy: map your reason codes; do not invent a parallel list.",
        "11:15  Human vs agent: the four classes that stay human.",
        "12:30  Lunch. No vendor pitches.",
        "13:15  First-slice vote. Heatmap if needed.",
        "14:00  Charter clinic (Intake L0 + Duplicate companion).",
        "15:15  Controls and kill-switch.",
        "15:45  30-day actions. Owners in the room.",
        "16:15  What we will not do.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "RULES OF THE ROOM", "How we work today")
    bullets(s, [
        "No live invoices with personal data on the projector unless already approved.",
        "No scoring hope. Score what you can show.",
        "Agents are never Accountable on the RACI.",
        "If someone says ‘just let it pay under $500’, write it on the parking lot and move on.",
        "Facilitator keeps the clock. Transformation lead owns the outcome.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "EXERCISE", "Human vs agent")
    bullets(s, [
        "For each activity, mark: Human only / Agent recommend / Agent prepare / Agent execute (guardrails).",
        "Payment release — Human only. Non-negotiable in this OS.",
        "Vendor bank-change — Human only.",
        "Draft a missing-PO email — Agent prepare, send-gate human (default).",
        "Classify EX-GR-001 — Agent recommend or prepare; posting a GR is human / ERP.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "CHARTER", "If it is not in the charter, it is not in scope")
    bullets(s, [
        "Purpose, exclusions, inputs, tools, output standard.",
        "L0–L4 decision rights. Today we fill L0 and L1 only.",
        "Failure handling and kill-switch.",
        "KPIs with formulas — n_min before we publish a percentage.",
        "Instruction skeleton is a starting operating instruction, not a magic prompt.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "CLOSE", "Tomorrow morning")
    bullets(s, [
        "Owner books observation slots (as-done, not a conference room reconstruction).",
        "Finance Systems opens a read-path ticket — not a write path.",
        "Privacy note started if images will leave the tenant.",
        "Steering date booked. No unsigned savings on that slide.",
    ])
    closer(prs, "Run the session. Then measure. Then ask for permission.")
    out = ROOT / "04_AP_AGENT_OS_TEAM" / "Workshop" / "ER_AP_Workshop_Deck.pptx"
    prs.save(out)
    return out


def steering_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cover(prs, "Steering committee", "Evidence, then a decision.", "A standing pack. Section 0 is actions. Unsigned savings are omitted.")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "SECTION 0", "Decisions this sitting")
    bullets(s, [
        "Stay / promote / demote / kill — for the named agent × entity × slice.",
        "Access or privacy blockers — who clears them.",
        "Any EX-PAY-001 (banking-change) items — listed even if count is zero.",
        "Do not add a ‘strategic AI’ slide in front of this one.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "SCOPE", "What is live")
    bullets(s, [
        "Agent, entity, channel, invoice type, current level, ceiling.",
        "Human owner present or represented.",
        "n for the period. If n < n_min, show counts only.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "RISK", "Risk-control can veto promotion")
    bullets(s, [
        "False-negative on duplicates. Control breaches. Overrides.",
        "Prompt-injection or tool-allow-list incidents.",
        "Applied autonomy > registered ceiling — treat as an incident.",
        "A pretty operational chart does not override a risk-control fail.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "ECONOMICS", "Two columns")
    bullets(s, [
        "Modelled (from the workbook) — labelled MODELLED.",
        "Validated (Controller-signed) — labelled VALIDATED.",
        "If the second column is empty, say so. Do not average them.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "ASK", "One increment or stop")
    bullets(s, [
        "Promotion requires the evidence pack in Autonomy Progression.",
        "Two signatures for L3+ (owner + Controller or Controls).",
        "Sampling plan and rollback written before the vote.",
        "If evidence is thin, the correct decision is stay.",
    ])
    closer(prs, "No proof, no promotion.")
    out = ROOT / "04_AP_AGENT_OS_TEAM" / "Executive" / "ER_Steering_Committee_Pack.pptx"
    prs.save(out)
    return out


def business_case_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cover(prs, "Business case", "A model is not a result.", "Conservative / Base / Upside on your volume. Capacity is not cash.")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "WHY A CASE", "Because 21% is the number that matters")
    bullets(s, [
        "Deloitte (8 Oct 2025): 63% deployed AI; 21% report clear, measurable ROI.",
        "This pack exists so you do not become the 42-point gap.",
        "We will not put Hackett’s 60% touchless on a slide as if it were yours.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "INPUTS", "You must own these")
    bullets(s, [
        "Invoice volume, AP ops FTE, loaded cost.",
        "Manual-touch and exception rates from your queue.",
        "Resolution minutes. Duplicate rate if measured.",
        "Late fees and missed discounts only if known — else zero.",
        "AI/tool run-cost and implementation cost.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "OUTPUTS", "What the workbook returns")
    bullets(s, [
        "Baseline labour proxy and exception hours.",
        "Hours released after a realisation haircut.",
        "Cash-like benefit only after a conversion percentage you set.",
        "Net after AI run-cost. Payback. Year-1 ROI (may be negative).",
        "Sensitivity grid: exception rate × efficiency.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "HOW TO PRESENT", "Three rules")
    bullets(s, [
        "Lead with Conservative. Show Base. Mention Upside last.",
        "Say ‘hours’ before you say ‘dollars’.",
        "Bring the Benefits tracker: VALIDATED column starts at zero.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title_bar(s, "ILLUSTRATIVE ACME", "8,400 invoices / month — overwrite in the workbook")
    bullets(s, [
        "Figures in the Excel default are ILLUSTRATIVE ACME Manufacturing.",
        "They exist so the model is not empty. They are not a benchmark.",
        "If you present ACME numbers as your own, you have failed the evidence standard.",
    ])
    closer(prs, "Open the workbook. Replace the yellow cells. Then talk.")
    out = ROOT / "03_AP_AGENT_OS_PRO" / "Business_Case" / "ER_Business_Case_Deck.pptx"
    prs.save(out)
    return out


def main():
    for fn in (cfo_deck, workshop_deck, steering_deck, business_case_deck):
        print(fn())


if __name__ == "__main__":
    main()
