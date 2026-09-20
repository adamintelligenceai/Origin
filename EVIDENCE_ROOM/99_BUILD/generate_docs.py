#!/usr/bin/env python3
"""Editable Word templates — example + blank where needed."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path("/workspace/EVIDENCE_ROOM")
INK = RGBColor(0x1A, 0x1C, 0x19)
MARK = RGBColor(0x8B, 0x69, 0x14)
SLATE = RGBColor(0x5C, 0x61, 0x58)


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def add_p(doc, text, size=11, bold=False, color=INK, italic=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size, bold, color, italic)
    return p


def header_block(doc, title, subtitle):
    add_p(doc, "EVIDENCE ROOM  ·  AP AGENT OS", 10, True, MARK, space_after=2)
    add_p(doc, title, 22, True, INK, space_after=4)
    add_p(doc, subtitle, 11, False, SLATE, True, space_after=12)
    add_p(doc, "Proof before permission. Agents earn responsibility. Evidence decides.", 10, False, MARK, True)
    add_p(doc, "This template is independently authored. Do not paste employer or client confidential data. ACME examples are ILLUSTRATIVE.", 9, False, SLATE, True)


def field_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        t.rows[i].cells[0].text = k
        t.rows[i].cells[1].text = v
        for cell in t.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, 10)
    return t


def save(doc, rel):
    path = ROOT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def sop():
    doc = Document()
    header_block(doc, "SOP template", "Current-state or target-state. One procedure. Named owner.")
    add_p(doc, "1. Header", 14, True)
    field_table(doc, [
        ("SOP ID", "[BUYER]"),
        ("Title", "[procedure name]"),
        ("Process owner (human, Accountable)", "[name / role]"),
        ("ERP / channel", "[system — agnostic]"),
        ("Version / date", "[ ]"),
        ("Related agent (if any)", "AG-__ at L__  ·  ceiling L__"),
        ("Kill-switch", "Where it lives / who can pull it"),
    ])
    add_p(doc, "")
    add_p(doc, "2. Purpose — what this procedure exists to do", 14, True)
    add_p(doc, "[ ]")
    add_p(doc, "3. Scope / out of scope", 14, True)
    add_p(doc, "In: [ ]")
    add_p(doc, "Out: payment release, vendor bank-change, policy exceptions, legal disputes remain human even if an agent drafts.")
    add_p(doc, "4. Triggers and inputs", 14, True)
    add_p(doc, "[document types, systems, frequency]")
    add_p(doc, "5. Steps (as-done, not as-designed unless this is target-state)", 14, True)
    add_p(doc, "1. [actor] does [action] in [system] using [evidence]. Failure: [what happens].")
    add_p(doc, "2. …")
    add_p(doc, "6. Decision tree", 14, True)
    add_p(doc, "If [condition] then [route / code / owner]. Else [ ].")
    add_p(doc, "7. Controls and evidence to retain", 14, True)
    add_p(doc, "[IDs, retention, who can retrieve]")
    add_p(doc, "8. Escalation", 14, True)
    add_p(doc, "[clock, amount, repeat]")
    add_p(doc, "9. Metrics", 14, True)
    add_p(doc, "[KPI IDs from the dictionary — n_min applies]")
    add_p(doc, "ILLUSTRATIVE ACME fragment (delete in live SOP)", 14, True, MARK)
    add_p(doc, "SOP-ACME-GR-01 Missing receipt chase. Owner: Plant AP liaison. Agent AG-05 recommends receiver; never posts GR. Channel: email PDF. Entity: ACME-US. Clock: 2 business days.")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_SOP_Template.docx")


def charter():
    doc = Document()
    header_block(doc, "Agent charter", "If it is not in this file, it is not in scope.")
    field_table(doc, [
        ("Agent ID / name", "AG-__  /  [ ]"),
        ("Entity × slice", "[legal entity] × [channel] × [invoice type]"),
        ("Human owner (Accountable)", "[named person]"),
        ("Current level / ceiling", "L0 / L1  (default)"),
        ("Effective date", "[ ]"),
    ])
    for h in [
        "Purpose", "In scope", "Explicit exclusions", "Inputs", "Tools (allow-list)",
        "Outputs and output standard", "Decision rights L0–L4", "Approvals",
        "Escalation", "Controls and audit evidence", "Failure handling / kill-switch",
        "Cost envelope", "KPIs", "First-90-day slice", "Instruction skeleton (starting instruction, not a magic prompt)",
    ]:
        add_p(doc, h, 13, True)
        add_p(doc, "[ ]")
    add_p(doc, "Standing exclusions (do not delete)", 13, True, HOLD if False else MARK)
    add_p(doc, "Payment release, vendor bank-change approval, policy exceptions, and legal disputes remain human. Agent 12 never releases payment. Agent 10 is not a fraud opinion. Agent 05 does not invent a goods receipt.")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Agent_Charter.docx")


def discovery():
    doc = Document()
    header_block(doc, "Process discovery record", "Observe as-done work. Do not reconstruct it in a conference room and call it truth.")
    field_table(doc, [
        ("Session ID", "[ ]"),
        ("Date / site / channel", "[ ]"),
        ("Observed role(s)", "[ ]"),
        ("Consent / recording note", "[ ]"),
        ("Redaction complete", "Yes / No"),
    ])
    add_p(doc, "Cases observed (IDs only — no personal data in the shared pack)", 13, True)
    add_p(doc, "1. [case_id]  2. …")
    add_p(doc, "Extract", 13, True)
    for item in ["Steps", "Systems", "Decisions", "Business rules", "Inputs", "Outputs", "Exceptions", "Controls", "Dependencies"]:
        add_p(doc, f"{item}: [ ]")
    add_p(doc, "Quotes that must survive (required for any rule you write)", 13, True)
    add_p(doc, "“[quote]” — role, case_id")
    add_p(doc, "What we still do not know", 13, True)
    add_p(doc, "[ ]")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Process_Discovery.docx")


def raci():
    doc = Document()
    header_block(doc, "RACI", "Agents may be Consulted or Informed. Agents are never Accountable.")
    add_p(doc, "R = Responsible (does the work)  A = Accountable (one human)  C = Consulted  I = Informed", 10, False, SLATE, True)
    headers = ["Activity", "AP Manager", "Processor", "Controller", "FinSys", "Agent (C/I only)"]
    rows = [
        ["Approve L0 charter", "A", "C", "I", "C", "I"],
        ["Classify exception", "A", "R", "I", "I", "C"],
        ["Release payment", "I", "R (prepare)", "A / Treasury", "I", "— (forbidden)"],
        ["Approve vendor bank-change", "I", "I", "A / Treasury", "C", "— (forbidden)"],
        ["Promote autonomy +1", "R", "C", "A (L3+)", "C", "I"],
        ["Pull kill-switch", "A", "R", "I", "R", "I"],
    ]
    t = doc.add_table(rows=1 + len(rows), cols=6)
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            t.rows[r].cells[c].text = val
    add_p(doc, "")
    add_p(doc, "Blank rows", 13, True)
    t2 = doc.add_table(rows=8, cols=6)
    t2.style = "Table Grid"
    for i, h in enumerate(headers):
        t2.rows[0].cells[i].text = h
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_RACI.docx")


def uat():
    doc = Document()
    header_block(doc, "UAT record", "Gold first. Do not score the agent against a label created after seeing the output.")
    field_table(doc, [
        ("Test pack ID", "[ ]"),
        ("Agent / level under test", "AG-__  /  L0–L2 (typical)"),
        ("n / n_min", "[ ] / [ ]"),
        ("Gold created before model output", "Yes — required"),
        ("Environment", "Historical / Shadow / Pilot allow-list"),
    ])
    add_p(doc, "Cases", 13, True)
    t = doc.add_table(rows=6, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["case_id", "Gold", "Agent output", "Match Y/N", "Notes"]):
        t.rows[0].cells[i].text = h
    add_p(doc, "")
    add_p(doc, "Result", 13, True)
    add_p(doc, "Accuracy (if n ≥ n_min): [ ]  |  Count-only: [ ]  |  Incidents: [ ]  |  Decision: stay / fix / promote increment")
    add_p(doc, "Sign-off (human owner)", 13, True)
    add_p(doc, "Name / date: [ ]")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_UAT.docx")


def risk():
    doc = Document()
    header_block(doc, "Risk assessment (agent slice)", "Buyer scores. Do not paste a generic AI-risk essay.")
    field_table(doc, [
        ("Slice", "[entity × channel × type]"),
        ("Agent", "AG-__"),
        ("Assessor / date", "[ ]"),
    ])
    add_p(doc, "Score each 1–5 (5 = worse). Likelihood × impact. Treatment owner must be human.", 10, False, SLATE, True)
    headers = ["Risk", "L", "I", "Score", "Treatment", "Owner"]
    risks = [
        "Wrong extract treated as fact",
        "False match / false exception",
        "Missed duplicate (FN)",
        "Unauthorised supplier send",
        "Payment verb / tool creep",
        "Bank-change missed on proposal",
        "Prompt injection via invoice text",
        "Privacy (image / personal data to model)",
        "Model swap without pin",
        "Autonomy applied > ceiling",
        "No fallback if vendor dark",
        "Unsigned savings in exec pack",
    ]
    t = doc.add_table(rows=1 + len(risks), cols=6)
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r, name in enumerate(risks, 1):
        t.rows[r].cells[0].text = name
        t.rows[r].cells[4].text = "[ ]"
        t.rows[r].cells[5].text = "[human]"
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Risk_Assessment.docx")


def meeting():
    doc = Document()
    header_block(doc, "Meeting guide", "Observation, interview, or steering. One purpose per sitting.")
    field_table(doc, [
        ("Type", "Observe / Interview / Steering / Working"),
        ("Purpose (one sentence)", "[ ]"),
        ("Must-leave-with", "[artefact]"),
        ("Do not discuss", "Autonomous payments as a success metric"),
        ("Data rules", "No live personal data unless approved"),
    ])
    add_p(doc, "Opening (5 minutes)", 13, True)
    add_p(doc, "Proof before permission. Today’s artefact. Timebox.")
    add_p(doc, "Questions (interview)", 13, True)
    add_p(doc, "1. Show me the last invoice that broke. What did you actually do?")
    add_p(doc, "2. Which system is true if they disagree?")
    add_p(doc, "3. Who can change a tolerance or a DoA line?")
    add_p(doc, "4. What do you still chase by inbox?")
    add_p(doc, "5. If a model drafted this email, who would you still want to press send?")
    add_p(doc, "Close", 13, True)
    add_p(doc, "Decisions, owners, dates. Parking lot. No unsigned dollars.")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Meeting_Guide.docx")


def impl_plan():
    doc = Document()
    header_block(doc, "Implementation plan (one agent)", "4–6 weeks is illustrative for one bounded L0 agent. Dependencies add time.")
    field_table(doc, [
        ("Agent × slice", "[ ]"),
        ("Human owner", "[ ]"),
        ("Sponsor", "[ ]"),
        ("Target level this increment", "L0 (default)"),
        ("Start / review", "[ ] / [ ]"),
    ])
    add_p(doc, "Phases 0–10 (tick only with evidence)", 13, True)
    for line in [
        "0 Baseline — diagnostic + registry row",
        "1 Discovery — as-done observation",
        "2 Specification — signed charter",
        "3 Access — read path + privacy note",
        "4 Prototype — boxed",
        "5 Historical test — gold first",
        "6 Shadow — no action permissions",
        "7 Controlled execution — only if evidence; allow-list; at-most-once",
        "8 Review — vs baseline",
        "9 Progression — one increment or stay",
        "10 Scale — next slice or next agent, not both",
    ]:
        add_p(doc, "☐  " + line)
    add_p(doc, "Dependencies that extend the calendar", 13, True)
    add_p(doc, "Integrations, DoA quality, image retrieval, vendor contracts, works-council / privacy, year-end freeze, shared-service shift patterns.")
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Implementation_Plan.docx")


def governance_std():
    doc = Document()
    header_block(doc, "Governance standard (summary)", "Full text: Governance/00_GOVERNANCE_FRAMEWORK.md. This is the signed one-pager.")
    add_p(doc, "We will", 13, True)
    add_p(doc, "Name a human Accountable for every agent. Start at L0. Promote with a file. Sample. Log overrides. Pin models. Keep a kill-switch. Retain packets. Terminate access when people leave. Recertify periodically.")
    add_p(doc, "We will not", 13, True)
    add_p(doc, "Grant L3 in a kickoff. Let an agent release payment or approve a bank-change. Treat Agent 10 as fraud detection. Publish unsigned savings. Send invoice images to a model without a privacy path.")
    add_p(doc, "Sign-off", 13, True)
    field_table(doc, [
        ("AP Manager", "Name / date"),
        ("Controller / Controls", "Name / date"),
        ("Finance Systems", "Name / date"),
    ])
    return save(doc, "03_AP_AGENT_OS_PRO/Templates/ER_Governance_Standard.docx")


def main():
    for fn in (sop, charter, discovery, raci, uat, risk, meeting, impl_plan, governance_std):
        print(fn())


if __name__ == "__main__":
    main()
