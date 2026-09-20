#!/usr/bin/env python3
"""Generate Evidence Room Word, PowerPoint, and PDF deliverables."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

INK = HexColor("#1A1F2E")
PAPER = HexColor("#F7F5F0")
MEASURE = HexColor("#2F6F5E")
LINE = HexColor("#D9D4C8")
MUTED = HexColor("#5C6370")

ROOT = Path("/workspace/EVIDENCE_ROOM")
TEMPLATES = ROOT / "03_AP_AGENT_OS_PRO" / "Templates"
TEMPLATES.mkdir(parents=True, exist_ok=True)
PPT_DIR = ROOT / "04_AP_AGENT_OS_TEAM" / "Executive"
PPT_DIR.mkdir(parents=True, exist_ok=True)
WS_DIR = ROOT / "04_AP_AGENT_OS_TEAM" / "Workshop"
WS_DIR.mkdir(parents=True, exist_ok=True)
PDF_DIR = ROOT / "_build" / "pdf"
PDF_DIR.mkdir(parents=True, exist_ok=True)


def set_doc_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1A, 0x1F, 0x2E)


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    run = p.add_run("EVIDENCE ROOM")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x2F, 0x6F, 0x5E)
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1F, 0x2E)
    if subtitle:
        s = doc.add_paragraph(subtitle)
        s.runs[0].italic = True
        s.runs[0].font.color.rgb = RGBColor(0x5C, 0x63, 0x70)
    doc.add_paragraph("—")


def field(doc: Document, label: str, hint: str = "") -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run("_" * 40 + (f"  ({hint})" if hint else ""))


def build_word_templates() -> None:
    doc = Document()
    set_doc_style(doc)
    add_title(doc, "Agent Charter Template", "Use one charter per agent. Approve before shadow mode.")
    for label, hint in [
        ("Agent ID / Name", "e.g. A04 Exception Triage"),
        ("Human Owner", "named role + person"),
        ("Sponsor", "AP Manager / Controller"),
        ("Purpose", "one sentence"),
        ("In scope", ""),
        ("Out of scope / Exclusions", "payment authorisation always human"),
        ("Systems / data access", "least privilege"),
        ("Inputs", ""),
        ("Outputs / artefacts", ""),
        ("Autonomy level at launch", "L0–L4; default L0/L1"),
        ("Promotion criteria", ""),
        ("Demotion / kill triggers", ""),
        ("Approval requirements", ""),
        ("Escalation path", ""),
        ("Primary KPIs", ""),
        ("Control references", "control matrix IDs"),
        ("Audit evidence to retain", ""),
        ("Cost / inference monitoring", ""),
        ("Effective date / version", ""),
    ]:
        field(doc, label, hint)
    doc.add_heading("Sign-off", level=2)
    field(doc, "Owner signature / date")
    field(doc, "Controls signature / date")
    field(doc, "Sponsor signature / date")
    doc.add_paragraph(
        "Disclaimer: Does not authorise autonomous payments or replace ERP controls."
    )
    doc.save(TEMPLATES / "ER_Agent_Charter_Template.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "AP SOP Template", "Standard operating procedure — blank")
    for label in [
        "SOP ID",
        "Process name",
        "Owner",
        "Version",
        "Effective date",
        "Related agents",
        "Systems",
    ]:
        field(doc, label)
    for h, body in [
        ("1. Purpose", "[Why this process exists]"),
        ("2. Scope", "[Inclusions / exclusions]"),
        ("3. Roles (RACI)", "Responsible / Accountable / Consulted / Informed"),
        ("4. Trigger", "[What starts the process]"),
        ("5. Steps", "1. …\n2. …\n3. …"),
        ("6. Decision points", "[If/then rules]"),
        ("7. Exception handling", "[Taxonomy codes + routing]"),
        ("8. Controls", "[Preventive / detective controls]"),
        ("9. Evidence & retention", "[What is retained for audit]"),
        ("10. Metrics", "[KPIs]"),
    ]:
        doc.add_heading(h, level=2)
        doc.add_paragraph(body)
    doc.save(TEMPLATES / "ER_SOP_Template.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "AP Agent RACI Template")
    table = doc.add_table(rows=8, cols=6)
    table.style = "Table Grid"
    hdr = ["Activity", "AP Ops", "AP Manager", "Controller", "IT/Systems", "Agent"]
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    for r, a in enumerate(
        [
            "Define agent charter",
            "Approve autonomy change",
            "Day-to-day monitoring",
            "Exception resolution",
            "Payment authorisation",
            "Model/prompt change",
            "Audit evidence pack",
        ],
        1,
    ):
        table.rows[r].cells[0].text = a
    doc.add_paragraph(
        "Legend: R/A/C/I. Payment authorisation must remain with a human Accountable role."
    )
    doc.save(TEMPLATES / "ER_RACI_Template.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "Process Discovery Interview Guide")
    field(doc, "Interviewee / role")
    field(doc, "Date / interviewer")
    doc.add_heading("Walkthrough questions", level=2)
    for q in [
        "Walk me through a typical invoice from receipt to payment.",
        "Where do most exceptions arise? Approximate mix?",
        "Which systems are touched, in order?",
        "What decisions require judgement vs rules?",
        "What evidence is retained for audit today?",
        "What breaks at month-end?",
        "If you had an analyst who never slept but needed approval, what would you give them first?",
    ]:
        doc.add_paragraph(q, style="List Number")
        doc.add_paragraph("Notes: ________________________________")
    doc.save(TEMPLATES / "ER_Process_Discovery_Guide.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "UAT & Shadow Mode Test Record")
    for label in [
        "Agent",
        "Test ID",
        "Scenario",
        "Data sample reference",
        "Expected result",
        "Actual result",
        "Pass / Fail",
        "Defect / follow-up",
        "Tester / date",
    ]:
        field(doc, label)
    doc.add_paragraph("Shadow mode: agent outputs recorded; no production action rights.")
    doc.save(TEMPLATES / "ER_UAT_Test_Record.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "AP Agent Risk Assessment Template")
    for label in [
        "Agent",
        "Inherent risk",
        "Risk scenarios",
        "Likelihood",
        "Impact",
        "Controls",
        "Residual risk",
        "Owner",
        "Review date",
        "Treatment decision",
    ]:
        field(doc, label)
    doc.save(TEMPLATES / "ER_Risk_Assessment_Template.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(
        doc,
        "AP Agent Governance Standard (Editable)",
        "Organisation-specific adoption of Evidence Room framework",
    )
    for h, b in [
        ("1. Purpose & scope", "[Adopt / adapt]"),
        ("2. Accountability", "Named human owners for each agent. No orphan agents."),
        (
            "3. Autonomy policy",
            "Default L0/L1. L3+ requires certification pack and steering approval. "
            "Payment authorisation excluded from agent autonomy.",
        ),
        (
            "4. Change control",
            "Prompt, model, tool, and workflow changes require tickets and regression tests.",
        ),
        ("5. Monitoring & audit", "[Local requirements]"),
    ]:
        doc.add_heading(h, level=2)
        doc.add_paragraph(b)
    doc.save(TEMPLATES / "ER_Governance_Standard.docx")

    doc = Document()
    set_doc_style(doc)
    add_title(doc, "90-Day Implementation Plan Template")
    field(doc, "Organisation")
    field(doc, "First agent")
    field(doc, "Sponsor")
    doc.add_heading("Days 1–30 — Baseline & design", level=2)
    doc.add_paragraph("• Diagnostic\n• Process discovery\n• Charter & controls\n• KPI baseline")
    doc.add_heading("Days 31–60 — Shadow", level=2)
    doc.add_paragraph("• Historical tests\n• Shadow mode\n• Weekly scorecard")
    doc.add_heading("Days 61–90 — Controlled pilot", level=2)
    doc.add_paragraph(
        "• Limited scope execution (L2)\n• Steering review\n• Go/No-go for expansion"
    )
    doc.save(TEMPLATES / "ER_90_Day_Implementation_Plan.docx")
    print("DOCX", len(list(TEMPLATES.glob("*.docx"))))


def add_blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=PRGB(0x1A, 0x1F, 0x2E)):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = PPt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    return box


INK_R = PRGB(0x1A, 0x1F, 0x2E)
MEAS_R = PRGB(0x2F, 0x6F, 0x5E)
WHITE = PRGB(0xFF, 0xFF, 0xFF)
MUTED_R = PRGB(0x5C, 0x63, 0x70)
SUB = PRGB(0xC5, 0xC9, 0xD1)


def cover(prs, title, subtitle):
    s = add_blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, INK_R)
    add_rect(s, 0, 0, PInches(0.2), prs.slide_height, MEAS_R)
    add_text(s, PInches(0.8), PInches(1.8), PInches(11), PInches(0.4), "EVIDENCE ROOM", 14, True, MEAS_R)
    add_text(s, PInches(0.8), PInches(2.4), PInches(11), PInches(1.5), title, 36, True, WHITE)
    add_text(s, PInches(0.8), PInches(4.2), PInches(11), PInches(1), subtitle, 18, False, SUB)
    add_text(
        s,
        PInches(0.8),
        PInches(6.6),
        PInches(11),
        PInches(0.4),
        "AP Agent Operating System  |  Working materials",
        12,
        False,
        MUTED_R,
    )


def bullets(prs, title, items, footer="Evidence Room — Operating Evidence"):
    s = add_blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, PInches(0.15), MEAS_R)
    add_text(s, PInches(0.7), PInches(0.4), PInches(12), PInches(0.6), title, 28, True, INK_R)
    body = s.shapes.add_textbox(PInches(0.7), PInches(1.3), PInches(12), PInches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.size = PPt(18)
        p.font.color.rgb = INK_R
        p.font.name = "Calibri"
        p.space_after = PPt(12)
    add_text(s, PInches(0.7), PInches(7.0), PInches(12), PInches(0.3), footer, 10, False, MUTED_R)


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    cover(
        prs,
        "AP Agent Transformation\nSteering Pack",
        "From experimentation to governed, measurable agent operations",
    )
    bullets(
        prs,
        "The problem we are solving",
        [
            "AP automation tools process transactions; they rarely define how agents earn responsibility.",
            "Exception work still consumes disproportionate AP capacity.",
            "AI pilots stall without governance, KPIs, and control design.",
            "Finance leaders need an operating system — not another prompt pack.",
        ],
    )
    bullets(
        prs,
        "What Evidence Room is",
        [
            "An AP Agent Operating System: design, govern, measure, and scale AI agents across AP.",
            "ERP-agnostic — sits above Dynamics, SAP, Oracle, NetSuite, Workday and peers.",
            "Human accountability first. Payment authorisation remains human-controlled.",
            "Responsibility is earned: Observe → Recommend → Prepare → Execute in guardrails → Managed autonomy.",
        ],
    )
    bullets(
        prs,
        "The 16-agent AP stack (summary)",
        [
            "Intake, Validation, Matching, Exception Triage, Goods Receipt, PO Quality, Approval",
            "Supplier Resolution, Internal Follow-Up, Duplicate & Anomaly, Vendor Statement, Payment Proposal Review",
            "AP Close, AP Reporting, Root Cause, AP Manager / Orchestrator",
            "Start with one bounded agent — typically Exception Triage, Intake, or Missing GR follow-up.",
        ],
    )
    bullets(
        prs,
        "Governance principles",
        [
            "Named human owner for every agent — no orphans.",
            "Segregation of duties: agents never authorise payments.",
            "Shadow mode before production rights.",
            "Version control for prompts, models, tools, and workflows.",
            "Audit evidence packs for autonomy progression.",
        ],
    )
    bullets(
        prs,
        "How we will measure success",
        [
            "Operational: STP rate, exception ageing, resolution time, human intervention rate",
            "Financial: cost per invoice, hours released (validated vs estimated), inference cost",
            "Risk/control: false positive/negative rates, control breaches, rework, escalations",
            "No vanity metrics. Capacity released ≠ cash saved until redesign is evidenced.",
        ],
    )
    bullets(
        prs,
        "Illustrative 90-day path (one agent)",
        [
            "Days 1–30: Diagnostic, discovery, charter, controls, KPI baseline",
            "Days 31–60: Historical tests + shadow mode",
            "Days 61–90: Controlled pilot (L2), steering go/no-go",
            "Duration depends on systems, data quality, integrations, and governance appetite.",
        ],
    )
    bullets(
        prs,
        "Decisions requested",
        [
            "Confirm sponsor and first-agent scope",
            "Approve shadow-mode data access (read-only)",
            "Agree KPI baseline and success gates",
            "Schedule steering review at end of pilot window",
        ],
    )
    prs.save(PPT_DIR / "ER_CFO_AP_Transformation_Deck.pptx")

    prs2 = Presentation()
    prs2.slide_width = PInches(13.333)
    prs2.slide_height = PInches(7.5)
    cover(prs2, "AP Agent OS Workshop", "Designing governed AI agents for Accounts Payable")
    for title, items in [
        (
            "Workshop objectives",
            [
                "Map current AP pain and exception load",
                "Select first agent candidate",
                "Draft charter, controls, and KPIs",
                "Agree 90-day pilot plan",
            ],
        ),
        (
            "Agenda",
            [
                "Context & principles (30m)",
                "Current-state mapping (45m)",
                "Exception taxonomy exercise (40m)",
                "Agent selection & charter (45m)",
                "Controls & KPIs (30m)",
                "Pilot plan & owners (30m)",
            ],
        ),
        (
            "Rules of engagement",
            [
                "Payment authorisation stays human",
                "No tool-vendor bake-off in this session",
                "Evidence over anecdotes",
                "Decide owners before features",
            ],
        ),
        (
            "Exception taxonomy drill",
            [
                "Code top 20 exception samples",
                "Find concentration and repeat causes",
                "Identify agent-suitable vs human-judgment clusters",
            ],
        ),
        (
            "Responsibility model L0–L4",
            [
                "L0 Observe",
                "L1 Recommend",
                "L2 Prepare (approval gates)",
                "L3 Execute in guardrails",
                "L4 Managed autonomy — earned only",
            ],
        ),
        (
            "Pilot definition",
            [
                "Population limits (entity, supplier set, value cap)",
                "Success gates",
                "Steering date",
            ],
        ),
        ("Close-out", ["Owners named", "Next 10 working days actions", "Risks logged"]),
    ]:
        bullets(prs2, title, items, "Evidence Room Workshop")
    prs2.save(WS_DIR / "ER_AP_Workshop_Deck.pptx")

    prs3 = Presentation()
    prs3.slide_width = PInches(13.333)
    prs3.slide_height = PInches(7.5)
    cover(prs3, "AP Agent Business Case", "Illustrative model — validate with your baselines")
    bullets(
        prs3,
        "How to use this model",
        [
            "Replace assumptions with local volume, FTE, exception, and cost data",
            "Run Conservative / Base / Upside — present Conservative to steering by default",
            "Separate capacity released from cash savings",
            "Re-forecast after shadow mode with measured accuracy and cycle-time deltas",
        ],
    )
    bullets(
        prs3,
        "Evidence we can cite externally (independent)",
        [
            "Ardent Partners: Best-in-Class AP teams show materially lower cost/invoice, faster cycle times, and lower exception rates vs peers (State of ePayables research series).",
            "Use as context for ambition — not as a promise that Evidence Room delivers the same delta.",
        ],
    )
    bullets(
        prs3,
        "Ask",
        [
            "Approve baseline measurement sprint",
            "Fund first-agent shadow pilot",
            "Defer autonomy expansion until gates are met",
        ],
    )
    prs3.save(PPT_DIR / "ER_Business_Case_Deck.pptx")
    print("PPTX ok")


def make_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="H1ER",
            fontName="Helvetica-Bold",
            fontSize=18,
            textColor=INK,
            spaceAfter=6,
            leading=22,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H2ER",
            fontName="Helvetica-Bold",
            fontSize=13,
            textColor=MEASURE,
            spaceAfter=4,
            leading=16,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H3ER",
            fontName="Helvetica-Bold",
            fontSize=11,
            textColor=INK,
            spaceAfter=3,
            leading=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyER",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=INK,
            leading=13,
            alignment=TA_JUSTIFY,
        )
    )
    return styles


def md_to_flowables(text: str, styles):
    story = []
    lines = text.split("\n")
    i = 0
    buf: list[str] = []

    def flush_para() -> None:
        nonlocal buf
        if buf:
            para = " ".join(buf).strip()
            if para:
                para = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", para)
                para = re.sub(r"`(.+?)`", r'<font face="Courier" size="9">\1</font>', para)
                story.append(Paragraph(para, styles["BodyER"]))
                story.append(Spacer(1, 6))
            buf = []

    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            flush_para()
            story.append(Spacer(1, 12))
            story.append(Paragraph(line[2:].strip(), styles["H1ER"]))
            story.append(HRFlowable(width="100%", thickness=1, color=LINE, spaceAfter=10))
        elif line.startswith("## "):
            flush_para()
            story.append(Spacer(1, 10))
            story.append(Paragraph(line[3:].strip(), styles["H2ER"]))
        elif line.startswith("### "):
            flush_para()
            story.append(Spacer(1, 8))
            story.append(Paragraph(line[4:].strip(), styles["H3ER"]))
        elif line.strip() == "---":
            flush_para()
            story.append(HRFlowable(width="100%", thickness=0.5, color=LINE, spaceBefore=8, spaceAfter=8))
        elif re.match(r"^\|.*\|$", line) and i + 1 < len(lines) and re.match(
            r"^\|[-:| ]+\|$", lines[i + 1]
        ):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if re.match(r"^\|[-:| ]+\|$", lines[i]):
                    i += 1
                    continue
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(
                    [
                        Paragraph(re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", c), styles["BodyER"])
                        for c in cells
                    ]
                )
                i += 1
            if rows:
                n = len(rows[0])
                t = Table(rows, hAlign="LEFT", colWidths=[(A4[0] - 50 * mm) / n] * n)
                t.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), INK),
                            ("TEXTCOLOR", (0, 0), (-1, 0), white),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("FONTSIZE", (0, 0), (-1, -1), 8),
                            ("GRID", (0, 0), (-1, -1), 0.4, LINE),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 4),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                            ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [PAPER, white]),
                        ]
                    )
                )
                story.append(t)
                story.append(Spacer(1, 10))
            continue
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            flush_para()
            items = []
            while i < len(lines) and (
                lines[i].strip().startswith("- ") or lines[i].strip().startswith("* ")
            ):
                item = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", lines[i].strip()[2:])
                items.append(
                    ListItem(Paragraph(item, styles["BodyER"]), leftIndent=10, bulletColor=MEASURE)
                )
                i += 1
            story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=15))
            story.append(Spacer(1, 6))
            continue
        elif re.match(r"^\d+\. ", line.strip()):
            flush_para()
            items = []
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                item = re.sub(
                    r"\*\*(.+?)\*\*",
                    r"<b>\1</b>",
                    re.sub(r"^\d+\. ", "", lines[i].strip()),
                )
                items.append(
                    ListItem(Paragraph(item, styles["BodyER"]), leftIndent=10, bulletColor=MEASURE)
                )
                i += 1
            story.append(ListFlowable(items, bulletType="1", leftIndent=15))
            story.append(Spacer(1, 6))
            continue
        elif line.strip() == "":
            flush_para()
        else:
            buf.append(line.strip())
        i += 1
    flush_para()
    return story


def build_pdf_from_md(md_path: Path, pdf_path: Path, title: str, subtitle: str, tier: str) -> None:
    if not md_path.exists():
        print("SKIP missing", md_path)
        return
    styles = make_styles()

    def first_page(c, doc):
        c.setFillColor(INK)
        c.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
        c.setFillColor(MEASURE)
        c.rect(0, 0, 8 * mm, A4[1], fill=1, stroke=0)
        c.setFillColor(MEASURE)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(25 * mm, A4[1] - 40 * mm, "EVIDENCE ROOM")
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 26)
        y = A4[1] - 60 * mm
        for line in title.split("\n"):
            c.drawString(25 * mm, y, line)
            y -= 12 * mm
        c.setFillColor(HexColor("#C5C9D1"))
        c.setFont("Helvetica", 11)
        c.drawString(25 * mm, y - 5 * mm, subtitle)
        c.drawString(25 * mm, 30 * mm, tier)
        c.drawString(25 * mm, 22 * mm, "Operating Evidence  ·  Not legal, tax, or accounting advice")

    def later_pages(c, doc):
        c.setStrokeColor(MEASURE)
        c.setLineWidth(2)
        c.line(15 * mm, A4[1] - 12 * mm, 15 * mm, 16 * mm)
        c.setFont("Helvetica", 8)
        c.setFillColor(MUTED)
        c.drawString(20 * mm, 10 * mm, "Evidence Room  ·  AP Agent OS")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    text = md_path.read_text(encoding="utf-8")
    story = [Spacer(1, 200 * mm), PageBreak()] + md_to_flowables(text, styles)
    doc.build(story, onFirstPage=first_page, onLaterPages=later_pages)
    print("PDF", pdf_path.name, pdf_path.stat().st_size)


def build_pdfs() -> None:
    builds = [
        (
            ROOT / "01_FREE_AP_AI_READINESS" / "01_AP_AI_READINESS_DIAGNOSTIC.md",
            PDF_DIR / "ER_AP_AI_Readiness_Diagnostic.pdf",
            "AP AI Readiness\nDiagnostic",
            "Score your operating readiness for governed AP agents",
            "Tier 0  ·  Free lead magnet",
        ),
        (
            ROOT / "02_AP_AGENT_STARTER" / "01_STARTER_KIT_GUIDE.md",
            PDF_DIR / "ER_AP_Agent_Starter_Kit.pdf",
            "AP Agent\nStarter Kit",
            "Operating model, top agents, taxonomy, and first-90-days path",
            "Tier 1  ·  US$79",
        ),
        (
            ROOT / "03_AP_AGENT_OS_PRO" / "Agent_Library" / "00_AGENT_STACK_OVERVIEW.md",
            PDF_DIR / "ER_AP_Agent_Stack_Overview.pdf",
            "AP Agent Stack\nOverview",
            "Sixteen agents. One operating system.",
            "Tier 2  ·  Professional",
        ),
        (
            ROOT / "03_AP_AGENT_OS_PRO" / "Governance" / "01_AP_AGENT_GOVERNANCE_FRAMEWORK.md",
            PDF_DIR / "ER_AP_Agent_Governance_Framework.pdf",
            "AP Agent\nGovernance Framework",
            "Accountability, controls, and earned autonomy",
            "Tier 2  ·  Professional",
        ),
        (
            ROOT / "03_AP_AGENT_OS_PRO" / "Controls" / "01_EXCEPTION_TAXONOMY.md",
            PDF_DIR / "ER_AP_Exception_Taxonomy.pdf",
            "AP Exception\nTaxonomy",
            "A shared language for exceptions, owners, and agents",
            "Tier 2  ·  Professional",
        ),
        (
            ROOT / "04_AP_AGENT_OS_TEAM" / "00_TEAM_EDITION_OVERVIEW.md",
            PDF_DIR / "ER_AP_Agent_OS_Team_Overview.pdf",
            "AP Agent OS\nTeam Edition",
            "Workshop-ready transformation system for AP leadership teams",
            "Tier 3  ·  US$499",
        ),
        (
            ROOT / "05_CUSTOM_BLUEPRINT" / "01_SERVICE_BROCHURE.md",
            PDF_DIR / "ER_AP_Transformation_Blueprint_Brochure.pdf",
            "AP Transformation\nBlueprint",
            "Productised assessment and roadmap service",
            "Tier 4  ·  US$1,500–3,000",
        ),
        (
            ROOT / "11_BRAND" / "01_BRAND_SYSTEM.md",
            PDF_DIR / "ER_Brand_System.pdf",
            "Brand System",
            "Operating Evidence",
            "Internal  ·  Brand",
        ),
    ]
    for args in builds:
        try:
            build_pdf_from_md(*args)
        except Exception as e:
            print("FAIL", args[1].name, e)

    # Copy into product folders where source PDF exists
    mapping = [
        (
            PDF_DIR / "ER_AP_AI_Readiness_Diagnostic.pdf",
            ROOT / "01_FREE_AP_AI_READINESS" / "ER_AP_AI_Readiness_Diagnostic.pdf",
        ),
        (
            PDF_DIR / "ER_AP_Agent_Starter_Kit.pdf",
            ROOT / "02_AP_AGENT_STARTER" / "ER_AP_Agent_Starter_Kit.pdf",
        ),
        (
            PDF_DIR / "ER_AP_Agent_OS_Team_Overview.pdf",
            ROOT / "04_AP_AGENT_OS_TEAM" / "ER_AP_Agent_OS_Team_Overview.pdf",
        ),
        (
            PDF_DIR / "ER_AP_Transformation_Blueprint_Brochure.pdf",
            ROOT / "05_CUSTOM_BLUEPRINT" / "ER_AP_Transformation_Blueprint_Brochure.pdf",
        ),
        (PDF_DIR / "ER_Brand_System.pdf", ROOT / "11_BRAND" / "ER_Brand_System.pdf"),
    ]
    for src, dst in mapping:
        if src.exists():
            shutil.copy(src, dst)
    pro_pdf = ROOT / "03_AP_AGENT_OS_PRO" / "_pdf"
    pro_pdf.mkdir(exist_ok=True)
    for n in [
        "ER_AP_Agent_Stack_Overview.pdf",
        "ER_AP_Agent_Governance_Framework.pdf",
        "ER_AP_Exception_Taxonomy.pdf",
    ]:
        src = PDF_DIR / n
        if src.exists():
            shutil.copy(src, pro_pdf / n)
    print("PDF packaging complete")


if __name__ == "__main__":
    build_word_templates()
    build_pptx()
    build_pdfs()
    print("All binary packaging complete")
